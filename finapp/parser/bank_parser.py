"""Парсер банковских выписок v2 (сгенерировано tools/build_parser_module.py из bank_parser_v2_clean.ipynb).

Не редактируйте вручную: правки алгоритма вносятся в ноутбук, затем модуль пересобирается.
Модуль хранит состояние в глобальных переменных (CLOCK, LLM_CALLS) — вызывать только через
finapp.parser.service.parse_file, который сериализует запуски блокировкой.
"""
# ruff: noqa
# настройки и импорты
from __future__ import annotations

import csv
import io
import json
import re
import threading
import time
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from itertools import combinations
from pathlib import Path
from typing import Any

import pandas as pd

INPUT_FILE: str | None = None
INPUT_DIR: str | None = None
OUTPUT_DIR: str | None = None

OUTPUT_COLUMNS = [
    "date", "debit_inn", "debit_name", "credit_inn", "credit_name", "amount", "purpose",
]

OWNER_LLM_THRESHOLD = 0.75
TABLE_LLM_THRESHOLD = 0.70
ROW_UNCERTAIN_LIMIT = 60
REPAIR_MAX_ITERATIONS = 2
TEXT_CHUNK_LINES = 40
VERBOSE = True

LLM_CALL_TIMEOUT = 90
LLM_ALGO_BUDGET = 240
LLM_MAX_FAILS = 2
REVIEW_BUDGET = 90
AGENT_BUDGET = 420

GIGACHAT_MODEL = "GigaChat-3-Ultra"
LLM_CLIENT: Any | None = None
CACHE_PATH: str | None = None
LOG_HOOK = None


class RunClock:

    def __init__(self) -> None:
        self.reset()

    def reset(self) -> None:
        self.t0 = time.time()
        self.deadline = {"algo": self.t0 + LLM_ALGO_BUDGET, "agent": None, "review": None}
        self.fails = 0
        self.unavailable: str | None = None
        self.calls = 0
        self.timeline: list[dict[str, Any]] = []

    def elapsed(self) -> float:
        return time.time() - self.t0

    def start_agent(self) -> None:
        self.deadline["agent"] = time.time() + AGENT_BUDGET

    def start_review(self) -> None:
        self.deadline["review"] = time.time() + REVIEW_BUDGET

    def left(self, phase: str) -> float:
        d = self.deadline.get(phase)
        return float("inf") if d is None else d - time.time()

    def mark(self, stage: str, detail: str = "") -> None:
        self.timeline.append({"t": round(self.elapsed(), 1), "stage": stage, "detail": detail})
        log(stage, detail)


CLOCK = RunClock()


def log(stage: str, message: str) -> None:
    if VERBOSE:
        print(f"[+{CLOCK.elapsed():6.1f} с][{stage}] {message}", flush=True)
    if LOG_HOOK is not None:
        LOG_HOOK(f"[{stage}] {message}")


# чтение файлов в сетку
@dataclass(eq=False)
class Cell:
    r: int
    c: int
    r2: int
    c2: int
    text: str
    value: Any = None

    @property
    def width(self) -> int:
        return self.c2 - self.c + 1


class Grid:
    def __init__(self, name: str, cells: list[Cell], source: str):
        self.name = name
        self.source = source
        self.cells = sorted(cells, key=lambda x: (x.r, x.c))
        self._anchor = {(x.r, x.c): x for x in self.cells}
        self._cover: dict[tuple[int, int], Cell] = {}
        self._rows: dict[int, list[Cell]] = defaultdict(list)
        for x in self.cells:
            self._rows[x.r].append(x)
            if x.r2 > x.r or x.c2 > x.c:
                for rr in range(x.r, x.r2 + 1):
                    for cc in range(x.c, x.c2 + 1):
                        self._cover[(rr, cc)] = x
        self.n_rows = max((x.r2 for x in self.cells), default=0) + 1
        self.n_cols = max((x.c2 for x in self.cells), default=0) + 1

    def row(self, r: int) -> list[Cell]:
        return self._rows.get(r, [])

    def at(self, r: int, c: int) -> Cell | None:
        return self._anchor.get((r, c)) or self._cover.get((r, c))

    def text_at(self, r: int, c: int) -> str:
        cell = self.at(r, c)
        return cell.text if cell else ""

    def row_text(self, r: int, sep: str = " | ") -> str:
        return sep.join(x.text for x in self.row(r))

    def rows_range(self) -> range:
        return range(self.n_rows)

    def is_text_only(self) -> bool:
        rows = [r for r in self._rows if self._rows[r]]
        if not rows:
            return True
        multi = sum(1 for r in rows if len(self._rows[r]) > 1)
        return multi / len(rows) < 0.15

    def __repr__(self) -> str:
        return f"Grid({self.name!r}, rows={self.n_rows}, cols={self.n_cols}, cells={len(self.cells)})"


@dataclass
class Document:
    path: Path
    grids: list[Grid]
    warnings: list[str] = field(default_factory=list)


_WS_RE = re.compile(r"[ \t   ]+")


def clean_cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y") if value.time() == datetime.min.time() else value.strftime("%d.%m.%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e21:
            return str(int(value))
        return repr(value)
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    lines = [_WS_RE.sub(" ", line).strip() for line in text.split("\n")]
    out = "\n".join(line for line in lines if line)
    if out and re.fullmatch(r"[▼▲▾▴▽△◀▶⏷⏶↓↑⌄˅˄▿▵\s]+", out):
        return ""
    return out


def _cells_from_matrix(matrix: list[list[Any]], merges: list[tuple[int, int, int, int]] = ()) -> list[Cell]:
    span = {(r, c): (r2, c2) for r, c, r2, c2 in merges}
    covered = {
        (rr, cc)
        for r, c, r2, c2 in merges
        for rr in range(r, r2 + 1)
        for cc in range(c, c2 + 1)
        if (rr, cc) != (r, c)
    }
    cells = []
    for r, row in enumerate(matrix):
        for c, value in enumerate(row):
            if (r, c) in covered:
                continue
            text = clean_cell_text(value)
            if not text:
                continue
            r2, c2 = span.get((r, c), (r, c))
            cells.append(Cell(r, c, r2, c2, text, value))
    return cells


def _read_xlsx(path: Path) -> list[Grid]:
    import openpyxl

    read_only = path.stat().st_size > 30 * 1024 * 1024
    wb = openpyxl.load_workbook(path, data_only=True, read_only=read_only)
    grids = []
    try:
        for ws in wb.worksheets:
            matrix = [list(row) for row in ws.iter_rows(values_only=True)]
            merges = []
            if not read_only:
                for m in ws.merged_cells.ranges:
                    merges.append((m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1))
            cells = _cells_from_matrix(matrix, merges)
            if cells:
                grids.append(Grid(ws.title, cells, f"xlsx:{ws.title}"))
    finally:
        wb.close()
    return grids


def _read_xls(path: Path) -> list[Grid]:
    try:
        import xlrd
    except ImportError:
        converted = _convert_with_libreoffice(path, "xlsx")
        if converted is None:
            raise RuntimeError("Для .xls нужен пакет xlrd (pip install xlrd) или LibreOffice")
        return _read_xlsx(converted)
    book = xlrd.open_workbook(str(path), formatting_info=True)
    grids = []
    for sheet in book.sheets():
        matrix = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    except Exception:
                        pass
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    value = None
                row.append(value)
            matrix.append(row)
        merges = [(r1, c1, r2 - 1, c2 - 1) for r1, r2, c1, c2 in sheet.merged_cells]
        cells = _cells_from_matrix(matrix, merges)
        if cells:
            grids.append(Grid(sheet.name, cells, f"xls:{sheet.name}"))
    return grids


def _convert_with_libreoffice(path: Path, target: str) -> Path | None:
    import shutil
    import subprocess
    import tempfile

    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        return None
    out_dir = Path(tempfile.mkdtemp(prefix="stmt_convert_"))
    subprocess.run([office, "--headless", "--convert-to", target, "--outdir", str(out_dir), str(path)],
                   check=False, capture_output=True, timeout=180)
    result = out_dir / (path.stem + "." + target)
    return result if result.exists() else None


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _read_csv(path: Path) -> list[Grid]:
    text = _read_text_file(path)
    sample = "\n".join(text.splitlines()[:50])
    delimiter = max([";", ",", "\t", "|"], key=lambda d: sample.count(d))
    matrix = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return [Grid(path.stem, _cells_from_matrix(matrix), "csv")]


def _read_txt(path: Path) -> list[Grid]:
    lines = _read_text_file(path).splitlines()
    return [Grid(path.stem, _cells_from_matrix([[line] for line in lines]), "txt")]


def _html_table_matrix(table) -> tuple[list[list[Any]], list[tuple[int, int, int, int]]]:
    matrix: dict[tuple[int, int], str] = {}
    merges = []
    occupied: set[tuple[int, int]] = set()
    for r, tr in enumerate(table.find_all("tr")):
        c = 0
        for td in tr.find_all(["td", "th"], recursive=False):
            while (r, c) in occupied:
                c += 1
            colspan = int(td.get("colspan", 1) or 1)
            rowspan = int(td.get("rowspan", 1) or 1)
            matrix[(r, c)] = td.get_text("\n", strip=True)
            for rr in range(r, r + rowspan):
                for cc in range(c, c + colspan):
                    occupied.add((rr, cc))
            if colspan > 1 or rowspan > 1:
                merges.append((r, c, r + rowspan - 1, c + colspan - 1))
            c += colspan
    if not matrix:
        return [], []
    n_rows = max(r for r, _ in occupied) + 1
    n_cols = max(c for _, c in occupied) + 1
    grid = [[matrix.get((r, c)) for c in range(n_cols)] for r in range(n_rows)]
    return grid, merges


def _read_html(path: Path) -> list[Grid]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_read_text_file(path), "lxml")
    for tag in soup(["script", "style"]):
        tag.decompose()
    tables = [t for t in soup.find_all("table") if t.find_parent("table") is None]
    parsed = {}
    for i, table in enumerate(tables):
        parsed[i] = _html_table_matrix(table)
        table.replace_with(soup.new_string(f"\n@@TABLE{i}@@\n"))
    matrix: list[list[Any]] = []
    merges: list[tuple[int, int, int, int]] = []
    for line in soup.get_text("\n").splitlines():
        line = line.strip()
        if not line:
            continue
        m = re.fullmatch(r"@@TABLE(\d+)@@", line)
        if m:
            rows, table_merges = parsed[int(m.group(1))]
            offset = len(matrix)
            matrix.extend(rows)
            merges.extend((r + offset, c, r2 + offset, c2) for r, c, r2, c2 in table_merges)
        else:
            matrix.append([line])
    return [Grid(path.stem, _cells_from_matrix(matrix, merges), "html")]


def _docx_text(element, qn) -> str:
    parts = []
    for node in element.iter():
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag in (qn("w:br"), qn("w:cr")):
            parts.append("\n")
        elif node.tag == qn("w:tab"):
            parts.append(" ")
    return "".join(parts)


def _read_docx(path: Path) -> list[Grid]:
    import docx
    from docx.oxml.ns import qn

    document = docx.Document(str(path))
    cells: list[Cell] = []
    row_index = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = clean_cell_text(_docx_text(child, qn))
            if text:
                cells.append(Cell(row_index, 0, row_index, 0, text, text))
                row_index += 1
        elif child.tag == qn("w:tbl"):
            open_vertical: dict[int, Cell] = {}
            for tr in child.findall(qn("w:tr")):
                col = 0
                for tc in tr.findall(qn("w:tc")):
                    props = tc.find(qn("w:tcPr"))
                    span = 1
                    vmerge = None
                    if props is not None:
                        grid_span = props.find(qn("w:gridSpan"))
                        if grid_span is not None:
                            span = int(grid_span.get(qn("w:val"), 1))
                        vm = props.find(qn("w:vMerge"))
                        if vm is not None:
                            vmerge = vm.get(qn("w:val"), "continue")
                    text = clean_cell_text("\n".join(_docx_text(p, qn) for p in tc.findall(qn("w:p"))))
                    if vmerge == "continue" and col in open_vertical:
                        open_vertical[col].r2 = row_index
                    else:
                        cell = Cell(row_index, col, row_index, col + span - 1, text, text)
                        if text:
                            cells.append(cell)
                        if vmerge == "restart":
                            open_vertical[col] = cell
                        else:
                            open_vertical.pop(col, None)
                    col += span
                row_index += 1
    return [Grid(path.stem, cells, "docx")]


def _pdf_layout_rows(words: list[dict]) -> list[list[tuple[float, str]]]:
    lines: list[list[dict]] = []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if lines and abs(lines[-1][0]["top"] - w["top"]) <= 3:
            lines[-1].append(w)
        else:
            lines.append([w])
    result = []
    for line in lines:
        line.sort(key=lambda w: w["x0"])
        chunks: list[list[dict]] = [[line[0]]]
        for prev, cur in zip(line, line[1:]):
            char_w = (prev["x1"] - prev["x0"]) / max(len(prev["text"]), 1)
            if cur["x0"] - prev["x1"] > max(3 * char_w, 8):
                chunks.append([cur])
            else:
                chunks[-1].append(cur)
        result.append([(chunk[0]["x0"], " ".join(w["text"] for w in chunk)) for chunk in chunks])
    return result


def _read_pdf(path: Path) -> list[Grid]:
    import pdfplumber

    matrix: list[list[Any]] = []
    merges: list[tuple[int, int, int, int]] = []
    empty_pages = 0
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            tables = page.find_tables()
            items: list[tuple[float, list[list[Any]], bool]] = []
            bboxes = [t.bbox for t in tables]
            for table in tables:
                rows = table.extract()
                items.append((table.bbox[1], [[None if v is None else clean_cell_text(v) for v in row] for row in rows], True))

            def outside(obj, boxes=bboxes):
                x = (obj["x0"] + obj["x1"]) / 2
                y = (obj["top"] + obj["bottom"]) / 2
                return not any(b[0] <= x <= b[2] and b[1] <= y <= b[3] for b in boxes)

            words = [w for w in page.extract_words(keep_blank_chars=False, use_text_flow=False) if outside(w)]
            if not words and not tables:
                empty_pages += 1
            layout = _pdf_layout_rows(words)
            anchors: list[float] = []
            for x0 in sorted(x for row in layout for x, _ in row):
                if not anchors or x0 - anchors[-1] > 12:
                    anchors.append(x0)
            line_tops = sorted({round(w["top"]) for w in words})
            for row, top in zip(layout, line_tops):
                cells_row: list[Any] = [None] * max(len(anchors), 1)
                for x0, text in row:
                    col = max(i for i, a in enumerate(anchors) if a <= x0 + 0.5) if anchors else 0
                    cells_row[col] = text if cells_row[col] is None else f"{cells_row[col]} {text}"
                items.append((top, [cells_row], False))
            for _, rows, is_table in sorted(items, key=lambda item: item[0]):
                for row in rows:
                    offset = len(matrix)
                    if is_table:
                        for c, value in enumerate(row):
                            if value is not None:
                                end = c
                                while end + 1 < len(row) and row[end + 1] is None:
                                    end += 1
                                if end > c:
                                    merges.append((offset, c, offset, end))
                    matrix.append(row)
    grids = [Grid(path.stem, _cells_from_matrix(matrix, merges), "pdf")]
    if empty_pages and not grids[0].cells:
        raise RuntimeError("PDF без текстового слоя (скан): нужен OCR")
    return grids


READERS = {
    ".xlsx": _read_xlsx, ".xlsm": _read_xlsx, ".xls": _read_xls,
    ".csv": _read_csv, ".txt": _read_txt, ".html": _read_html, ".htm": _read_html,
    ".docx": _read_docx, ".pdf": _read_pdf,
}


def normalize_serial_dates(grid: Grid) -> int:
    by_col: dict[int, list[Cell]] = defaultdict(list)
    for cell in grid.cells:
        v = cell.value
        if isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer() and 36526 <= v <= 55153:
            by_col[cell.c].append(cell)
    fixed = 0
    for c, cells in by_col.items():
        if len(cells) < 3:
            continue
        numeric_all = [x for x in grid.cells if x.c == c and isinstance(x.value, (int, float)) and not isinstance(x.value, bool)]
        if len(cells) < 0.8 * len(numeric_all):
            continue
        first = min(x.r for x in cells)
        head = [x for r in range(max(0, first - 12), first) for x in grid.row(r) if x.c <= c <= x.c2]
        if not any(re.search(r"дат|date", x.text, re.I) for x in head):
            continue
        for x in cells:
            d = datetime(1899, 12, 30) + timedelta(days=int(x.value))
            x.value, x.text = d, d.strftime("%d.%m.%Y")
            fixed += 1
    return fixed


def read_document(path: str | Path) -> Document:
    path = Path(path)
    reader = READERS.get(path.suffix.lower())
    if reader is None:
        raise ValueError(f"Формат {path.suffix} не поддерживается")
    grids = [g for g in reader(path) if g.cells]
    for g in grids:
        if normalize_serial_dates(g):
            log("чтение", f"{g.name}: даты, записанные числом Excel, переведены в даты")
    log("чтение", f"{path.name}: {len(grids)} лист(ов): " + ", ".join(map(repr, grids)))
    return Document(path=path, grids=grids)


# типы значений и валидаторы
def inn_ok(value: str) -> bool:
    if not value or not value.isdigit() or len(value) not in (10, 12) or set(value) == {"0"}:
        return False
    d = [int(ch) for ch in value]

    def check(weights: list[int]) -> int:
        return sum(w * x for w, x in zip(weights, d)) % 11 % 10

    if len(value) == 10:
        return check([2, 4, 10, 3, 5, 9, 4, 6, 8]) == d[9]
    return (check([7, 2, 4, 10, 3, 5, 9, 4, 6, 8]) == d[10]
            and check([3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8]) == d[11])


def fix_inn(value: str) -> str:
    if not value or not value.isdigit() or value.startswith("0") or len(value) not in (8, 9, 10, 11) or inn_ok(value):
        return value
    for zeros in (1, 2):
        cand = "0" * zeros + value
        if len(cand) in (10, 12) and inn_ok(cand):
            return cand
    return value


def kpp_ok(value: str) -> bool:
    return bool(re.fullmatch(r"\d{4}[\dA-Z]{2}\d{3}", value or "")) and value[:4] != "0000"


def bik_ok(value: str) -> bool:
    return bool(re.fullmatch(r"0\d{8}", value or ""))


def account_ok(value: str, bik: str | None = None) -> bool:
    if not re.fullmatch(r"\d{20}", value or ""):
        return False
    if bik and bik_ok(bik):
        prefix = bik[-3:] if bik[6:] not in ("000", "001", "002") else "0" + bik[4:6]
        digits = prefix + value
        weights = [7, 1, 3] * 8
        return sum(int(a) * w for a, w in zip(digits, weights)) % 10 == 0
    return value[0] in "1234"


_MONTHS = {
    "янв": 1, "фев": 2, "мар": 3, "апр": 4, "мая": 5, "май": 5, "июн": 6, "июл": 7,
    "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12,
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8,
    "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}
_DATE_PATTERNS = [
    (re.compile(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{4})(?:[ T]\d{1,2}:\d{2}(?::\d{2})?)?$"), "dmy"),
    (re.compile(r"^(\d{1,2})\.(\d{1,2})\.(\d{2})(?:\s+\d{1,2}:\d{2}(?::\d{2})?)?$"), "dmy2"),
    (re.compile(r"^(\d{4})[./-](\d{1,2})[./-](\d{1,2})(?:[ T]\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?)?$"), "ymd"),
    (re.compile(r"^(19|20)\d{2}(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])$"), "ymdcompact"),
    (re.compile(r"^(\d{1,2})\s+([а-яa-z]{3})[а-яa-z]*\.?\s+(\d{4})(?:\s*г\.?)?$", re.I), "dMy"),
]


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_cell_text(value).split("\n")[0].strip()
    if not text or len(text) > 25:
        return None
    for pattern, kind in _DATE_PATTERNS:
        m = pattern.match(text)
        if not m:
            continue
        try:
            if kind == "dmy":
                d, mth, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
            elif kind == "dmy2":
                d, mth, y = int(m.group(1)), int(m.group(2)), 2000 + int(m.group(3))
            elif kind == "ymd":
                y, mth, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
            elif kind == "ymdcompact":
                y, mth, d = int(m.group(0)[:4]), int(m.group(0)[4:6]), int(m.group(0)[6:8])
            else:
                mth = _MONTHS.get(m.group(2).lower()[:3])
                if mth is None:
                    return None
                d, y = int(m.group(1)), int(m.group(3))
            if not 1990 <= y <= 2100:
                return None
            return date(y, mth, d)
        except ValueError:
            return None
    return None


_AMOUNT_RE = re.compile(
    r"^(?P<sign>[-−–+])?\s*(?P<paren>\()?\s*(?P<num>\d{1,3}(?:[  '’]\d{3})+|\d+)"
    r"(?:(?P<sep>[.,-])(?P<frac>\d{1,2}))?\s*\)?\s*(?P<tail>[-−])?\s*"
    r"(?:руб(?:л\w*)?\.?|rub|rur|₽|р\.?)?\s*(?:(?P<dk>д|к|дт|кт|cr|dr)\.?)?$",
    re.I,
)


_EN_THOUSANDS_RE = re.compile(r"(?<![\d.,])(\d{1,3}(?:,\d{3})+)(\.\d{1,2})(?!\d)")


def parse_amount(value: Any) -> Decimal | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"))
        except InvalidOperation:
            return None
    text = clean_cell_text(value).replace("\n", " ").strip()
    if not text or len(text) > 30:
        return None
    text = _EN_THOUSANDS_RE.sub(lambda mm: mm.group(1).replace(",", "") + mm.group(2), text)
    m = _AMOUNT_RE.match(text)
    if not m:
        return None
    num = re.sub(r"[  '’]", "", m.group("num"))
    frac = m.group("frac") or "0"
    if m.group("sep") == "-" and len(frac) != 2:
        return None
    try:
        amount = Decimal(f"{num}.{frac}")
    except InvalidOperation:
        return None
    negative = bool(m.group("sign") in ("-", "−", "–") or m.group("paren") or m.group("tail"))
    return -amount if negative else amount


def looks_like_identifier(text: str) -> bool:
    digits = re.sub(r"\s", "", text)
    return digits.isdigit() and len(digits) >= 9 and not re.search(r"[.,]", text)


OPF_SHORT = r"ООО|ОАО|ЗАО|ПАО|АО|НАО|ИП|НКО|АНО|ГУП|МУП|ФГУП|ФГБУ|ГБУ|МБУ|ФКУ|КФХ|СНТ|ТСЖ|ПК|LLC|LTD|JSC|INC|GMBH"
OPF_FULL = (
    r"общество\s+с\s+ограниченной\s+ответственностью|акционерное\s+общество|"
    r"публичное\s+акционерное\s+общество|индивидуальный\s+предприниматель|"
    r"некоммерческая\s+организация|товарищество|кооператив|учреждение|"
    r"предприятие|фонд|ассоциация|управление\s+федерального\s+казначейства|уфк|казначейство"
)
OPF_RE = re.compile(rf"(?<![А-ЯЁA-Zа-яё])(?:{OPF_SHORT})(?![А-ЯЁA-Zа-яё])|{OPF_FULL}", re.I)
PERSON_RE = re.compile(r"^[А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?\s+[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+(?:вич|вна|ична|оглы|кызы)$")
PERSON_INITIALS_RE = re.compile(r"^[А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.\s?[А-ЯЁ]\.$")
PURPOSE_WORD_RE = re.compile(
    r"(?<![а-яё])(?:оплат\w*|предоплат\w*|перевод\w*|возврат\w*|погашени\w*|выплат\w*|заработн\w*|аренд\w*|"
    r"комисси\w*|пополнени\w*|согласно|по\s+(?:сч[её]ту|договор\w*|накладн\w*|акт\w*)|сч[её]т[\s-]*фактур\w*|"
    r"(?:без\s+)?ндс|за\s+(?:услуг\w*|товар\w*|работ\w*|поставк\w*|февраль|январь|март|апрель|май|июнь|июль|август|"
    r"сентябрь|октябрь|ноябрь|декабрь))(?![а-яё])", re.I)


def looks_like_purpose(text: str | None) -> bool:
    t = " ".join((text or "").split())
    if len(t) < 12:
        return False
    head = OPF_RE.match(t)
    if head or PERSON_RE.match(t) or PERSON_INITIALS_RE.match(t):
        return False
    if PURPOSE_WORD_RE.search(t):
        return True
    return len(t) > 70 and len(t.split()) >= 8 and not OPF_RE.search(t)


BANK_WORD_RE = re.compile(r"банк(?!ет)|\bбик\b|корр?\.?\s*сч|к/с|swift|кредитн\w+\s+организ|(?<![а-яё])втб(?![а-яё])", re.I)


def looks_like_party_name(text: str) -> bool:
    t = text.strip()
    return bool(OPF_RE.search(t) or PERSON_RE.match(t) or PERSON_INITIALS_RE.match(t)
                or re.search(r"[«\"].{2,}[»\"]", t))


_ID_LABELS = {
    "inn": r"инн",
    "kpp": r"кпп",
    "account": r"(?:р/с|р\.с\.|рс|сч[её]т|счет\s*№|л/с|account)",
    "bik": r"бик",
}


def split_digit_run(digits: str) -> list[tuple[str, str]]:
    n = len(digits)
    best: dict[int, tuple[float, list[tuple[str, str]]]] = {n: (0.0, [])}
    for i in range(n - 1, -1, -1):
        options = []
        for size, kind in ((12, "inn"), (10, "inn"), (20, "account"), (9, "kpp"), (9, "bik")):
            j = i + size
            if j > n or j not in best:
                continue
            piece = digits[i:j]
            ok = {"inn": inn_ok, "kpp": kpp_ok, "account": account_ok, "bik": bik_ok}[kind](piece)
            if ok:
                gain = {"inn": 3.0, "account": 1.5, "kpp": 1.0, "bik": 0.8}[kind]
                options.append((best[j][0] + gain, [(kind, piece)] + best[j][1]))
        if options:
            best[i] = max(options, key=lambda o: o[0])
    return best.get(0, (0.0, []))[1]


def find_ids_in_text(text: str) -> dict[str, list[str]]:
    found: dict[str, list[str]] = defaultdict(list)
    for m in re.finditer(r"(?<![\d.,])(\d{10}|\d{12}|\d{20})(?!\d)", re.sub(r"[\u00a0\n\r\t]", " ", text)):
        val = m.group(1)
        if len(val) == 20 and account_ok(val) and val not in found["account"]:
            found["account"].append(val)
        elif len(val) in (10, 12) and inn_ok(val) and val not in found["inn"]:
            found["inn"].append(val)
    compact = re.sub(r"(?<=\d)[  ](?=\d)", "", text)
    for kind, label in _ID_LABELS.items():
        for m in re.finditer(rf"{label}\s*[:№#]?\s*(\d{{9,20}})", compact, re.I):
            val = m.group(1)
            ok = {"inn": inn_ok, "kpp": kpp_ok, "account": account_ok, "bik": bik_ok}[kind](val)
            if ok and val not in found[kind]:
                found[kind].append(val)
    for m in re.finditer(r"(\d{10}|\d{12})\s*/\s*(\d{9})", compact):
        if inn_ok(m.group(1)) and m.group(1) not in found["inn"]:
            found["inn"].append(m.group(1))
        if kpp_ok(m.group(2)) and m.group(2) not in found["kpp"]:
            found["kpp"].append(m.group(2))
    for m in re.finditer(r"(?<!\d)(\d{19,41})(?!\d)", compact):
        if len(m.group(1)) == 20:
            continue
        for kind, val in split_digit_run(m.group(1)):
            if val not in found[kind]:
                found[kind].append(val)
    for m in re.finditer(r"(?<!\d)(\d{10}|\d{12}|\d{20})(?!\d)", compact):
        val = m.group(1)
        if len(val) == 20 and account_ok(val) and val not in found["account"]:
            found["account"].append(val)
        elif len(val) in (10, 12) and inn_ok(val) and val not in found["inn"]:
            found["inn"].append(val)
    return dict(found)


def cell_kinds(text: str, value: Any = None) -> set[str]:
    kinds: set[str] = set()
    if not text:
        return {"empty"}
    t = text.strip()
    if t in ("-", "—", "–", "--", "х", "x"):
        return {"zero"}
    digits = re.sub(r"\s", "", t)
    if parse_date(value if value is not None else t) is not None:
        kinds.add("date")
    pair = re.fullmatch(r"(\d{10}|\d{12})\s*[/\\,;\n]\s*(\d{9})", t)
    if pair and inn_ok(pair.group(1)):
        kinds.add("inn")
    if digits.isdigit():
        if len(digits) in (10, 12) and inn_ok(digits):
            kinds.add("inn")
        elif len(digits) in (8, 9, 10, 11) and fix_inn(digits) != digits:
            kinds.add("inn_nozero")
        if len(digits) == 9 and kpp_ok(digits):
            kinds.add("kpp")
        if len(digits) == 9 and bik_ok(digits):
            kinds.add("bik")
        if len(digits) == 20:
            kinds.add("account")
        if len(digits) <= 7:
            kinds.add("small_int")
    if not looks_like_identifier(t) and "date" not in kinds and parse_amount(value if isinstance(value, (int, float)) else t) is not None:
        kinds.add("amount")
    if re.search(r"[А-Яа-яЁёA-Za-z]{2,}", t):
        kinds.add("text")
        if looks_like_party_name(t) and not BANK_WORD_RE.search(t):
            kinds.add("name")
        ids = find_ids_in_text(t)
        if ids.get("inn"):
            kinds.add("composite")
        if len(t) >= 25:
            kinds.add("long_text")
    return kinds


# таблица операций и шапка
class KindCache:

    def __init__(self) -> None:
        self._cache: dict[int, set[str]] = {}

    def __call__(self, cell: Cell | None) -> set[str]:
        if cell is None:
            return {"empty"}
        key = id(cell)
        if key not in self._cache:
            self._cache[key] = cell_kinds(cell.text, cell.value)
        return self._cache[key]


@dataclass
class ColumnProfile:
    col: int
    header: str
    n: int
    fill: float
    frac: dict[str, float]
    negatives: float
    decimals: float
    distinct: float
    avg_len: float
    examples: list[str]

    def card(self) -> str:
        top = ", ".join(f"{k}={v:.0%}" for k, v in sorted(self.frac.items(), key=lambda kv: (-kv[1], kv[0])) if v >= 0.05)
        ex = " ; ".join(e[:40] for e in self.examples[:4])
        return (f"C{self.col}: шапка «{self.header or '—'}» | заполнено {self.fill:.0%} | типы: {top or 'пусто'}"
                f" | отрицательных {self.negatives:.0%} | примеры: {ex}")


@dataclass
class Region:
    grid: Grid
    header_rows: list[int]
    first_anchor: int
    last_anchor: int
    end_row: int
    columns: list[int]
    headers: dict[int, str]
    anchors: list[int]
    profiles: dict[int, ColumnProfile] = field(default_factory=dict)
    kinds: KindCache = field(default_factory=KindCache)

    @property
    def start_row(self) -> int:
        return self.header_rows[0] if self.header_rows else self.first_anchor


def _row_has(grid: Grid, kinds: KindCache, r: int, kind: str, own_only: bool = True) -> list[int]:
    cols = []
    for c in range(grid.n_cols):
        cell = grid.at(r, c)
        if cell is None or cell.c != c:
            continue
        if own_only and cell.r != r:
            continue
        if kind in kinds(cell):
            cols.append(c)
    return cols


def _date_cols(grid: Grid, kinds: KindCache, r: int) -> list[int]:
    return [c for c in range(grid.n_cols)
            if (cell := grid.at(r, c)) is not None and cell.c == c and "date" in kinds(cell)]


def find_anchor_rows(grid: Grid, kinds: KindCache) -> tuple[list[int], int | None]:
    votes: Counter = Counter()
    per_row: dict[int, list[int]] = {}
    for r in range(grid.n_rows):
        dcols = _date_cols(grid, kinds, r)
        if not dcols:
            continue
        own = grid.row(r)
        if own and len(own) <= 6 and FOOTER_WORD_RE.search(own[0].text):
            continue
        amounts = set(_row_has(grid, kinds, r, "amount")) | set(_row_has(grid, kinds, r + 1, "amount"))
        amounts -= set(dcols)
        if amounts:
            per_row[r] = dcols
            for c in dcols:
                votes[c] += 1
    if not votes:
        return [], None
    date_col = votes.most_common(1)[0][0]
    anchors = [r for r, cols in per_row.items() if date_col in cols]
    if anchors:
        sig = {c.c for c in grid.row(anchors[-1])}
        r, quiet = anchors[-1] + 1, 0
        tail = []
        while r < grid.n_rows and sig and quiet <= 6:
            row = grid.row(r)
            if row and (FOOTER_WORD_RE.search(grid.row_text(r)) or any(c.c == date_col for c in row)):
                break
            if row and _row_has(grid, kinds, r, "amount"):
                if any("text" in kinds(c) for c in row) and len({c.c for c in row} & sig) / len(sig) >= 0.6:
                    tail.append(r)
                    quiet = 0
                else:
                    break
            else:
                quiet += 1
            r += 1
        anchors += tail
    extra = []
    for r in anchors:
        cell = grid.at(r, date_col)
        if cell is not None and cell.r2 > cell.r:
            extra.extend(rr for rr in range(cell.r + 1, cell.r2 + 1) if _row_has(grid, kinds, rr, "amount"))
    return sorted(set(anchors) | set(extra)), date_col


HEADER_WORD_RE = re.compile(
    r"дата|сумм|дебет|кредит|назнач|плательщ|получат|контрагент|корреспондент|инн|кпп|бик|сч[её]т|"
    r"наименован|№|номер|оборот|остаток|приход|расход|списан|зачисл|операц|документ|описан|"
    r"date|amount|debit|credit|description|purpose|payer|payee|balance|account",
    re.I,
)


def _is_numbering_row(cells: list[Cell]) -> bool:
    values = [c.text for c in cells]
    nums = [int(v) for v in values if v.isdigit() and len(v) <= 2]
    others = len(values) - len(nums)
    if len(nums) < 3 or not (others == 0 or (len(nums) >= 4 and others <= 3 and others / len(values) <= 0.34)):
        return False
    return nums == sorted(nums) and nums[0] in (0, 1)


def detect_header(grid: Grid, kinds: KindCache, first_anchor: int, data_cols: list[int]) -> list[int]:
    header: list[int] = []
    r = first_anchor - 1
    skipped = 0
    dcols = set(data_cols)

    def _blank(rr: int) -> bool:
        return not any(c.text.strip() and any(x in dcols for x in range(c.c, c.c2 + 1)) for c in grid.row(rr))

    def _service(rr: int) -> bool:
        row = grid.row(rr)
        return 0 < len(row) <= 4 and bool(FOOTER_WORD_RE.search(row[0].text))

    while r >= 0 and skipped < 10 and (_blank(r) or _service(r)):
        r -= 1
        skipped += 1
    while r >= 0 and len(header) < 12:
        own = grid.row(r)
        covering = {grid.at(r, c) for c in data_cols} - {None}
        if not own and not covering:
            break
        texts = [c for c in covering]
        if not texts:
            break
        if _is_numbering_row(own):
            header.insert(0, r)
            r -= 1
            continue
        if own and len(own) <= 3 and own[0].text.rstrip().endswith(":") and any(
                re.fullmatch(r"[\d ]{9,22}", c.text.strip()) or OPF_RE.search(c.text) or "kpp" in kinds(c) or "inn" in kinds(c) for c in own[1:]):
            break
        has_data = any("amount" in kinds(c) and "small_int" not in kinds(c) for c in own) or \
            any("date" in kinds(c) for c in own)
        coverage = sum(1 for c in data_cols if grid.at(r, c) is not None) / max(len(data_cols), 1)
        header_words = sum(1 for c in texts if HEADER_WORD_RE.search(unhyphenate(c.text)))
        strong_header = header_words >= 2 and len(own) >= 2
        if has_data or (coverage < 0.4 and not strong_header) or (header_words == 0 and len(own) < 3):
            break
        header.insert(0, r)
        r -= 1
    return header


def unhyphenate(text: str) -> str:
    return re.sub(r"(?<=[а-яё])[-­]\s*(?=[а-яё])", "", text)


def header_paths(grid: Grid, header_rows: list[int], columns: list[int]) -> dict[int, str]:
    paths = {}
    for c in columns:
        parts: list[str] = []
        for r in header_rows:
            cell = grid.at(r, c)
            if cell is None:
                continue
            text = unhyphenate(cell.text.replace("\n", " "))
            if text.isdigit() and len(text) <= 2:
                continue
            if not parts or parts[-1] != text:
                parts.append(text)
        paths[c] = " › ".join(parts)
    return paths


FOOTER_WORD_RE = re.compile(r"итог|всего|оборот|остат|сальдо|исходящ|входящ|перенос|страниц", re.I)


def leaf_spans(grid: Grid, header_rows: list[int]) -> list[tuple[int, int]]:
    cells = [c for r in header_rows for c in grid.row(r) if not (c.text.isdigit() and len(c.text) <= 2)]
    leaves = []
    for c in cells:
        deeper = any(o is not c and o.r > c.r2 and o.c <= c.c2 and o.c2 >= c.c for o in cells)
        if not deeper:
            leaves.append((c.c, c.c2))
    return sorted(set(leaves))


def snap_grid(grid: Grid, header_rows: list[int], first_row: int, last_row: int) -> Grid | None:
    spans = leaf_spans(grid, header_rows)
    if len(spans) < 3:
        return None
    data = [c for r in range(first_row, last_row + 1) for c in grid.row(r)]
    if not data:
        return None
    span_of = {}
    for a, z in spans:
        for x in range(a, z + 1):
            span_of.setdefault(x, (a, z))
    col_count = Counter(c.c for c in data if c.c not in span_of)
    left_edge, right_edge = min(a for a, _ in spans), max(z for _, z in spans)
    rows_n = last_row - first_row + 1
    moved = dropped = 0
    new_cells = []
    occupied = {(c.r, c.c) for c in data}
    kinds = KindCache()
    data_rows = {c.r for c in data if "date" in kinds(c)}
    for c in grid.cells:
        if not (first_row <= c.r <= last_row):
            new_cells.append(c)
            continue
        span = span_of.get(c.c)
        if span and c.c2 <= span[1] and c.c != span[0] and (c.r, span[0]) not in occupied:
            occupied.add((c.r, span[0]))
            new_cells.append(Cell(c.r, span[0], c.r2, span[1] if c.c2 > c.c or span[1] > span[0] else span[0], c.text, c.value))
            moved += 1
        elif span is None and all(x not in span_of for x in range(c.c, c.c2 + 1)) \
                and col_count[c.c] / max(rows_n, 1) < 0.25 \
                and not (kinds(c) & {"date", "inn", "account", "bik"}) \
                and not ("amount" in kinds(c) and re.search(r"[.,-]\d{2}\s*(?:₽|р\.?|руб\.?|rub)?$", c.text, re.I) and len(c.text) >= 4) \
                and ((c.r in data_rows and len(c.text) <= 12)
                     or ((c.c > right_edge or c.c2 < left_edge) and len(c.text) <= 40)):
            dropped += 1
        else:
            new_cells.append(c)
    if not (moved or dropped):
        return None
    log("сетка", f"{grid.name}: {moved} значений перенесено в колонки по шапке, {dropped} шумовых ячеек вне шапки убрано")
    return Grid(grid.name, new_cells, grid.source)


def find_header_block(grid: Grid, kinds: "KindCache") -> list[int]:
    def header_cells(r):
        return [c for c in grid.row(r) if len(c.text) >= 3 and HEADER_WORD_RE.search(unhyphenate(c.text)) and not c.text[:1].isdigit()]
    cands: list[tuple[int, int, int, list[int]]] = []
    starts: list[int] = []
    r0 = 0
    while r0 < grid.n_rows:
        if len(header_cells(r0)) < 3:
            r0 += 1
            continue
        starts.append(r0)
        last = r0
        covered = {x for c in grid.row(r0) for x in range(c.c, c.c2 + 1)}
        for rr in range(r0 + 1, min(r0 + 16, grid.n_rows)):
            own = grid.row(rr)
            if any("date" in kinds(c) for c in own):
                break
            data_kinds = {"amount", "date", "inn", "account", "bik", "kpp", "composite", "long_text"}
            sub_band = bool(own) and all(c.c >= min(covered) and c.c2 <= max(covered) and not (kinds(c) & data_kinds) for c in own) \
                and any(not c.text[:1].isdigit() for c in own)
            if len(header_cells(rr)) >= 2 or _is_numbering_row(own) or sub_band:
                last = rr
                covered |= {x for c in own for x in range(c.c, c.c2 + 1)}
            elif own:
                break
        last = max([last] + [c.r2 for rr in range(r0, last + 1) for c in grid.row(rr)])
        top = min([r0] + [c.r for rr in range(max(r0 - 12, 0), r0) for c in grid.row(rr) if c.r2 >= r0])
        dated = [rr for rr in range(last + 1, min(last + 45, grid.n_rows)) if any("date" in kinds(c) for c in grid.row(rr))]
        if len(dated) >= 3 and dated[0] - last <= 10:
            cands.append((last, dated[0], r0, list(range(top, last + 1))))
        r0 = last + 1
    for last, first_date, start, rows in cands:
        if not any(last < x < first_date for x in starts):
            return rows
    return []


def find_regions(grid: Grid, _snapped: bool = False) -> list[Region]:
    kinds = KindCache()
    if not _snapped:
        block = find_header_block(grid, kinds)
        if block:
            snapped = snap_grid(grid, block, block[-1] + 1, grid.n_rows - 1)
            if snapped is not None:
                return find_regions(snapped, _snapped=True)
    anchors, date_col = find_anchor_rows(grid, kinds)
    if not anchors:
        return []
    if not _snapped:
        first_cols = sorted({cell.c for cell in grid.row(anchors[0])} | {date_col})
        header = detect_header(grid, kinds, anchors[0], first_cols)
        if header:
            snapped = snap_grid(grid, header, anchors[0], min(anchors[-1] + 3, grid.n_rows - 1))
            if snapped is not None:
                return find_regions(snapped, _snapped=True)
    def _dense_between(a: int, b: int) -> bool:
        rows = range(a + 1, b)
        return (sum(1 for rr in rows if grid.row(rr)) >= 0.6 * len(rows)
                and sum(1 for rr in rows if _row_has(grid, kinds, rr, "amount")) >= 3)

    groups: list[list[int]] = [[anchors[0]]]
    for r in anchors[1:]:
        if r - groups[-1][-1] > 30 and not _dense_between(groups[-1][-1], r):
            groups.append([r])
        else:
            groups[-1].append(r)
    sig_cache: dict[int, set[int]] = {}

    def _block_sig(a: int) -> set[int]:
        if a not in sig_cache:
            cols: set[int] = set()
            for rr in range(a, min(a + 3, grid.n_rows)):
                row = grid.row(rr)
                if not row:
                    break
                cols |= {cell.c for cell in row}
            sig_cache[a] = cols
        return sig_cache[a]

    merged: list[list[int]] = []
    for g in groups:
        if len(g) == 1 and merged:
            prev = merged[-1]
            prev_cols = set().union(*(_block_sig(a) for a in prev))
            if _block_sig(g[0]) and _block_sig(g[0]) <= (prev_cols | {date_col}):
                prev.append(g[0])
                continue
        merged.append(list(g))
    groups = merged
    regions = []
    for gi, group in enumerate(groups):
        next_start = groups[gi + 1][0] if gi + 1 < len(groups) else grid.n_rows
        block_rows = op_block_rows(grid, group, group[-1])
        col_hits = Counter(cell.c for r in block_rows for cell in grid.row(r))
        top = max(col_hits.values(), default=0)
        dense = [c for c, n in col_hits.items() if n >= 0.6 * top]
        lo_d, hi_d = (min(dense), max(dense)) if dense else (0, 10 ** 6)
        def _data_like(col: int) -> bool:
            hits = 0
            for r in block_rows:
                cell = grid.at(r, col)
                if cell is not None and cell.c == col and (kinds(cell) & {"date", "inn"} or ("amount" in kinds(cell) and re.search(r"[.,-]\d{2}\s*(?:₽|р\.?|руб\.?|rub)?$", cell.text, re.I))):
                    hits += 1
            return hits >= 2
        def _noise_col(col: int, n: int) -> bool:
            if lo_d <= col <= hi_d or n > 0.5 * top or _data_like(col):
                return False
            return all(len(cell.text) <= 40 for r in block_rows if (cell := grid.at(r, col)) is not None and cell.c == col)
        cols = sorted({c for c, n in col_hits.items() if not _noise_col(c, n)} | {date_col})
        header_rows = detect_header(grid, kinds, group[0], cols)
        for r in header_rows:
            for cell in grid.row(r):
                if cell.c not in cols and cell.width == 1:
                    cols.append(cell.c)
        cols = sorted(set(cols))
        end = group[-1]
        empty_streak = 0
        for r in range(group[-1] + 1, min(next_start, group[-1] + 16)):
            if not grid.row(r):
                empty_streak += 1
                if empty_streak >= 3:
                    break
                continue
            empty_streak = 0
            end = r
        region = Region(grid=grid, header_rows=header_rows, first_anchor=group[0], last_anchor=group[-1],
                        end_row=end, columns=cols, headers=header_paths(grid, header_rows, cols),
                        anchors=group, kinds=kinds)
        region.profiles = build_profiles(region)
        regions.append(region)
    if len(regions) > 1:
        regions = [r for r in regions if len(r.anchors) >= 2] or regions
    return regions


TRANSPOSED_LABEL_RE = re.compile(
    r"^\s*(?:дата|сумма|назначение\s+платежа|назначение|контрагент|корреспондент|инн|бик|сч[её]т|направление|"
    r"вид\s+операц|тип\s+операц|плательщик|получатель|дебет|кредит|остаток|приход|расход|номер|№)\b", re.I)


def looks_transposed(grid: Grid) -> int | None:
    if grid.n_cols < 5 or grid.n_rows < 5:
        return None
    label_col = None
    for c in range(min(3, grid.n_cols)):
        hits = sum(1 for cell in grid.cells
                   if cell.c == c and cell.r2 == cell.r and TRANSPOSED_LABEL_RE.match(cell.text))
        if hits >= 4:
            label_col = c
            break
    if label_col is None:
        return None
    for r in range(grid.n_rows):
        if sum(1 for c in range(grid.n_cols)
               if (cell := grid.at(r, c)) is not None and "date" in cell_kinds(cell.text, cell.value)) >= 3:
            return label_col
    return None


def transposed_grid(grid: Grid, label_col: int) -> Grid | None:
    first_label = min((cell.r for cell in grid.cells
                       if cell.c == label_col and TRANSPOSED_LABEL_RE.match(cell.text)), default=None)
    if first_label is None or grid.n_rows - first_label < 3:
        return None
    cells = [Cell(x.c, x.r - first_label, x.c2, x.r2 - first_label, x.text, x.value)
             for x in grid.cells if x.r >= first_label]
    if not cells:
        return None
    return Grid(f"{grid.name} (трансп.)", cells, f"transposed:{grid.source}")


def op_block_rows(grid: Grid, anchors: list[int], last_anchor: int) -> list[int]:
    anchor_set = set(anchors)
    rows_set = set(anchor_set)
    for a in anchors:
        for k in (1, 2):
            rr = a + k
            if rr in anchor_set or rr > last_anchor or not grid.row(rr):
                break
            if FOOTER_WORD_RE.search(grid.row_text(rr)):
                break
            rows_set.add(rr)
    return sorted(rows_set)


def build_profiles(region: Region) -> dict[int, ColumnProfile]:
    grid, kinds = region.grid, region.kinds
    rows = op_block_rows(grid, region.anchors, region.last_anchor)
    profiles = {}
    for c in region.columns:
        own = []
        for r in rows:
            cell = grid.at(r, c)
            if cell is not None and cell.c == c and cell.r == r:
                own.append(cell)
        counts: Counter = Counter()
        negatives = decimals = 0
        for cell in own:
            ks = kinds(cell)
            counts.update(ks)
            if "amount" in ks:
                amount = parse_amount(cell.value if isinstance(cell.value, (int, float)) else cell.text)
                if amount is not None and amount < 0:
                    negatives += 1
                if re.search(r"[.,-]\d{2}\b", cell.text) or (isinstance(cell.value, float) and not cell.value.is_integer()):
                    decimals += 1
        n = len(own)
        if counts.get("inn_nozero") and counts.get("inn", 0) >= 0.5 * n:
            counts["inn"] += counts["inn_nozero"]
        if counts.get("zero"):
            counts["amount_compatible_zero"] = counts["zero"]
        texts = [cell.text for cell in own]
        profiles[c] = ColumnProfile(
            col=c, header=region.headers.get(c, ""), n=n,
            fill=n / max(len(rows), 1),
            frac={k: v / n for k, v in counts.items()} if n else {},
            negatives=negatives / max(counts.get("amount", 0), 1),
            decimals=decimals / max(counts.get("amount", 0), 1),
            distinct=len(set(texts)) / n if n else 0.0,
            avg_len=sum(map(len, texts)) / n if n else 0.0,
            examples=list(dict.fromkeys(texts))[:6],
        )
    return profiles


# схема колонок
H_DATE_OP = re.compile(r"операц|провод|исполн|списан|зачисл|date", re.I)
H_DATE_DOC = re.compile(r"документ|валютир|поступ\w*\s+в\s+банк", re.I)
H_AMOUNT = re.compile(r"сумм|дебет|кредит|(?<![а-яё])(?:дт|кт|дб|кр)(?![а-яё])|списан|зачисл|приход|расход|оборот|поступл|amount|debit|credit|withdraw|deposit|руб", re.I)
H_NOT_AMOUNT = re.compile(r"№|номер|код|шифр|\bво\b|вид|бик|кпп|инн|сч[её]т(?!.*сумм)|очередн|n п/п|#", re.I)
H_FX = re.compile(r"в\s+валют|валютн|currency|foreign|в\s+вал\.", re.I)
H_RUB = re.compile(r"руб|rub|rur|₽|национальн|эквивалент", re.I)
H_BALANCE = re.compile(r"остат|сальдо|баланс|balance", re.I)
H_DEBIT = re.compile(r"дебет|(?<![а-яё])(?:дт|дб)(?![а-яё])|списан|расход|debit|withdraw|снят", re.I)
H_CREDIT = re.compile(r"кредит|(?<![а-яё])(?:кт|кр)(?![а-яё])|зачисл|поступ|приход|credit|deposit", re.I)
H_PAYER = re.compile(r"плательщ|отправител|payer|sender", re.I)
H_PAYEE = re.compile(r"получател|бенефициар|payee|recipient|beneficiary", re.I)
H_CP = re.compile(r"контрагент|корреспондент|counterpart|реквизит", re.I)
H_BANK = re.compile(r"банк|bank|бик", re.I)
H_PURPOSE = re.compile(r"назнач|описан|основание(?![а-яё])|содержан|комментар|детал|примечан|purpose|description|details|narrative", re.I)
H_NAME = re.compile(r"наим|назван|ф\.?\s?и\.?\s?о|name|клиент|организац", re.I)
H_INN = re.compile(r"инн|tax|inn", re.I)
H_DIRECTION = re.compile(r"направлен|тип|вид\s+операц|признак|дт/кт|д/к|приход/расход", re.I)
H_DOCNO = re.compile(r"^№|номер|№\s*док|n\s*п/п|шифр|\bво\b|вид\s+оп", re.I)

DIRECTION_OUT = {"д", "дт", "дб", "дебет", "debit", "dr", "d", "-", "списание", "расход", "out", "outgoing", "исходящий", "снятие"}
DIRECTION_IN = {"к", "кт", "кр", "кредит", "credit", "cr", "c", "+", "зачисление", "приход", "in", "incoming", "входящий", "поступление"}


@dataclass
class TableSpec:
    date: int | None = None
    debit: int | None = None
    credit: int | None = None
    signed: int | None = None
    amount: int | None = None
    direction: int | None = None
    direction_out: list[str] = field(default_factory=list)
    direction_in: list[str] = field(default_factory=list)
    payer: dict[str, int | None] = field(default_factory=lambda: {"inn": None, "name": None, "composite": None})
    payee: dict[str, int | None] = field(default_factory=lambda: {"inn": None, "name": None, "composite": None})
    counterparty: dict[str, int | None] = field(default_factory=lambda: {"inn": None, "name": None, "composite": None})
    purpose: list[int] = field(default_factory=list)
    balance: list[int] = field(default_factory=list)
    confidence: float = 1.0
    notes: list[str] = field(default_factory=list)
    source: str = "algorithm"

    def as_dict(self) -> dict[str, Any]:
        return {k: v for k, v in self.__dict__.items()}

    def used_columns(self) -> set[int]:
        cols = {self.date, self.debit, self.credit, self.signed, self.amount, self.direction}
        for side in (self.payer, self.payee, self.counterparty):
            cols |= set(side.values())
        cols |= set(self.purpose)
        return {c for c in cols if c is not None}


def _last(header: str) -> str:
    return header.split("›")[-1].strip()


def _is_bank_col(p: ColumnProfile) -> bool:
    if H_BANK.search(p.header):
        return True
    if p.frac.get("composite", 0.0) >= 0.4 or p.frac.get("inn", 0.0) >= 0.4:
        return False
    values = [e for e in p.examples if re.search(r"[А-Яа-яA-Za-z]{3,}", e)]
    party_header = (H_PAYER.search(_last(p.header)) or H_PAYEE.search(_last(p.header)) or H_CP.search(_last(p.header))) and not H_BANK.search(p.header)
    if party_header:
        return False
    threshold = 0.85 if H_NAME.search(_last(p.header)) and not H_BANK.search(p.header) else 0.5
    return bool(values) and sum(1 for e in values if BANK_WORD_RE.search(e)) / len(values) >= threshold


def _looks_like_purpose(p: ColumnProfile) -> bool:
    return p.frac.get("text", 0.0) >= 0.5 and p.avg_len >= 12 and p.distinct >= 0.3


def _f(p: ColumnProfile, kind: str) -> float:
    if kind == "amount":
        return p.frac.get("amount", 0.0) + p.frac.get("zero", 0.0)
    return p.frac.get(kind, 0.0)


def _is_direction_column(p: ColumnProfile) -> bool:
    values = {e.strip().casefold().rstrip(".") for e in p.examples}
    return bool(values) and len(values) <= 6 and values <= (DIRECTION_OUT | DIRECTION_IN) and p.fill > 0.2


def _complementary(region: Region, a: int, b: int) -> float:
    grid, kinds = region.grid, region.kinds
    one = both = 0
    for r in region.anchors:
        va = parse_amount(grid.text_at(r, a)) if "amount" in kinds(grid.at(r, a)) else None
        vb = parse_amount(grid.text_at(r, b)) if "amount" in kinds(grid.at(r, b)) else None
        fa = va not in (None, Decimal(0))
        fb = vb not in (None, Decimal(0))
        if fa != fb:
            one += 1
        elif fa and fb:
            both += 1
    return one / max(one + both, 1)


def _running_balance(region: Region, bal: int, amount_cols: list[int]) -> bool:
    grid = region.grid
    hits = total = 0
    prev = None
    for r in region.anchors:
        b = parse_amount(grid.text_at(r, bal))
        if b is None:
            continue
        if prev is not None:
            total += 1
            delta = abs(b - prev)
            for c in amount_cols:
                v = parse_amount(grid.text_at(r, c))
                if v is not None and abs(abs(v) - delta) < Decimal("0.01"):
                    hits += 1
                    break
        prev = b
    return total >= 2 and hits / total >= 0.7


def _looks_like_counter(region: Region, c: int) -> bool:
    grid = region.grid
    values = []
    for r in region.anchors:
        cell = grid.at(r, c)
        if cell is None or cell.r != r:
            continue
        t = cell.text.strip()
        if not re.fullmatch(r"\d{1,7}", t):
            return False
        values.append((t, int(t)))
    if len(values) < 2:
        return bool(values) and (values[0][0].startswith("0") or H_DOCNO.search(region.headers.get(c, "") or ""))
    nums = [v for _, v in values]
    if all(t.startswith("0") or v <= 99 for t, v in values) and len(set(nums)) <= max(3, len(nums) // 3):
        return True
    steps = [b - a for a, b in zip(nums, nums[1:])]
    increasing = sum(1 for s in steps if 0 < s <= 5) / len(steps)
    return increasing >= 0.7


def infer_table_spec(region: Region) -> TableSpec:
    spec = TableSpec()
    P = region.profiles
    cols = [c for c in region.columns if P[c].n > 0]

    date_cands = [c for c in cols if _f(P[c], "date") >= 0.6]
    if date_cands:
        def date_score(c):
            h, last = P[c].header, _last(P[c].header)
            op = bool(H_DATE_OP.search(last)) or (last == h and bool(H_DATE_OP.search(h)))
            doc = bool(H_DATE_DOC.search(h)) and not op
            not_date = bool(H_DOCNO.search(last)) and not re.search(r"дат|date", last, re.I)
            return (_f(P[c], "date") * P[c].fill + (0.8 if op else (0.25 if re.search(r"дат|date", h, re.I) else 0))
                    - (0.6 if doc else 0) - (0.8 if not_date else 0))
        spec.date = max(date_cands, key=date_score)
        if len(date_cands) > 1:
            spec.notes.append(f"несколько колонок дат {date_cands}, выбрана C{spec.date}")
    else:
        spec.confidence -= 0.6
        spec.notes.append("колонка даты не найдена")

    def amount_score(c):
        p = P[c]
        s = _f(p, "amount") - _f(p, "date")
        if H_AMOUNT.search(p.header):
            s += 0.5
        if H_NOT_AMOUNT.search(p.header) and not H_AMOUNT.search(p.header):
            s -= 1.0
        if not H_AMOUNT.search(p.header) and _looks_like_counter(region, c):
            s -= 0.8
        s += 0.3 * p.decimals
        return s

    amount_cols = [c for c in cols if c != spec.date and _f(P[c], "amount") >= 0.5 and amount_score(c) >= 0.6
                   and not _is_direction_column(P[c])]
    if any(P[c].decimals >= 0.5 for c in amount_cols):
        dropped = [c for c in amount_cols if P[c].decimals == 0 and not H_AMOUNT.search(P[c].header)]
        if dropped:
            spec.notes.append(f"колонки {dropped} — целые числа без копеек, не суммы")
        amount_cols = [c for c in amount_cols if c not in dropped]
    balance_cols = [c for c in amount_cols if H_BALANCE.search(P[c].header)]
    rest = [c for c in amount_cols if c not in balance_cols]
    fx_cols = [c for c in rest if H_FX.search(P[c].header) and not H_RUB.search(P[c].header)]
    if fx_cols and any(H_RUB.search(P[c].header) for c in rest):
        rest = [c for c in rest if c not in fx_cols]
        spec.notes.append(f"колонки {fx_cols} — суммы в валюте, взяты рубли")
    if len(rest) >= 2:
        for c in list(rest):
            others = [o for o in rest if o != c]
            if _running_balance(region, c, others):
                balance_cols.append(c)
                rest.remove(c)
                spec.notes.append(f"C{c} — текущий остаток (меняется на сумму операции)")
    spec.balance = balance_cols
    debit_h = [c for c in rest if H_DEBIT.search(P[c].header) and not H_CREDIT.search(_last(P[c].header))]
    credit_h = [c for c in rest if H_CREDIT.search(_last(P[c].header)) and c not in debit_h]

    if debit_h and credit_h:
        spec.debit, spec.credit = debit_h[0], credit_h[0]
    elif len(rest) >= 2:
        pairs = [(x, y) for i, x in enumerate(rest) for y in rest[i + 1:]]
        a, b = max(pairs, key=lambda xy: (_complementary(region, *xy), -abs(xy[0] - xy[1])))
        if _complementary(region, a, b) >= 0.8:
            spec.debit, spec.credit = a, b
            spec.confidence -= 0.25
            spec.notes.append(f"дебет/кредит C{a}/C{b} — по соглашению «левая = списание»")
        else:
            spec.amount = rest[0]
            spec.confidence -= 0.3
            spec.notes.append(f"несколько колонок сумм {rest}, взята C{rest[0]}")
    elif len(rest) == 1:
        c = rest[0]
        if P[c].negatives >= 0.05:
            spec.signed = c
        else:
            spec.amount = c
    else:
        d_cols = [c for c in cols if c != spec.date and H_DEBIT.search(_last(P[c].header)) and not H_CREDIT.search(_last(P[c].header))]
        c_cols = [c for c in cols if c != spec.date and H_CREDIT.search(_last(P[c].header))]
        if d_cols or c_cols:
            spec.debit = d_cols[0] if d_cols else None
            spec.credit = c_cols[0] if c_cols else None
            spec.confidence -= 0.4
            spec.notes.append(f"дебет/кредит C{spec.debit}/C{spec.credit} — по шапке; в них не числа, суммы нет, направление по тому, какая колонка заполнена")
        else:
            spec.confidence -= 0.6
            spec.notes.append("колонка суммы не найдена")

    if spec.amount is not None:
        dir_cols = [c for c in cols if c not in (spec.date, spec.amount) and _is_direction_column(P[c])]
        if dir_cols:
            spec.direction = dir_cols[0]
            spec.direction_out = sorted(DIRECTION_OUT)
            spec.direction_in = sorted(DIRECTION_IN)
        else:
            spec.confidence -= 0.35
            spec.notes.append("одна колонка суммы без знака и без направления — направление по сторонам")

    used = {c for c in (spec.date, spec.debit, spec.credit, spec.signed, spec.amount, spec.direction) if c is not None}
    used |= set(balance_cols)

    def side_of(c) -> str | None:
        h = P[c].header
        last = _last(h)
        if H_PAYER.search(last) != H_PAYEE.search(last) and last != h:
            return "payer" if H_PAYER.search(last) else "payee"
        if H_PAYER.search(h) and H_PAYEE.search(h):
            return "counterparty"
        if H_PAYER.search(h):
            return "payer"
        if H_PAYEE.search(h):
            return "payee"
        if H_CP.search(h):
            return "counterparty"
        return None

    inn_like = {c for c in cols if _f(P[c], "inn") >= 0.5 or _f(P[c], "composite") >= 0.4}

    def _next_to_inn(c: int) -> bool:
        i = cols.index(c)
        return any(cols[j] in inn_like for j in (i - 1, i + 1) if 0 <= j < len(cols))

    bank_cols = [c for c in cols if c not in used and _is_bank_col(P[c])
                 and (H_BANK.search(P[c].header) or not _next_to_inn(c))]
    if bank_cols:
        spec.notes.append(f"колонки банка контрагента {bank_cols} исключены из сторон")
    used |= set(bank_cols)
    inn_cols = [c for c in cols if c not in used and (_f(P[c], "inn") >= 0.5 or (H_INN.search(P[c].header) and _f(P[c], "inn") >= 0.3))]
    comp_cols = [c for c in cols if c not in used and c not in inn_cols and _f(P[c], "composite") >= 0.4]
    name_cols = [c for c in cols if c not in used and c not in inn_cols and c not in comp_cols
                 and not H_BANK.search(_last(P[c].header))
                 and not H_PURPOSE.search(P[c].header)
                 and (_f(P[c], "name") >= 0.4 or ((H_NAME.search(P[c].header) or H_CP.search(_last(P[c].header)) or H_PAYER.search(_last(P[c].header))
                                                or H_PAYEE.search(_last(P[c].header))) and _f(P[c], "text") >= 0.5))]
    inn_cols = [c for c in inn_cols if not H_BANK.search(_last(P[c].header))]

    groups: list[dict[str, Any]] = []
    party_rank = {c: i for i, c in enumerate(sorted(set(inn_cols) | set(comp_cols) | set(name_cols)))}
    for kind, clist in (("inn", inn_cols), ("composite", comp_cols), ("name", name_cols)):
        for c in clist:
            side = side_of(c)
            target = None
            for g in groups:
                if kind in g["cols"]:
                    continue
                near = min(abs(c - x) for x in g["cols"].values()) <= 3 or \
                    min(abs(party_rank[c] - party_rank[x]) for x in g["cols"].values()) <= 1
                if side is not None and g["side"] == side:
                    target = g
                    break
                if near and (side is None or g["side"] is None):
                    others_same_side = any(o is not g and o["side"] == side and side is not None for o in groups)
                    if not others_same_side:
                        target = g
                        if g["side"] is None:
                            g["side"] = side
                        break
            if target is None:
                target = {"side": side, "cols": {}}
                groups.append(target)
            target["cols"][kind] = c

    groups = [g for g in groups if g["cols"]]
    sides_known = {g["side"] for g in groups if g["side"]}
    if len(groups) == 1:
        g = groups[0]
        side = g["side"] if g["side"] in ("payer", "payee") and (spec.debit is None and spec.credit is None and spec.signed is None and spec.direction is None) else "counterparty"
        if g["side"] in ("payer", "payee"):
            spec.notes.append(f"одна сторона с шапкой «{g['side']}» считается контрагентом (вторая — владелец)")
        getattr(spec, side).update(g["cols"])
    elif len(groups) >= 2:
        unknown = [g for g in groups if g["side"] is None]
        if "payer" not in sides_known and unknown:
            unknown.pop(0)["side"] = "payer"
            spec.confidence -= 0.15
            spec.notes.append("сторона «плательщик» назначена по порядку колонок")
        if "payee" not in sides_known and unknown:
            unknown.pop(0)["side"] = "payee"
            spec.confidence -= 0.15
            spec.notes.append("сторона «получатель» назначена по порядку колонок")
        for g in groups:
            if g["side"] in ("payer", "payee", "counterparty"):
                target = getattr(spec, g["side"])
                for k, v in g["cols"].items():
                    if target.get(k) is None:
                        target[k] = v
    else:
        spec.notes.append("колонок контрагента нет (например, карточная выписка)")

    used |= spec.used_columns()

    header_hits = [c for c in cols if c not in used and H_PURPOSE.search(P[c].header)]
    rejected = [c for c in header_hits if not _looks_like_purpose(P[c])]
    if rejected:
        spec.notes.append(f"колонки {rejected} названы как назначение, но содержат коды/короткие значения — отклонены")
    purpose = sorted((c for c in header_hits if _looks_like_purpose(P[c])),
                     key=lambda c: (not re.search(r"назнач", P[c].header, re.I), -P[c].avg_len))
    if not purpose:
        texty = [c for c in cols if c not in used and _f(P[c], "text") >= 0.5 and not H_BANK.search(P[c].header)
                 and not H_DOCNO.search(P[c].header)]
        if texty:
            best = max(texty, key=lambda c: P[c].avg_len * P[c].fill)
            if P[best].avg_len >= 12:
                purpose = [best]
                spec.notes.append(f"назначение — самая «длинная» текстовая колонка C{best}")
                spec.confidence -= 0.05
    spec.purpose = purpose[:1]
    spec.confidence = round(max(spec.confidence, 0.0), 2)
    return spec


def check_spec(spec: TableSpec, region: Region) -> list[str]:
    P = region.profiles
    errors = []

    def exists(c, what):
        if c is not None and c not in P:
            errors.append(f"{what}: колонки C{c} нет в таблице")
            return False
        return c is not None

    if not exists(spec.date, "date"):
        errors.append("не указана колонка даты")
    elif _f(P[spec.date], "date") < 0.5:
        errors.append(f"date=C{spec.date}: даты только в {_f(P[spec.date], 'date'):.0%} значений")
    amount_like = [(spec.debit, "debit"), (spec.credit, "credit"), (spec.signed, "signed"), (spec.amount, "amount")]
    if all(c is None for c, _ in amount_like):
        errors.append("не указана ни одна колонка суммы")
    for c, name in amount_like:
        if exists(c, name) and _f(P[c], "amount") < 0.5:
            errors.append(f"{name}=C{c}: суммы только в {_f(P[c], 'amount'):.0%} значений")
    two_sided = any(v is not None for v in spec.payer.values()) and any(v is not None for v in spec.payee.values())
    if spec.amount is not None and spec.direction is None and not two_sided:
        errors.append("amount без знака требует колонку direction, либо обе стороны payer и payee")
    if spec.signed is not None and spec.signed in P and P[spec.signed].n >= 5 and P[spec.signed].negatives == 0 and not two_sided:
        errors.append(f"signed=C{spec.signed}: нет ни одного отрицательного значения — это не сумма со знаком")
    for side in ("payer", "payee", "counterparty"):
        for k, c in getattr(spec, side).items():
            if not exists(c, f"{side}.{k}"):
                continue
            if k == "inn" and _f(P[c], "inn") < 0.3 and _f(P[c], "composite") < 0.3:
                errors.append(f"{side}.inn=C{c}: верный ИНН только в {_f(P[c], 'inn'):.0%} значений")
            if k == "name" and _f(P[c], "text") < 0.5:
                errors.append(f"{side}.name=C{c}: текста только в {_f(P[c], 'text'):.0%} значений")
    single = [c for c in (spec.date, spec.debit, spec.credit, spec.signed, spec.amount) if c is not None]
    if len(single) != len(set(single)):
        errors.append("одна колонка назначена сразу нескольким полям даты/суммы")
    roles: dict[int, list[str]] = {}
    for label, c in ([("date", spec.date), ("debit", spec.debit), ("credit", spec.credit), ("signed", spec.signed),
                      ("amount", spec.amount), ("direction", spec.direction)]
                     + [(f"{side}.{k}", c) for side in ("payer", "payee", "counterparty")
                        for k, c in getattr(spec, side).items()]
                     + [("purpose", c) for c in spec.purpose]):
        if c is not None:
            roles.setdefault(c, []).append(label)
    for c, labels in sorted(roles.items()):
        if len(labels) > 1:
            errors.append(f"колонка C{c} назначена сразу нескольким ролям {labels} — у каждой роли своя колонка")
    return errors


# классификация строк
OPERATION, CONTINUATION = "OPERATION", "CONTINUATION"
HEADER_REPEAT, COLUMN_NUMBERS, DATE_GROUP = "HEADER_REPEAT", "COLUMN_NUMBERS", "DATE_GROUP"
TOTAL, SUBTOTAL, BALANCE_IN, BALANCE_OUT = "TOTAL", "SUBTOTAL", "BALANCE_IN", "BALANCE_OUT"
SERVICE, NOTE, EMPTY, UNCERTAIN = "SERVICE", "NOTE", "EMPTY", "UNCERTAIN"
ROW_LABELS = [OPERATION, CONTINUATION, HEADER_REPEAT, COLUMN_NUMBERS, DATE_GROUP, TOTAL, SUBTOTAL,
              BALANCE_IN, BALANCE_OUT, SERVICE, NOTE, EMPTY, UNCERTAIN]

TOTAL_RE = re.compile(r"итог|всего|оборот|total|turnover", re.I)
SUBTOTAL_RE = re.compile(r"за\s+день|за\s+дату|по\s+страниц|по\s+листу|перенос|за\s+\d{2}\.\d{2}", re.I)
BAL_IN_RE = re.compile(r"(входящ|начальн|на\s+начало|opening)\w*.{0,20}(остат|сальдо|balance)|(остат|сальдо)\w*\s+(входящ|на\s+начало|начальн)", re.I)
BAL_OUT_RE = re.compile(r"(исходящ|конечн|на\s+конец|closing)\w*.{0,20}(остат|сальдо|balance)|(остат|сальдо)\w*\s+(исходящ|на\s+конец|конечн)", re.I)
SERVICE_RE = re.compile(
    r"страниц|^стр\.?\s*\d|лист\s+\d|продолжение|подпись|м\.\s*п\.|главный\s+бухгалтер|исполнитель|"
    r"сформирован|дата\s+печати|отпечатан|выписка\s+(?:по|из)|page\s+\d",
    re.I,
)


@dataclass
class RowVerdict:
    r: int
    label: str
    confidence: float
    reason: str
    decided_by: str = "rule"
    inherited_date: bool = False


@dataclass
class OpRecord:
    rows: list[int]
    date: date | None
    date_inherited: bool
    values: dict[str, str]
    amounts: dict[str, Decimal | None]
    direction_raw: str | None
    provenance: dict[str, tuple[int, int]]
    grid_name: str


def _row_amounts(region: Region, spec: TableSpec, r: int) -> dict[str, Decimal | None]:
    grid, kinds = region.grid, region.kinds
    out = {}
    for key in ("debit", "credit", "signed", "amount"):
        c = getattr(spec, key)
        if c is None:
            continue
        cell = grid.at(r, c)
        if cell is None or cell.r != r or "amount" not in kinds(cell):
            out[key] = None
            continue
        out[key] = parse_amount(cell.value if isinstance(cell.value, (int, float)) else cell.text)
    return out


def _nonzero(amounts: dict[str, Decimal | None]) -> bool:
    return any(v not in (None, Decimal(0)) for v in amounts.values())


def _header_similarity(region: Region, r: int) -> float:
    header_tokens = set()
    for hr in region.header_rows:
        header_tokens |= set(normalize_tokens(region.grid.row_text(hr)))
    if not header_tokens:
        return 0.0
    tokens = set(normalize_tokens(region.grid.row_text(r)))
    if not tokens:
        return 0.0
    return len(tokens & header_tokens) / len(tokens)


def _header_cells_match(region: Region, r: int) -> int:
    grid, n = region.grid, 0
    for cell in grid.row(r):
        toks = set(normalize_tokens(cell.text))
        if not toks:
            continue
        for hr in region.header_rows:
            hc = grid.at(hr, cell.c)
            if hc is not None and len(toks & set(normalize_tokens(hc.text))) / len(toks) >= 0.6:
                n += 1
                break
    return n


def normalize_tokens(text: str) -> list[str]:
    return [t for t in re.split(r"[^0-9a-zа-яё]+", text.casefold().replace("ё", "е")) if len(t) >= 3]


def _text_cols(spec: TableSpec) -> set[int]:
    cols = set(spec.purpose)
    for side in (spec.payer, spec.payee, spec.counterparty):
        cols |= {c for c in side.values() if c is not None}
    return cols


def _text_only_row(region: Region, spec: TableSpec, r: int, text_cols: set[int]) -> bool:
    grid = region.grid
    if r > region.end_row or not grid.row(r):
        return False
    if _nonzero(_row_amounts(region, spec, r)):
        return False
    date_cell = grid.at(r, spec.date) if spec.date is not None else None
    if date_cell is not None and date_cell.r == r and parse_date(date_cell.value if date_cell.value is not None else date_cell.text):
        return False
    if FOOTER_WORD_RE.search(grid.row_text(r)):
        return False
    return any((cell := grid.at(r, c)) is not None and cell.r == r for c in text_cols)


def classify_rows(region: Region, spec: TableSpec) -> list[RowVerdict]:
    grid = region.grid
    verdicts: list[RowVerdict] = []
    text_cols = _text_cols(spec)
    last_op_amounts: dict[str, Decimal | None] | None = None
    group_sum: dict[str, Decimal] = defaultdict(Decimal)
    current_date: date | None = None

    first_row = max(region.first_anchor - 6, region.header_rows[-1] + 1) if region.header_rows else region.first_anchor
    first_row = min(first_row, region.first_anchor)
    for r in range(first_row, region.end_row + 1):
        own = grid.row(r)
        row_text = grid.row_text(r)
        amounts = _row_amounts(region, spec, r)
        date_cell = grid.at(r, spec.date) if spec.date is not None else None
        row_date = parse_date(date_cell.value if date_cell is not None and date_cell.value is not None else (date_cell.text if date_cell else ""))
        date_is_own = date_cell is not None and date_cell.r == r
        text_present = any(grid.at(r, c) is not None and grid.at(r, c).r == r for c in text_cols)

        def add(label, conf, reason, **kw):
            verdicts.append(RowVerdict(r, label, conf, reason, **kw))

        if not own:
            add(EMPTY, 1.0, "пустая строка")
            continue
        if _is_numbering_row(own):
            add(COLUMN_NUMBERS, 1.0, "строка нумерации колонок")
            continue
        if region.header_rows and row_date is None and len(own) >= 3 and _header_similarity(region, r) >= 0.6 \
                and _header_cells_match(region, r) >= 3:
            add(HEADER_REPEAT, 0.95, "повтор шапки (перенос страницы)")
            continue
        first_texts = " ".join(c.text for c in own[:3])
        if row_date is None or not date_is_own:
            if BAL_IN_RE.search(row_text):
                add(BALANCE_IN, 0.95, "входящий остаток")
                continue
            if BAL_OUT_RE.search(row_text):
                add(BALANCE_OUT, 0.95, "исходящий остаток")
                continue
            if TOTAL_RE.search(first_texts) or (TOTAL_RE.search(row_text) and not text_present):
                label = SUBTOTAL if SUBTOTAL_RE.search(row_text) else TOTAL
                add(label, 0.95, "ключевое слово итога")
                group_sum = defaultdict(Decimal)
                continue
            long_purpose = any((pc := grid.at(r, c)) is not None and pc.r == r and len(pc.text) >= 15
                               and not SERVICE_RE.search(pc.text) for c in spec.purpose)
            if SERVICE_RE.search(row_text) and not _nonzero(amounts) and not long_purpose:
                add(SERVICE, 0.9, "служебная строка")
                continue

        party_cols = {c for side in (spec.payer, spec.payee, spec.counterparty) for c in side.values() if c is not None}
        party_present = any(grid.at(r, c) is not None and grid.at(r, c).r == r for c in party_cols)
        if row_date is not None and not party_present and TOTAL_RE.search(row_text):
            label = SUBTOTAL if (SUBTOTAL_RE.search(row_text) or "день" in row_text.casefold()) else TOTAL
            add(label, 0.9, "ключевое слово итога (строка с датой, без контрагента)")
            group_sum = defaultdict(Decimal)
            continue

        if row_date is not None and date_is_own:
            other = [c for c in own if c.c != spec.date]
            if not other:
                add(DATE_GROUP, 0.9, "строка-заголовок дня (только дата)")
                current_date = row_date
                continue
        if row_date is not None:
            current_date = row_date

        if row_date is not None and (date_is_own or _nonzero(amounts)):
            if _nonzero(amounts) or not date_is_own:
                add(OPERATION, 0.95, "дата и сумма")
            else:
                add(OPERATION, 0.75, "дата без суммы — сумма ожидается в строке-продолжении")
            last_op_amounts = amounts
            for k, v in amounts.items():
                if v is not None:
                    group_sum[k] += v
            continue

        if _nonzero(amounts):
            prev = verdicts[-1] if verdicts else None
            prev_is_op_without_amount = (prev is not None and prev.label in (OPERATION, CONTINUATION)
                                         and last_op_amounts is not None and not _nonzero(last_op_amounts))
            if prev_is_op_without_amount:
                add(CONTINUATION, 0.85, "сумма операции во второй строке")
                last_op_amounts = amounts
                for k, v in amounts.items():
                    if v is not None:
                        group_sum[k] += v
                continue
            matches_group = all(
                v is None or v == 0 or abs(v - group_sum.get(k, Decimal(0))) < Decimal("0.01")
                for k, v in amounts.items()
            ) and any(group_sum.values())
            if matches_group and not text_present:
                add(SUBTOTAL, 0.9, "сумма равна сумме операций группы")
                group_sum = defaultdict(Decimal)
                continue
            dir_cell = grid.at(r, spec.direction) if spec.direction is not None else None
            has_direction = dir_cell is not None and dir_cell.r == r and bool(dir_cell.text.strip())
            next_text_only = _text_only_row(region, spec, r + 1, text_cols)
            if (text_present or has_direction or next_text_only) and current_date is not None:
                why = ("сумма и контрагент/назначение без даты" if text_present else
                       "сумма с направлением без даты" if has_direction else
                       "сумма без даты, текст операции в следующей строке")
                add(OPERATION, 0.7, why + " — дата унаследована", inherited_date=True)
                last_op_amounts = amounts
                for k, v in amounts.items():
                    if v is not None:
                        group_sum[k] += v
                continue
            add(UNCERTAIN, 0.3, "сумма без даты, ключевых слов и текста — не сходится с группой")
            continue

        tail_continuation = (r == region.last_anchor + 1 and bool(verdicts) and verdicts[-1].label == OPERATION
                             and any(v.label == CONTINUATION for v in verdicts)
                             and any((pc := grid.at(r, c)) is not None and pc.r == r and pc.text.strip() for c in spec.purpose))
        if r > region.last_anchor and not tail_continuation and re.match(r"\s*(?:клиент|владелец|в\s+отношении|наименование\s+клиента|идентификационн|инн\b|организация\s*[:\-])", own[0].text, re.I):
            add(NOTE, 0.8, "реквизиты клиента под таблицей — не продолжение операции")
            continue
        prev_ops = [v for v in verdicts if v.label not in (EMPTY,)]
        if prev_ops and prev_ops[-1].label in (OPERATION, CONTINUATION) and (text_present or len(own) <= 3):
            add(CONTINUATION, 0.8, "текст без даты и суммы после операции")
            continue
        add(NOTE, 0.6, "текстовая строка вне операции")
    return verdicts


def build_records(region: Region, spec: TableSpec, verdicts: list[RowVerdict]) -> list[OpRecord]:
    grid = region.grid
    field_cols: dict[str, int] = {}
    for side in ("payer", "payee", "counterparty"):
        for k, c in getattr(spec, side).items():
            if c is not None:
                field_cols[f"{side}.{k}"] = c
    if spec.purpose:
        field_cols["purpose"] = spec.purpose[0]
    if spec.direction is not None:
        field_cols["direction"] = spec.direction

    records: list[OpRecord] = []
    current: OpRecord | None = None
    last_date: date | None = None
    for v in verdicts:
        if v.label == DATE_GROUP:
            cell = grid.at(v.r, spec.date)
            last_date = parse_date(cell.value if cell.value is not None else cell.text)
            continue
        if v.label == OPERATION:
            cell = grid.at(v.r, spec.date) if spec.date is not None else None
            raw_date = cell.value if cell is not None and cell.value is not None else (cell.text if cell else "")
            d = parse_date(raw_date)
            date_invalid = d is None and bool(re.search(r"\d", str(raw_date or "")))
            inherited = d is None and not date_invalid
            if d is None and not date_invalid:
                d = last_date
            last_date = d or last_date
            current = OpRecord(rows=[v.r], date=d, date_inherited=inherited, values={}, amounts={},
                               direction_raw=None, provenance={}, grid_name=grid.name)
            if spec.date is not None and cell is not None:
                current.provenance["date"] = (cell.r, cell.c)
            records.append(current)
        elif v.label == CONTINUATION and current is not None:
            current.rows.append(v.r)
        else:
            if v.label != EMPTY:
                current = None
            continue
        r = v.r
        if (spec.debit is not None or spec.credit is not None) and "_dir_hint" not in current.values:
            dc = grid.at(r, spec.debit) if spec.debit is not None else None
            cc = grid.at(r, spec.credit) if spec.credit is not None else None
            d_txt = dc.text.strip() if dc is not None and dc.r == r else ""
            c_txt = cc.text.strip() if cc is not None and cc.r == r else ""
            if bool(d_txt) != bool(c_txt) and not any(v is not None for v in _row_amounts(region, spec, r).values()):
                current.values["_dir_hint"] = "out" if d_txt else "in"
        for key, amount in _row_amounts(region, spec, r).items():
            if amount is not None and current.amounts.get(key) is None:
                current.amounts[key] = amount
                current.provenance[key] = (r, getattr(spec, key))
        for name, c in field_cols.items():
            cell = grid.at(r, c)
            if cell is None or cell.r != r or cell.c != c:
                continue
            text = cell.text.replace("\n", " ").strip()
            if not text:
                continue
            if name.endswith((".name", ".composite")) and looks_like_purpose(text):
                name = "purpose"
            if name in current.values:
                if text not in current.values[name]:
                    current.values[name] += " " + text
            else:
                current.values[name] = text
                current.provenance[name] = (r, c)
    for rec in records:
        rec.direction_raw = rec.values.pop("direction", None)
    return records


# вызов GigaChat
LLM_CALLS: list[dict[str, Any]] = []


class LLMUnavailable(Exception):
    """Модель недоступна для этого файла: бюджет времени вышел или было слишком много сбоев подряд."""


def guarded_call(fn, label: str, phase: str = "algo"):
    if CLOCK.unavailable:
        raise LLMUnavailable(CLOCK.unavailable)
    left = CLOCK.left(phase)
    if left <= 1:
        raise LLMUnavailable(f"бюджет времени фазы «{phase}» исчерпан")
    timeout = min(LLM_CALL_TIMEOUT, left)
    box: dict[str, Any] = {}

    def work() -> None:
        try:
            box["value"] = fn()
        except BaseException as exc:
            box["error"] = exc

    CLOCK.calls += 1
    started = time.time()
    log("gigachat", f"→ {label}: запрос отправлен (ждём до {timeout:.0f} с, вызов №{CLOCK.calls})")
    thread = threading.Thread(target=work, daemon=True)
    thread.start()
    thread.join(timeout)
    took = time.time() - started
    if thread.is_alive():
        error: BaseException = TimeoutError(f"нет ответа за {timeout:.0f} с")
    else:
        error = box.get("error")
    if error is not None:
        CLOCK.fails += 1
        log("gigachat", f"✗ {label}: {type(error).__name__}: {error} (через {took:.1f} с; сбоев подряд {CLOCK.fails}/{LLM_MAX_FAILS})")
        if CLOCK.fails >= LLM_MAX_FAILS:
            CLOCK.unavailable = f"{CLOCK.fails} сбоя подряд, последний: {type(error).__name__}"
            log("gigachat", f"модель считается недоступной для этого файла ({CLOCK.unavailable}) — дальше только алгоритм")
        raise error
    CLOCK.fails = 0
    log("gigachat", f"← {label}: ответ за {took:.1f} с")
    return box["value"]


def _llm_invoke(llm: Any, messages: list[tuple[str, str]], label: str = "запрос", phase: str = "algo") -> str:
    return guarded_call(lambda: _llm_invoke_raw(llm, messages), label, phase)


def _llm_invoke_raw(llm: Any, messages: list[tuple[str, str]]) -> str:
    if SystemMessage is not None and HumanMessage is not None and hasattr(llm, "invoke") and not getattr(llm, "plain_text", False):
        from langchain_core.messages import AIMessage

        lc = []
        for role, text in messages:
            lc.append({"system": SystemMessage, "user": HumanMessage, "assistant": AIMessage}[role](content=text))
        response = llm.invoke(lc)
    elif hasattr(llm, "invoke"):
        response = llm.invoke(messages)
    else:
        response = llm(messages)
    return getattr(response, "content", response) if not isinstance(response, str) else response


class CachedLLM:

    def __init__(self, client: Any, path: str | Path, model: str = ""):
        self.client, self.path, self.model = client, Path(path), model or getattr(client, "model", "") or ""
        self.plain_text = getattr(client, "plain_text", False)
        self.hits = self.misses = 0
        try:
            self.store: dict[str, str] = json.loads(self.path.read_text(encoding="utf-8")) if self.path.exists() else {}
        except (OSError, ValueError):
            self.store = {}

    @staticmethod
    def _text(messages) -> list[list[str]]:
        out = []
        for m in messages:
            if isinstance(m, (tuple, list)):
                out.append([str(m[0]), str(m[1])])
            elif isinstance(m, dict):
                out.append([str(m.get("role", "")), str(m.get("content", ""))])
            else:
                out.append([str(getattr(m, "type", type(m).__name__)), str(getattr(m, "content", m))])
        return out

    def invoke(self, messages):
        import hashlib
        key = hashlib.sha256(json.dumps([self.model, self._text(messages)], ensure_ascii=False).encode()).hexdigest()
        if key in self.store:
            self.hits += 1
            return type("CachedResponse", (), {"content": self.store[key]})()
        response = self.client.invoke(messages)
        content = getattr(response, "content", response)
        if isinstance(content, str):
            self.store[key] = content
            self.misses += 1
            try:
                self.path.write_text(json.dumps(self.store, ensure_ascii=False), encoding="utf-8")
            except OSError:
                pass
        return response


if LLM_CLIENT is not None and CACHE_PATH and not isinstance(LLM_CLIENT, CachedLLM):
    LLM_CLIENT = CachedLLM(LLM_CLIENT, CACHE_PATH, GIGACHAT_MODEL)


def extract_json(raw: str) -> tuple[Any, str | None]:
    text = (raw or "").strip()
    fence = re.search(r"```(?:json)?\s*(.*?)```", text, re.S)
    if fence:
        text = fence.group(1).strip()
    start = min([i for i in (text.find("{"), text.find("[")) if i >= 0], default=-1)
    if start < 0:
        return None, "в ответе нет JSON"
    opener = text[start]
    closer = "}" if opener == "{" else "]"
    end = text.rfind(closer)
    if end <= start:
        return None, "JSON не закрыт"
    chunk = text[start:end + 1]
    try:
        return json.loads(chunk), None
    except json.JSONDecodeError as exc:
        fixed = re.sub(r",\s*([}\]])", r"\1", chunk)
        fixed = re.sub(r"//[^\n]*", "", fixed)
        try:
            return json.loads(fixed), None
        except json.JSONDecodeError:
            return None, f"JSON не разбирается: {exc.msg} (позиция {exc.pos})"


def call_llm_json(task: str, system: str, user: str, validate, llm: Any | None = None,
                  retries: int = 1, phase: str = "algo") -> Any | None:
    llm = llm if llm is not None else LLM_CLIENT
    if llm is None:
        return None
    messages = [("system", system), ("user", user)]
    entry: dict[str, Any] = {"task": task, "attempts": [], "accepted": False,
                             "prompt_chars": len(system) + len(user)}
    LLM_CALLS.append(entry)
    for attempt in range(retries + 1):
        started = time.time()
        try:
            raw = _llm_invoke(llm, messages, label=f"{task}, попытка {attempt + 1}/{retries + 1}, {sum(len(t) for _, t in messages)} симв.", phase=phase)
        except Exception as exc:
            entry["attempts"].append({"error": f"{type(exc).__name__}: {exc}"})
            return None
        data, error = extract_json(raw)
        if error is None:
            data, error = validate(data)
        entry["attempts"].append({"raw": raw[:3000], "error": error, "seconds": round(time.time() - started, 2)})
        if error is None:
            entry["accepted"] = True
            entry["result"] = data
            log("gigachat", f"{task}: ответ принят проверкой (попытка {attempt + 1})")
            return data
        log("gigachat", f"{task}: ответ отклонён проверкой — {error}")
        messages += [("assistant", raw), ("user",
                      f"Ответ не прошёл проверку: {error}\n"
                      f"Исправь и верни ТОЛЬКО JSON в том же формате, без пояснений.")]
    return None


OWNER_SYSTEM = """Ты определяешь владельца банковской выписки — клиента, по счёту которого она сформирована.
Тебе даны ячейки шапки выписки с координатами (R — строка, C — колонка) и кандидаты, которые уже нашёл алгоритм.

ПРАВИЛА
1. Владелец — клиент банка (организация, ИП или физлицо), а НЕ банк, НЕ филиал банка, НЕ БИК/корр.счёт и НЕ контрагент из операций.
2. Выбирай только из списков кандидатов по их ID. Сам ИНН не пиши.
3. Если наименования владельца нет среди кандидатов, но оно есть в ячейке — верни координату этой ячейки в name_cell (например "R4C2").
4. Сигналы: метка рядом («Владелец счёта», «Клиент», «Организация», «Плательщик» в шапке), близость ИНН и наименования, частота в операциях.
5. Если надёжно определить нельзя — верни null в соответствующем поле.

Ответ — ТОЛЬКО JSON без пояснений:
{"inn_id": "I1" | null, "name_id": "N2" | null, "name_cell": "R4C2" | null, "confidence": 0.0-1.0, "reason": "до 15 слов"}"""

TABLE_SYSTEM = """Ты сопоставляешь колонки таблицы операций банковской выписки с полями.
Для каждой колонки дана карточка: путь шапки, доли типов значений (date, amount, inn, name, composite, text…) и примеры.

ПОЛЯ (укажи номер колонки или null)
- date — дата проведения операции (не дата документа, если есть обе).
- debit — сумма СПИСАНИЯ со счёта владельца (дебет, расход, списано).
- credit — сумма ЗАЧИСЛЕНИЯ на счёт владельца (кредит, приход, поступило).
- signed — одна колонка суммы со знаком (минус = списание). Используй вместо debit/credit.
- amount + direction — одна сумма без знака и отдельная колонка направления; тогда перечисли значения
  направления для списания в direction_out и для зачисления в direction_in (как они написаны в колонке).
- payer / payee — плательщик и получатель, если в таблице есть обе стороны; counterparty — если одна сторона-контрагент.
  В каждой стороне: inn (колонка ИНН), name (колонка наименования), composite (одна ячейка с ИНН+названием+счётом).
- purpose — назначение платежа / описание операции.
- balance — колонки остатка (их НЕ использовать как сумму).

ПРАВИЛА
1. Номера колонок бери только из карточек (C0, C1… → пиши число 0, 1…).
2. Не путай: номер документа, код вида операции, БИК, КПП, счёт и остаток — это НЕ сумма и НЕ ИНН.
3. Банк контрагента — не контрагент.
4. Тип значений важнее названия колонки: если в «ИНН» лежат 20-значные счета — это не ИНН.

Ответ — ТОЛЬКО JSON без пояснений, пример:
{"date": 0, "debit": 5, "credit": 6, "signed": null, "amount": null, "direction": null,
 "direction_out": [], "direction_in": [],
 "payer": {"inn": null, "name": null, "composite": null},
 "payee": {"inn": null, "name": null, "composite": null},
 "counterparty": {"inn": 3, "name": 4, "composite": null},
 "purpose": 7, "balance": [], "confidence": 0.9}"""

ROWS_SYSTEM = """Ты размечаешь строки таблицы банковской выписки, которые алгоритм не смог классифицировать.
Для каждой строки дан её текст по колонкам и соседние строки для контекста.

МЕТКИ
- OPERATION — отдельная операция (движение денег).
- CONTINUATION — продолжение предыдущей операции (перенос текста или суммы на следующую строку).
- SUBTOTAL — промежуточный итог (за день, по странице, перенос).
- TOTAL — итог/обороты за период.
- BALANCE_IN / BALANCE_OUT — входящий / исходящий остаток.
- HEADER_REPEAT — повтор шапки; SERVICE — служебная строка (страница, подписи); NOTE — прочий текст.

Размечай ТОЛЬКО строки из списка «РАЗМЕТИТЬ». Ответ — ТОЛЬКО JSON без пояснений:
{"R57": "SUBTOTAL", "R58": "OPERATION"}"""

REPAIR_SYSTEM = """Ты исправляешь разметку строк банковской выписки. Сверка с итогами документа не сошлась.
Дано: расхождение, текущая разметка строк-кандидатов и их суммы.
Предложи, какие строки надо переразметить, чтобы сумма операций совпала с итогом документа.
Меняй только строки из списка кандидатов. Допустимые метки: OPERATION, SUBTOTAL, TOTAL, NOTE, BALANCE_IN, BALANCE_OUT.

Ответ — ТОЛЬКО JSON без пояснений: {"R57": "OPERATION", "R60": "SUBTOTAL"}"""

TEXT_SYSTEM = """Ты извлекаешь операции из фрагмента банковской выписки, который не является таблицей.
Строки фрагмента пронумерованы (L12: …).

Для каждой операции верни:
- lines — номера строк, где она описана;
- date — дату РОВНО как в тексте;
- amount — сумму РОВНО как в тексте;
- direction — "out" (списание со счёта владельца) или "in" (зачисление);
- counterparty_name, counterparty_inn — контрагент РОВНО как в тексте (или null);
- purpose — назначение РОВНО как в тексте (или null).

ПРАВИЛА
1. Копируй значения из текста дословно. Ничего не придумывай и не вычисляй.
2. Итоги, остатки, обороты — не операции.
3. Если в фрагменте нет операций — верни [].

Ответ — ТОЛЬКО JSON-массив без пояснений:
[{"lines": [12, 13], "date": "01.02.2024", "amount": "1 500,00", "direction": "out",
  "counterparty_name": "ООО «Ромашка»", "counterparty_inn": "7701234567", "purpose": "Оплата по счёту 5"}]"""


# владелец выписки
OWNER_LBL_RE = re.compile(
    r"владел|клиент|организац|наименование(?!\s+банк)|плательщик|получатель|предприят|абонент|заемщик|"
    r"account\s+holder|customer|\bфио\b|лицевой\s+сч|в\s+отношении",
    re.I,
)
OTHER_LBL_RE = re.compile(r"^(?:инн|кпп|огрн|бик|сч[её]т|р/с|адрес|телефон|валюта|период|дата|№|номер|банк)\b", re.I)
CAPTION_RE = re.compile(
    r"^\(?(?:сокращенное|полное|наименование|ф\.?\s?и\.?\s?о|владелец|депонент|бенефициар|должность|подпись|"
    r"дата|номер|вид\s+сч|физического\s+лица|индивидуального|цифровой\s+код|сведения)", re.I)


def is_caption(text: str) -> bool:
    t = text.strip()
    if not t:
        return True
    if CAPTION_RE.match(t) and t[0].islower():
        return True
    return bool(re.search(r"наименование\s*/\s*ф\.?\s?и\.?\s?о|\(указывается|представляющего\s+информацию|"
                          r"именительном\s+падеже", t, re.I))


NONE_LIKE = {"", "none", "null", "nan", "n/a", "na", "-", "—", "–", "нет", "неизвестно", "не определён", "не определен", "unknown", "owner", "владелец"}
TAX_BODY_RE = re.compile(r"налогов|\bфнс\b|\bифнс\b|инспекц|межрайон|код\s+(?:налогового\s+)?орган|\bкод\s+ифнс", re.I)


def owner_problems(owner: dict[str, Any]) -> list[str]:
    problems: list[str] = []
    name = (owner.get("name") or "").strip()
    inn = (owner.get("inn") or "").strip()
    if name:
        if name.casefold() in NONE_LIKE:
            problems.append(f"наименование владельца {name!r} — пустое значение, а не имя")
        elif not re.search(r"[А-Яа-яA-Za-z]{2,}", name):
            problems.append(f"наименование владельца {name!r} не содержит букв")
        elif TAX_BODY_RE.search(name):
            problems.append(f"наименование владельца {name!r} — налоговый орган/подпись формы, а не клиент банка")
        elif is_caption(name):
            problems.append(f"наименование владельца {name!r} — подпись формы, а не имя")
        elif BANK_WORD_RE.search(name):
            problems.append(f"наименование владельца {name!r} — банк, а не клиент банка")
    if inn:
        if not inn.isdigit() or len(inn) not in (10, 12) or not inn_ok(inn):
            problems.append(f"ИНН владельца {inn!r} не проходит проверку (10/12 цифр и контрольная сумма)")
    return problems


NAME_TAIL_RE = re.compile(r"[,;]?\s*\b(?:инн|кпп|огрн|огрнип|р/с|р/сч|р\.с\.|сч[её]т|адрес|бик|тел\.?|телефон)\b.*$", re.I | re.S)


@dataclass
class OwnerCandidate:
    id: str
    kind: str
    value: str
    grid: str
    r: int
    c: int
    c2: int
    source: str
    score: float = 0.0
    features: list[str] = field(default_factory=list)

    def where(self) -> str:
        return f"{self.grid}!R{self.r}C{self.c}"


def owner_zone_cells(doc: Document, regions_by_grid: dict[str, list[Region]]) -> list[tuple[Grid, Cell, str]]:
    out = []
    for grid in doc.grids:
        regions = regions_by_grid.get(grid.name, [])
        if not regions:
            out.extend((grid, cell, "other_sheet") for cell in grid.cells)
            continue
        spans = [(reg.start_row, reg.end_row) for reg in regions]
        first = spans[0][0]
        for cell in grid.cells:
            if any(a <= cell.r <= b for a, b in spans):
                if not (any(reg.last_anchor < cell.r <= reg.end_row for reg in regions)
                        and (OWNER_LBL_RE.search(cell.text) or re.search(
                            r"инн|идентификационн|налогоплательщик|владел|клиент|в\s+отношении|организаци|наименование\s+клиента|плательщик\s+по\s+сч", cell.text, re.I))):
                    continue
            out.append((grid, cell, "preamble" if cell.r < first else "footer"))
    return out


def _split_digit_run(digits: str) -> list[tuple[str, str]]:
    return split_digit_run(digits)


def _digit_runs(zone: list[tuple[Grid, Cell, str]]) -> list[tuple[Grid, list[Cell], str]]:
    rows: dict[tuple[str, int], list[tuple[Grid, Cell, str]]] = defaultdict(list)
    for grid, cell, zname in zone:
        rows[(grid.name, cell.r)].append((grid, cell, zname))
    runs = []
    for key, items in rows.items():
        items.sort(key=lambda t: t[1].c)
        run: list[Cell] = []
        for grid, cell, zname in items + [(None, None, None)]:
            is_digit = cell is not None and re.fullmatch(r"[\d ]+", cell.text) and len(cell.text.replace(" ", "")) <= 22
            if is_digit and (not run or cell.c - run[-1].c2 <= 6 or (len(cell.text.strip()) == 1 and len(run[-1].text.strip()) == 1)):
                run.append(cell)
                continue
            if run:
                runs.append((items[0][0], run, items[0][2]))
            run = [cell] if is_digit else []
    return runs


SECT_BANK_RE = re.compile(r"наименование\s+банка|банка\s*\(филиала|представляющего\s+информацию|банк\s*\(филиал", re.I)
SECT_OWNER_RE = re.compile(r"в\s+отношении|наименование\s+организации|организации\)?,?\s*ф\.?\s?и\.?\s?о|"
                           r"индивидуального\s+предпринимателя|физического\s+лица|владелец|клиент|ин+/кио", re.I)


def section_above(grid: Grid, r: int, depth: int = 14) -> str:
    for dr in range(1, depth + 1):
        texts = [x.text for x in grid.row(r - dr)]
        for t in texts:
            if SECT_OWNER_RE.search(t) and not SECT_BANK_RE.search(t):
                return "owner"
        for t in texts:
            if SECT_BANK_RE.search(t):
                return "bank"
    return ""


def _context_text(grid: Grid, cell: Cell) -> tuple[str, str, str]:
    left = [x for x in grid.row(cell.r) if x.c2 < cell.c and not re.fullmatch(r"[\d ]+", x.text)]
    left_label = left[-1].text if left and cell.c - left[-1].c2 <= 12 else ""
    if left_label and (len(left_label) > 40 or OPF_RE.search(left_label)):
        left_label = ""
    above = ""
    if not left_label:
        for dr in (1, 2):
            cands = [x for x in grid.row(cell.r - dr) if x.c <= cell.c2 and x.c2 >= cell.c
                     and not re.fullmatch(r"[\d ]+", x.text)]
            if cands:
                if len(cands[0].text) <= 30:
                    above = cands[0].text
                break
    return left_label, above, grid.row_text(cell.r)


def _norm_name(text: str) -> str:
    t = text.casefold().replace("ё", "е")
    t = re.sub(OPF_FULL, " ", t, flags=re.I)
    t = re.sub(rf"(?<![а-яa-z])(?:{OPF_SHORT.lower()})(?![а-яa-z])", " ", t)
    return " ".join(re.findall(r"[0-9a-zа-я]+", t))


def names_match(a: str | None, b: str | None) -> bool:
    if not a or not b:
        return False
    na, nb = _norm_name(a), _norm_name(b)
    if not na or not nb:
        return False
    if na == nb or na in nb or nb in na:
        return True
    ta, tb = set(na.split()), set(nb.split())
    return len(ta & tb) / max(len(ta | tb), 1) >= 0.6


LABEL_WORDS = {"владелец", "владельца", "счета", "счёта", "счет", "счёт", "клиент", "клиента", "организация",
               "организации", "наименование", "полное", "сокращенное", "сокращённое", "плательщик", "получатель",
               "фио", "лицевого", "расчетного", "расчётного", "по", "предприятие", "предприятия", "абонент", "заемщик",
               "карты", "карта", "держатель", "держателя", "в", "отношении", "счету", "счёту", "account", "holder", "customer", "name"}


def clean_party_name(text: str) -> str:
    t = text.replace("\n", " ")
    m_opf = OPF_RE.search(t)
    if m_opf and 0 < m_opf.start() <= 40 and not re.search(r"[\d(«\"]", t[:m_opf.start()]) \
            and all(w.strip(":;,.").casefold() in LABEL_WORDS | {"расчетный", "расчётный", "текущий", "лицевой", "специальный"}
                    for w in t[:m_opf.start()].split()):
        t = t[m_opf.start():]
    t = NAME_TAIL_RE.sub("", t)
    tokens = t.split()
    i = 0
    while i < len(tokens) and all(part.strip(":;,.-–").casefold() in LABEL_WORDS
                                  for part in tokens[i].split("/") if part.strip(":;,.-–")):
        i += 1
    t = " ".join(tokens[i:])
    t = re.sub(r"(?<!\d)(?:\d{9,12}|\d{20})(?!\d)(?:\s*/\s*\d{9})?", " ", t)
    t = re.sub(r"\s+", " ", t).strip(" ,;:|-–/")
    return t


def table_party_stats(records_ops: list[dict[str, Any]]) -> dict[str, Any]:
    n = len(records_ops)
    inn_any: Counter = Counter()
    inn_cp: Counter = Counter()
    names_any: Counter = Counter()
    two_sided_rows = 0
    for op in records_ops:
        seen = set()
        for side in ("payer", "payee", "counterparty"):
            inn = op.get(f"{side}_inn")
            name = op.get(f"{side}_name")
            if inn and inn not in seen:
                inn_any[inn] += 1
                seen.add(inn)
                if side == "counterparty":
                    inn_cp[inn] += 1
            if name:
                names_any[_norm_name(name)] += 1
        if (op.get("payer_inn") or op.get("payer_name")) and (op.get("payee_inn") or op.get("payee_name")):
            two_sided_rows += 1
    return {"n": n, "inn_any": inn_any, "inn_cp": inn_cp, "names_any": names_any,
            "two_sided": n > 0 and two_sided_rows / n >= 0.6}


def detect_owner(doc: Document, regions_by_grid: dict[str, list[Region]],
                 table_ops: list[dict[str, Any]], llm: Any | None = None, allow_llm: bool = True) -> dict[str, Any]:
    zone = owner_zone_cells(doc, regions_by_grid)
    grids = {g.name: g for g in doc.grids}
    stats = table_party_stats(table_ops)
    n_ops = max(stats["n"], 1)

    raw_inns: list[OwnerCandidate] = []
    for grid, run, zname in _digit_runs(zone):
        digits = "".join(c.text.replace(" ", "") for c in run)
        if len(run) == 1 and len(digits) == 20:
            continue
        pieces = _split_digit_run(digits)
        if not any(k == "inn" for k, _ in pieces):
            pieces = [("inn", c.text.replace(" ", "")) for c in run if inn_ok(c.text.replace(" ", ""))]
        for kind, value in pieces:
            if kind == "inn":
                raw_inns.append(OwnerCandidate("", "inn", value, grid.name, run[0].r, run[0].c, run[-1].c2,
                                               "digits" if len(run) > 1 else "cell", features=[zname]))
    for grid, cell, zname in zone:
        if re.fullmatch(r"[\d ]+", cell.text):
            continue
        for inn in find_ids_in_text(cell.text).get("inn", []):
            pos = re.sub(r"(?<=\d)[  ](?=\d)", "", cell.text).find(inn)
            before = cell.text[:max(pos, 0)]
            feats = [zname, "in_text"]
            if re.search(r"инн[^\d]{0,15}$", before, re.I):
                feats.append("inn_label_in_text")
            last_bank = max((m.end() for m in BANK_WORD_RE.finditer(before)), default=-1)
            last_owner = max((m.end() for m in OWNER_LBL_RE.finditer(before)), default=-1)
            if last_bank > last_owner:
                feats.append("bank_in_text")
            elif last_owner >= 0:
                feats.append("owner_in_text")
            raw_inns.append(OwnerCandidate("", "inn", inn, grid.name, cell.r, cell.c, cell.c2, "text", features=feats))

    if not raw_inns:
        for grid, cell, zname in zone:
            left, above, _ = _context_text(grid, cell)
            digits = re.sub(r"\D", "", cell.text) if re.fullmatch(r"[\d ]+", cell.text) else ""
            labelled = re.search(r"инн", f"{left} {above}", re.I)
            if not digits:
                m = re.search(r"инн[^\d\n]{0,25}?(\d{10}|\d{12})(?!\d)", cell.text, re.I)
                digits, labelled = (m.group(1), True) if m else ("", False)
            if labelled and len(digits) in (10, 12):
                raw_inns.append(OwnerCandidate("", "inn", digits, grid.name, cell.r, cell.c, cell.c2, "checksum_failed",
                                               features=[zname, "контрольная сумма ИНН не сходится"]))

    inn_cands: dict[str, OwnerCandidate] = {}
    for cand in raw_inns:
        grid = grids[cand.grid]
        anchor = grid.at(cand.r, cand.c)
        left, above, row_text = _context_text(grid, anchor) if anchor else ("", "", "")
        score = 1.0
        feats = list(cand.features)
        if cand.source == "text":
            left = above = ""
        label_text = f"{left} {above}"
        if re.search(r"инн", label_text, re.I) or "inn_label_in_text" in feats:
            score += 2; feats.append("метка ИНН")
        owner_ctx = bool(OWNER_LBL_RE.search(label_text) or OWNER_LBL_RE.search(row_text) or "owner_in_text" in feats)
        bank_ctx = bool(BANK_WORD_RE.search(label_text) or "bank_in_text" in feats
                        or (not left and BANK_WORD_RE.search(row_text) and not OWNER_LBL_RE.search(row_text)))
        if owner_ctx and not bank_ctx:
            score += 4; feats.append("контекст владельца")
        if bank_ctx:
            score -= 8; feats.append("контекст банка")
        if cand.source in ("digits", "cell") and not owner_ctx and not bank_ctx:
            sect = section_above(grid, cand.r)
            if sect == "owner":
                score += 3; feats.append("раздел владельца")
            elif sect == "bank":
                score -= 3; feats.append("раздел банка")
        if re.search(r"инн\s*/\s*кио", label_text, re.I):
            score += 2; feats.append("ИНН/КИО")
        if cand.features[0] == "footer":
            score -= 1
        if cand.source == "checksum_failed":
            score -= 1.5
        freq = stats["inn_any"].get(cand.value, 0) / n_ops
        if stats["two_sided"] and freq >= 0.6:
            score += 10 if freq >= 0.8 else 6
            feats.append(f"в {freq:.0%} операций")
        elif stats["inn_cp"].get(cand.value):
            score -= 3; feats.append("встречается как контрагент")
        score -= min(0.01 * cand.r, 0.5)
        cand.score, cand.features = score, feats
        prev = inn_cands.get(cand.value)
        if prev is None or cand.score > prev.score:
            if prev is not None:
                cand.score += 0.5
            inn_cands[cand.value] = cand
        else:
            prev.score += 0.5
    if stats["two_sided"] and stats["inn_any"]:
        top_inn, count = stats["inn_any"].most_common(1)[0]
        if count / n_ops >= 0.6 and top_inn not in inn_cands:
            inn_cands[top_inn] = OwnerCandidate("", "inn", top_inn, "table", -1, -1, -1, "table",
                                                score=6.0, features=[f"в {count / n_ops:.0%} операций"])
    inns = sorted(inn_cands.values(), key=lambda x: -x.score)
    for i, cand in enumerate(inns, 1):
        cand.id = f"I{i}"
    best_inn = inns[0] if inns and inns[0].score > 0 else None

    raw_names: list[OwnerCandidate] = []
    zone_by_grid: dict[str, list[Cell]] = defaultdict(list)
    for grid, cell, zname in zone:
        zone_by_grid[grid.name].append(cell)

    def is_label_like(text: str) -> bool:
        t = text.strip()
        return bool(OTHER_LBL_RE.match(t) or is_caption(t) or re.fullmatch(r"[а-яa-z]\)", t)
                    or (OWNER_LBL_RE.search(t) and len(clean_party_name(t)) < 3))

    def usable_value(x: Cell) -> bool:
        return (not is_label_like(x.text) and not re.fullmatch(r"[\d ]+", x.text)
                and re.search(r"[А-Яа-яA-Za-z]{3,}", x.text) is not None and not BANK_WORD_RE.search(x.text))

    for grid, cell, zname in zone:
        text = cell.text
        if OWNER_LBL_RE.search(text) and not BANK_WORD_RE.search(text):
            inline = clean_party_name(text)
            if len(inline) >= 3 and re.search(r"[А-Яа-яA-Za-z]{2,}", inline) and not OWNER_LBL_RE.fullmatch(inline):
                if OWNER_LBL_RE.match(text.strip()):
                    raw_names.append(OwnerCandidate("", "name", inline, grid.name, cell.r, cell.c, cell.c2,
                                                    "inline_label", features=[zname, f"метка «{text[:20]}»"]))
            elif is_label_like(text) or len(text) <= 40:
                right = sorted((x for x in grid.row(cell.r) if x.c > cell.c2), key=lambda x: x.c)
                right_value = next((x for x in right[:1] if usable_value(x)), None)
                above = [x for dr in (1, 2) for x in grid.row(cell.r - dr) if x.c <= cell.c2 and x.c2 >= cell.c]
                below = [x for dr in (1, 2, 3) for x in grid.row(cell.r + dr) if x.c <= cell.c2 and x.c2 >= cell.c]
                if is_caption(text):
                    value = next((x for x in above if usable_value(x)), None) or right_value
                else:
                    value = right_value or next((x for x in below if usable_value(x)), None)
                if value is not None:
                    raw_names.append(OwnerCandidate("", "name", clean_party_name(value.text), grid.name, value.r, value.c,
                                                    value.c2, "label", features=[zname, f"метка «{text[:25]}»"]))
        m = OPF_RE.search(text)
        cell_left, _, _ = _context_text(grid, cell)
        if m and not BANK_WORD_RE.search(text) and not BANK_WORD_RE.search(cell_left) and not is_caption(text):
            raw_names.append(OwnerCandidate("", "name", clean_party_name(text[m.start():]), grid.name, cell.r, cell.c,
                                            cell.c2, "opf", features=[zname]))
        elif PERSON_RE.match(text.strip()) or PERSON_INITIALS_RE.match(text.strip()):
            raw_names.append(OwnerCandidate("", "name", text.strip(), grid.name, cell.r, cell.c, cell.c2,
                                            "person", features=[zname]))

    name_cands: dict[str, OwnerCandidate] = {}
    for cand in raw_names:
        if len(cand.value) < 3 or not re.search(r"[А-Яа-яA-Za-z]{2,}", cand.value) or not _norm_name(cand.value) \
                or is_caption(cand.value):
            continue
        score = 0.0
        feats = list(cand.features)
        if cand.source in ("label", "inline_label"):
            val = cand.value.strip()
            if looks_like_party_name(val):
                score += 4
            elif re.search(r"\d", val) or val[:1].islower():
                score -= 1
            else:
                score += 1
        if OPF_RE.search(cand.value):
            score += 2; feats.append("форма собственности")
        if cand.source == "person":
            score += 2; feats.append("ФИО")
        if best_inn is not None and best_inn.grid == cand.grid and best_inn.r >= 0:
            dist = abs(cand.r - best_inn.r) * 3 + max(0, cand.c - best_inn.c2, best_inn.c - cand.c2)
            if cand.r < best_inn.r:
                dist += 1
            score -= min(0.3 * dist, 3.0)
            if dist == 0:
                feats.append("в одной ячейке с ИНН")
        if cand.features[0] == "footer":
            score -= 1
        norm = _norm_name(cand.value)
        freq = stats["names_any"].get(norm, 0) / n_ops
        if stats["two_sided"] and freq >= 0.6:
            score += 8 if freq >= 0.8 else 5; feats.append(f"в {freq:.0%} операций")
        elif freq > 0 and not stats["two_sided"]:
            score -= 3; feats.append("встречается как контрагент")
        cand.score, cand.features = score, feats
        prev = name_cands.get(norm)
        if prev is None or cand.score > prev.score:
            if prev is not None:
                cand.score += 0.5
            name_cands[norm] = cand
    if stats["two_sided"] and stats["names_any"]:
        top_name, count = stats["names_any"].most_common(1)[0]
        if count / n_ops >= 0.6 and top_name not in name_cands:
            original = next((op.get(f"{side}_name") for op in table_ops for side in ("payer", "payee")
                             if op.get(f"{side}_name") and _norm_name(op[f"{side}_name"]) == top_name), top_name)
            name_cands[top_name] = OwnerCandidate("", "name", original, "table", -1, -1, -1, "table",
                                                  score=5.0, features=[f"в {count / n_ops:.0%} операций"])
    names = sorted(name_cands.values(), key=lambda x: -x.score)
    for i, cand in enumerate(names, 1):
        cand.id = f"N{i}"
    best_name = names[0] if names else None

    conflicts = []
    labelled = [c for c in names if c.source in ("label", "inline_label")]
    table_name = next((c for c in names if any(f.startswith("в ") and "операций" in f for f in c.features)), None)
    if labelled and table_name and not names_match(labelled[0].value, table_name.value):
        conflicts.append(f"в шапке владелец «{labelled[0].value}», а в каждой операции стоит «{table_name.value}»")
        best_name = labelled[0]
    labelled_inn = [c for c in inns if "контекст владельца" in c.features or "метка ИНН" in c.features]
    table_inn = next((c for c in inns if any(f.startswith("в ") and "операций" in f for f in c.features)), None)
    if labelled_inn and table_inn and labelled_inn[0].value != table_inn.value and "контекст банка" not in labelled_inn[0].features:
        conflicts.append(f"в шапке ИНН {labelled_inn[0].value}, а в каждой операции {table_inn.value}")
        best_inn = labelled_inn[0]

    def conf(cands: list[OwnerCandidate]) -> float:
        if not cands or cands[0].score <= 0:
            return 0.0
        c = min(0.5 + 0.07 * cands[0].score, 0.97)
        if len(cands) > 1 and cands[0].score - cands[1].score < 2:
            c -= 0.25
        return round(max(c, 0.05), 2)

    inn_conf, name_conf = conf(inns), conf(names)
    if conflicts:
        inn_conf, name_conf = min(inn_conf, 0.5), min(name_conf, 0.5)
    result = {
        "inn": best_inn.value if best_inn else None,
        "name": best_name.value if best_name else None,
        "inn_confidence": inn_conf,
        "name_confidence": name_conf,
        "method": "algorithm",
        "evidence": {
            "inn": best_inn.where() + " " + ", ".join(best_inn.features) if best_inn else None,
            "name": best_name.where() + " " + ", ".join(best_name.features) if best_name else None,
        },
        "conflicts": conflicts,
        "inn_candidates": [(c.id, c.value, round(c.score, 1), c.where(), c.features) for c in inns[:8]],
        "name_candidates": [(c.id, c.value, round(c.score, 1), c.where(), c.features) for c in names[:8]],
    }
    need_llm = min(inn_conf if inns else 1.0, name_conf) < OWNER_LLM_THRESHOLD or (not inns and not names)
    if need_llm and allow_llm and llm is not None and (inns or names or zone):
        _owner_arbiter(result, zone, inns, names, stats, llm)
    result["confidence"] = round(min(result["inn_confidence"] if result["inn"] else 0.5,
                                     result["name_confidence"] if result["name"] else 0.5), 2)
    if allow_llm:
        log("владелец", f"ИНН={result['inn']} ({result['inn_confidence']}), "
                        f"наименование={result['name']!r} ({result['name_confidence']}), метод={result['method']}")
    return result


def _owner_arbiter(result: dict[str, Any], zone, inns: list[OwnerCandidate], names: list[OwnerCandidate],
                   stats: dict[str, Any], llm: Any | None) -> None:
    cells_by_ref = {}
    lines = []
    for grid, cell, zname in zone[:400]:
        ref = f"R{cell.r}C{cell.c}"
        if len(zone) and len({g.name for g, _, _ in zone}) > 1:
            ref = f"{grid.name}!{ref}"
        cells_by_ref[ref] = cell
        if len(lines) < 120:
            lines.append(f"{ref}: {cell.text[:80]}")
    inn_ids = {c.id: c for c in inns[:8]}
    name_ids = {c.id: c for c in names[:8]}
    top_parties = ", ".join(f"{inn}×{cnt}" for inn, cnt in stats["inn_any"].most_common(3))
    user = (
        "ЯЧЕЙКИ ШАПКИ ВЫПИСКИ:\n" + "\n".join(lines)
        + "\n\nКАНДИДАТЫ ИНН:\n" + ("\n".join(f"{c.id}: {c.value} ({c.where()}; {', '.join(c.features)})" for c in inn_ids.values()) or "нет")
        + "\n\nКАНДИДАТЫ НАИМЕНОВАНИЯ:\n" + ("\n".join(f"{c.id}: {c.value} ({c.where()}; {', '.join(c.features)})" for c in name_ids.values()) or "нет")
        + f"\n\nОПЕРАЦИЙ В ТАБЛИЦЕ: {stats['n']}; самые частые ИНН сторон: {top_parties or 'нет'}"
    )

    def validate(data):
        if not isinstance(data, dict):
            return None, "ожидался JSON-объект"
        inn_id, name_id, name_cell = data.get("inn_id"), data.get("name_id"), data.get("name_cell")
        if inn_id is not None and inn_id not in inn_ids:
            return None, f"inn_id={inn_id!r} нет среди кандидатов {list(inn_ids)}"
        if name_id is not None and name_id not in name_ids:
            return None, f"name_id={name_id!r} нет среди кандидатов {list(name_ids)}"
        def _in_ops(c) -> int:
            for f in c.features:
                m_ = re.match(r"в (\d+)% операций", f)
                if m_:
                    return int(m_.group(1))
            return 0

        for kind, chosen, pool in (("ИНН", inn_id, inn_ids), ("наименование", name_id, name_ids)):
            if chosen is None:
                continue
            strong = [i for i, c in pool.items() if _in_ops(c) >= 80]
            if strong and chosen not in strong and _in_ops(pool[chosen]) < 80:
                return None, f"{kind} {pool[strong[0]].value!r} ({strong[0]}) стоит в {_in_ops(pool[strong[0]])}% операций таблицы — это владелец; выбор {chosen} ему противоречит"
        if name_id is not None and not looks_like_party_name(name_ids[name_id].value) and \
                any(looks_like_party_name(c.value) for c in name_ids.values()):
            return None, (f"{name_id} ({name_ids[name_id].value!r}) не похоже на наименование организации или ФИО; "
                          f"есть кандидаты с формой собственности — выбери из них")
        if inn_id is not None and "контекст банка" in inn_ids[inn_id].features and _in_ops(inn_ids[inn_id]) < 80:
            others = [i for i, c in inn_ids.items() if "контекст банка" not in c.features]
            if others:
                return None, f"{inn_id} — ИНН банка (рядом банковские реквизиты). Выбери из {others} или null"
        if name_cell:
            if name_cell not in cells_by_ref:
                return None, f"ячейки {name_cell} нет в шапке"
            text = clean_party_name(cells_by_ref[name_cell].text)
            if not re.search(r"[А-Яа-яA-Za-z]{3,}", text) or BANK_WORD_RE.search(text):
                return None, f"в ячейке {name_cell} нет наименования клиента"
        try:
            confidence = float(data.get("confidence", 0.5))
        except (TypeError, ValueError):
            return None, "confidence должно быть числом"
        return {"inn_id": inn_id, "name_id": name_id, "name_cell": name_cell, "confidence": confidence,
                "reason": str(data.get("reason", ""))[:200]}, None

    answer = call_llm_json("owner", OWNER_SYSTEM, user, validate, llm=llm)
    if answer is None:
        result["llm"] = "не дал проверяемого ответа — оставлен результат алгоритма"
        return
    changed = []
    if answer["inn_id"]:
        new = inn_ids[answer["inn_id"]].value
        if new != result["inn"]:
            changed.append(f"ИНН {result['inn']} → {new}")
        result["inn"] = new
        result["inn_confidence"] = max(result["inn_confidence"], round(0.6 + 0.35 * answer["confidence"], 2))
    if answer["name_id"] or answer["name_cell"]:
        new = name_ids[answer["name_id"]].value if answer["name_id"] else clean_party_name(cells_by_ref[answer["name_cell"]].text)
        if new != result["name"]:
            changed.append(f"наименование {result['name']!r} → {new!r}")
        result["name"] = new
        result["name_confidence"] = max(result["name_confidence"], round(0.6 + 0.35 * answer["confidence"], 2))
    result["method"] = "algorithm+gigachat"
    result["llm"] = {"answer": answer, "changed": changed}


# сборка операций
def split_composite(text: str) -> tuple[str | None, str | None]:
    if not text:
        return None, None
    ids = find_ids_in_text(text)
    inn = ids.get("inn", [None])[0]
    name = text
    name = re.sub(r"\b(?:инн|кпп|огрн|бик|р/с|р/сч|к/с|сч[её]т|л/с)\b\s*[:№]?\s*[\d/ ]{5,}", " ", name, flags=re.I)
    name = re.sub(r"(?<!\d)\d{9,20}(?!\d)", " ", name)
    name = name.strip()
    cut = re.search(r"\s(?:в\s+)?(?:банк|пао\s+сбербанк|филиал)\b", name, re.I)
    if cut and cut.start() > 0 and not OPF_RE.match(name):
        name = name[:cut.start()]
    name = re.sub(r"\s+", " ", name).strip(" ,;:/|-–")
    return inn, (name or None)


def _party(values: dict[str, str], side: str) -> tuple[str | None, str | None]:
    inn_text = values.get(f"{side}.inn")
    name = values.get(f"{side}.name")
    comp = values.get(f"{side}.composite")
    inn = None
    if inn_text:
        digits = re.sub(r"\D", "", inn_text)
        digits = fix_inn(digits)
        inn = digits if len(digits) in (10, 12) else (find_ids_in_text(inn_text).get("inn", [None])[0])
    if comp:
        c_inn, c_name = split_composite(comp)
        inn = inn or c_inn
        name = name or c_name
    if name:
        name = clean_party_name(name) or name
    return inn, name


def record_direction(rec: OpRecord, spec: TableSpec) -> tuple[str | None, Decimal | None]:
    a = rec.amounts
    if spec.debit is not None or spec.credit is not None:
        d, c = a.get("debit"), a.get("credit")
        d_on = d not in (None, Decimal(0))
        c_on = c not in (None, Decimal(0))
        if d_on and not c_on:
            return "out", abs(d)
        if c_on and not d_on:
            return "in", abs(c)
        if d_on and c_on:
            return None, None
        return rec.values.get("_dir_hint"), None
    if spec.signed is not None:
        v = a.get("signed")
        if v in (None, Decimal(0)):
            return None, None
        return ("out" if v < 0 else "in"), abs(v)
    if spec.amount is not None:
        v = a.get("amount")
        if v in (None, Decimal(0)):
            return None, None
        raw = (rec.direction_raw or "").strip().casefold().rstrip(".")
        if raw in {x.casefold() for x in spec.direction_out}:
            return "out", abs(v)
        if raw in {x.casefold() for x in spec.direction_in}:
            return "in", abs(v)
        return None, abs(v)
    return None, None


def party_is_owner(inn: str | None, name: str | None, owner: dict[str, Any]) -> bool:
    if inn and owner.get("inn"):
        return inn == owner["inn"]
    return names_match(name, owner.get("name"))


def records_to_raw_ops(records: list[OpRecord], spec: TableSpec) -> list[dict[str, Any]]:
    ops = []
    for rec in records:
        direction, amount = record_direction(rec, spec)
        op = {"date": rec.date, "amount": amount, "direction": direction,
              "purpose": rec.values.get("purpose"), "_rows": rec.rows, "_grid": rec.grid_name,
              "_date_inherited": rec.date_inherited, "_provenance": rec.provenance,
              "_dir_by_amount": (spec.debit is not None or spec.credit is not None or spec.signed is not None)
                                 and not any("по соглашению" in n for n in spec.notes)}
        for side in ("payer", "payee", "counterparty"):
            op[f"{side}_inn"], op[f"{side}_name"] = _party(rec.values, side)
        ops.append(op)
    return ops


def finalize_ops(raw_ops: list[dict[str, Any]], owner: dict[str, Any]) -> tuple[list[dict[str, Any]], list[str]]:
    notes = []
    two_sided = any(op.get("payer_inn") or op.get("payer_name") for op in raw_ops) and \
        any(op.get("payee_inn") or op.get("payee_name") for op in raw_ops)

    for op in raw_ops:
        if op["direction"] is None and op["amount"] is not None:
            if party_is_owner(op.get("payer_inn"), op.get("payer_name"), owner):
                op["direction"] = "out"
            elif party_is_owner(op.get("payee_inn"), op.get("payee_name"), owner):
                op["direction"] = "in"

    if two_sided and (owner.get("inn") or owner.get("name")):
        agree = disagree = 0
        for op in raw_ops:
            if op["direction"] is None or op.get("_dir_by_amount"):
                continue
            payer_owner = party_is_owner(op.get("payer_inn"), op.get("payer_name"), owner)
            payee_owner = party_is_owner(op.get("payee_inn"), op.get("payee_name"), owner)
            if payer_owner == payee_owner:
                continue
            ok = (op["direction"] == "out") == payer_owner
            agree += ok
            disagree += not ok
        if disagree > agree and disagree >= 2:
            for op in raw_ops:
                if op["direction"] and not op.get("_dir_by_amount"):
                    op["direction"] = "in" if op["direction"] == "out" else "out"
            notes.append(f"направление перевёрнуто: владелец был плательщиком в зачислениях ({disagree} против {agree})")

    if two_sided and (owner.get("inn") or owner.get("name")):
        for op in raw_ops:
            if op.get("_dir_by_amount"):
                continue
            pay = party_is_owner(op.get("payer_inn"), op.get("payer_name"), owner)
            rec = party_is_owner(op.get("payee_inn"), op.get("payee_name"), owner)
            if pay != rec:
                op["direction"] = "out" if pay else "in"

    out = []
    for op in raw_ops:
        row = {k: None for k in OUTPUT_COLUMNS}
        row["date"] = op["date"]
        row["amount"] = op["amount"]
        row["purpose"] = op["purpose"]
        owner_inn, owner_name = owner.get("inn"), owner.get("name")
        if two_sided:
            row["debit_inn"], row["debit_name"] = op.get("payer_inn"), op.get("payer_name")
            row["credit_inn"], row["credit_name"] = op.get("payee_inn"), op.get("payee_name")
            cp_inn, cp_name = op.get("counterparty_inn"), op.get("counterparty_name")
            owner_side = {"out": "debit", "in": "credit"}.get(op["direction"])
            other = {"debit": "credit", "credit": "debit"}.get(owner_side)
            if owner_side and (cp_inn or cp_name) and not row[f"{other}_inn"] and not row[f"{other}_name"]:
                row[f"{other}_inn"], row[f"{other}_name"] = cp_inn, cp_name
        else:
            cp_inn = op.get("counterparty_inn") or op.get("payer_inn") or op.get("payee_inn")
            cp_name = op.get("counterparty_name") or op.get("payer_name") or op.get("payee_name")
            if op["direction"] == "out":
                row["debit_inn"], row["debit_name"] = owner_inn, owner_name
                row["credit_inn"], row["credit_name"] = cp_inn, cp_name
            elif op["direction"] == "in":
                row["debit_inn"], row["debit_name"] = cp_inn, cp_name
                row["credit_inn"], row["credit_name"] = owner_inn, owner_name
        row["_direction"] = op["direction"]
        row["_rows"] = op["_rows"]
        row["_grid"] = op["_grid"]
        row["_date_inherited"] = op["_date_inherited"]
        out.append(row)
    return out, notes


# проверки
TEXT_AMOUNT_RE = r"(-?\d{1,3}(?:[  ]\d{3})+(?:[.,]\d{1,2})?|-?\d+[.,]\d{2})(?!\d|[.,]\d)"
REF_PATTERNS = {
    "debit": re.compile(rf"(?:дебет\w*|списан\w*|расход\w*|debit)[^0-9\n]{{0,40}}?(?<![\d.,]){TEXT_AMOUNT_RE}", re.I),
    "credit": re.compile(rf"(?:кредит\w*|зачислен\w*|поступ\w*|приход\w*|credit)[^0-9\n]{{0,40}}?(?<![\d.,]){TEXT_AMOUNT_RE}", re.I),
    "balance_in": re.compile(rf"(?:(?:входящ|начальн)\w*\s+(?:остат|сальдо)\w*|(?:остат|сальдо)\w*\s+(?:на\s+начало|входящ)\w*)[^0-9\n]{{0,40}}?(?<![\d.,]){TEXT_AMOUNT_RE}", re.I),
    "balance_out": re.compile(rf"(?:(?:исходящ|конечн)\w*\s+(?:остат|сальдо)\w*|(?:остат|сальдо)\w*\s+(?:на\s+конец|исходящ)\w*)[^0-9\n]{{0,40}}?(?<![\d.,]){TEXT_AMOUNT_RE}", re.I),
}
PERIOD_RE = re.compile(r"(?:период|за\s+период|с)\s*[:]?\s*(\d{1,2}\.\d{1,2}\.\d{2,4})\s*(?:г\.?)?\s*(?:по|-|—)\s*(\d{1,2}\.\d{1,2}\.\d{2,4})", re.I)


def _amount_from_text(text: str) -> Decimal | None:
    return parse_amount(text.replace(" ", " "))


def find_references(doc: Document, regions: list[Region], verdicts_by_region: dict[int, list[RowVerdict]],
                    specs: dict[int, TableSpec]) -> dict[str, Any]:
    refs: dict[str, Any] = {}
    sources: dict[str, str] = {}
    for i, region in enumerate(regions):
        spec = specs[i]
        grid = region.grid
        totals = [v for v in verdicts_by_region[i] if v.label == TOTAL]
        if totals and spec.debit is not None and spec.credit is not None:
            finals = [v for v in totals if v.r >= region.last_anchor]
            chosen = finals or []
            if chosen:
                chosen.sort(key=lambda v: (bool(re.search(r"период|всего|оборот", grid.row_text(v.r), re.I)), v.r))
                v = chosen[-1]
                rows = [v]
            else:
                rows = []
            for v in rows:
                d = parse_amount(grid.text_at(v.r, spec.debit)) if grid.at(v.r, spec.debit) and grid.at(v.r, spec.debit).r == v.r else None
                c = parse_amount(grid.text_at(v.r, spec.credit)) if grid.at(v.r, spec.credit) and grid.at(v.r, spec.credit).r == v.r else None
                if d is not None or c is not None:
                    refs[f"debit#{i}"] = abs(d) if d is not None else Decimal(0)
                    refs[f"credit#{i}"] = abs(c) if c is not None else Decimal(0)
                    sources[f"turnover#{i}"] = f"{grid.name}!R{v.r}: {grid.row_text(v.r)[:80]}"
        for label, key in ((BALANCE_IN, "balance_in"), (BALANCE_OUT, "balance_out")):
            cands = []
            for v in verdicts_by_region[i]:
                if v.label != label:
                    continue
                amounts = [parse_amount(c.text) for c in grid.row(v.r) if "amount" in region.kinds(c)]
                amounts = [a for a in amounts if a is not None]
                if amounts:
                    cands.append((v.r, amounts[-1]))
            if not cands or f"{key}#{i}" in refs:
                continue
            if key == "balance_in":
                before = [a for r, a in cands if r <= region.first_anchor]
                r_take, a_take = (region.first_anchor, before[0]) if before else cands[0]
            else:
                after = [(r, a) for r, a in cands if r >= region.last_anchor]
                if not after:
                    continue
                r_take, a_take = after[-1]
            refs[f"{key}#{i}"] = a_take
            sources[f"{key}#{i}"] = f"{grid.name}!R{r_take}"
    in_table = {(r.grid.name, row) for r in regions for row in range(r.start_row, r.last_anchor + 1)}
    _table1_references(doc, in_table, refs, sources)
    for grid in doc.grids:
        for row in range(grid.n_rows):
            if (grid.name, row) in in_table:
                continue
            text = grid.row_text(row, sep=" ")
            if not text:
                continue
            m = PERIOD_RE.search(text)
            if m and "period" not in refs:
                start, end = parse_date(m.group(1)), parse_date(m.group(2))
                if start and end:
                    refs["period"] = (start, end)
            for key, pattern in REF_PATTERNS.items():
                if any(k.startswith(key + "#") for k in refs):
                    continue
                if key in ("debit", "credit") and not re.search(r"оборот|итог|всего", text, re.I):
                    continue
                m = pattern.search(text)
                if m:
                    value = _amount_from_text(m.group(1))
                    if value is not None:
                        refs[f"{key}#text"] = abs(value) if key in ("debit", "credit") else value
                        sources[f"{key}#text"] = f"{grid.name}!R{row}: {text[:80]}"
    refs["_sources"] = sources
    return refs


T1_PATTERNS = {
    "balance_in": re.compile(r"остат\w*[^\n]{0,60}на\s+начало|входящ\w*\s+(?:остат|сальдо)", re.I),
    "debit": re.compile(r"(?:сумма|оборот\w*)\s+по\s+дебету|дебетов\w+\s+оборот", re.I),
    "credit": re.compile(r"(?:сумма|оборот\w*)\s+по\s+кредиту|кредитов\w+\s+оборот", re.I),
    "balance_out": re.compile(r"остат\w*[^\n]{0,60}на\s+конец|исходящ\w*\s+(?:остат|сальдо)", re.I),
}


def _table1_references(doc: Document, in_table: set, refs: dict[str, Any], sources: dict[str, str]) -> None:
    for grid in doc.grids:
        for cell in grid.cells:
            if (grid.name, cell.r) in in_table or len(cell.text) < 10:
                continue
            for key, pattern in T1_PATTERNS.items():
                if not pattern.search(cell.text) or any(k.startswith(key + "#") for k in refs):
                    continue
                if key in ("debit", "credit") and re.search(r"операци", cell.text, re.I):
                    continue
                for dr in (1, 2, 3):
                    row = grid.row(cell.r2 + dr)
                    if _is_numbering_row(row) or (grid.name, cell.r2 + dr) in in_table:
                        continue
                    below = [x for x in row if x.c <= cell.c2 and x.c2 >= cell.c]
                    value = next((parse_amount(x.text) for x in below if "amount" in cell_kinds(x.text, x.value)), None)
                    if value is not None:
                        refs[f"{key}#t1"] = abs(value) if key in ("debit", "credit") else value
                        sources[f"{key}#t1"] = f"{grid.name}!R{cell.r}: {cell.text[:60]} → {value}"
                        break
                    if below:
                        break


def semantic_issues(ops: list[dict[str, Any]], owner: dict[str, Any]) -> list[tuple[str, str]]:
    issues = []
    n = len(ops)
    if not n:
        return issues
    purposes = [op["purpose"] for op in ops if op.get("purpose")]
    if purposes:
        good = sum(1 for p in purposes if len(p) >= 10 and re.search(r"[А-Яа-яA-Za-z]{3,}", p))
        if good / len(purposes) < 0.5:
            examples = ", ".join(sorted({p for p in purposes})[:3])
            issues.append(("purpose_codes", f"назначение похоже на коды/короткие значения ({examples}) — "
                                            f"вероятно, взята не та колонка"))
    leaked = [n for op in ops for n in (op.get("debit_name"), op.get("credit_name")) if n and looks_like_purpose(n)]
    if leaked:
        issues.append(("purpose_in_party", f"в наименованиях сторон {len(leaked)} значений похожи на назначение платежа "
                                           f"(например, «{leaked[0][:60]}») — текст назначения попал в колонку стороны"))
    same = sum(1 for op in ops if (op.get("debit_inn") and op["debit_inn"] == op.get("credit_inn"))
               or (op.get("debit_name") and op.get("debit_name") == op.get("credit_name")))
    if same / n >= 0.3 and same >= 2:
        issues.append(("same_party_both_sides", f"у {same} из {n} операций одна и та же сторона и в дебете, и в кредите — "
                                                f"плательщик и получатель взяты из одной колонки или владелец принят за контрагента"))
    if owner.get("inn"):
        with_dir = [op for op in ops if op["_direction"] in ("in", "out")]
        wrong = sum(1 for op in with_dir if (op["_direction"] == "in" and op.get("debit_inn") == owner["inn"])
                    or (op["_direction"] == "out" and op.get("credit_inn") == owner["inn"]))
        if len(with_dir) >= 3 and wrong / len(with_dir) >= 0.1:
            issues.append(("owner_on_wrong_side", f"у {wrong} из {len(with_dir)} операций владелец стоит не на своей стороне "
                                                  f"(при зачислении — в дебете или при списании — в кредите): направление или стороны определены неверно"))
    cps = []
    for op in ops:
        side = "credit" if op["_direction"] == "out" else "debit" if op["_direction"] == "in" else None
        if side:
            cps.append((op.get(f"{side}_inn"), op.get(f"{side}_name")))
    if len(cps) >= 3 and (owner.get("inn") or owner.get("name")):
        is_owner = sum(1 for inn, name in cps
                       if (inn and owner.get("inn") and inn == owner["inn"]) or (not inn and name and names_match(name, owner.get("name"))))
        if is_owner / len(cps) >= 0.3:
            issues.append(("owner_as_counterparty", f"у {is_owner} из {len(cps)} операций контрагент — сам владелец выписки: "
                                                    f"взяты колонки владельца вместо колонок противоположной стороны"))
        names = [n for _, n in cps if n]
        junk = sum(1 for n in names if sum(ch.isalpha() for ch in n) < 0.5 * len(n.replace(" ", "")))
        if names and junk / len(names) >= 0.3:
            issues.append(("party_not_text", f"у {junk} из {len(names)} контрагентов «наименование» — в основном цифры "
                                             f"(счёт, БИК, номер): взята не та колонка"))
        inns = [i for i, _ in cps if i]
        bad = sum(1 for i in inns if not inn_ok(i))
        if len(inns) >= 3 and bad / len(inns) >= 0.5:
            issues.append(("party_inn_invalid", f"у {bad} из {len(inns)} ИНН контрагентов не сходится контрольная сумма — "
                                                f"вероятно, в колонке ИНН не ИНН (КПП, счёт, номер документа)"))
    cp_names = []
    for op in ops:
        side = "credit" if op["_direction"] == "out" else "debit" if op["_direction"] == "in" else None
        if side and op[f"{side}_name"]:
            cp_names.append(op[f"{side}_name"])
    bank_share = sum(1 for x in cp_names if BANK_WORD_RE.search(x)) / max(len(cp_names), 1)
    branch_share = sum(1 for x in cp_names if re.search(r"филиал|отделение|головн|регионал|\bбанк\s+(?:пао|ао)\b", x, re.I)) / max(len(cp_names), 1)
    if cp_names and (bank_share >= 0.85 or (bank_share >= 0.5 and branch_share >= 0.4)):
        issues.append(("bank_as_counterparty", "в контрагентах в основном банки — вероятно, взята колонка "
                                               "«банк плательщика/получателя» вместо самого плательщика/получателя"))
    if owner.get("name") and is_caption(owner["name"]):
        issues.append(("owner_caption", f"наименование владельца «{owner['name']}» — подпись формы, а не имя"))
    totals = sum(1 for op in ops if op.get("purpose") and TOTAL_RE.search(op["purpose"])
                 and not (op["debit_inn"] and op["credit_inn"]))
    if totals:
        issues.append(("subtotals_in_ops", f"{totals} операций похожи на итоговые строки («итого», «оборот»)"))
    return issues


def run_checks(ops: list[dict[str, Any]], refs: dict[str, Any], verdicts_by_region: dict[int, list[RowVerdict]],
               regions: list[Region], specs: dict[int, TableSpec], owner: dict[str, Any], flipped: bool = False) -> dict[str, Any]:
    checks: list[dict[str, Any]] = []
    total_out = sum((op["amount"] for op in ops if op["_direction"] == "out" and op["amount"] is not None), Decimal(0))
    total_in = sum((op["amount"] for op in ops if op["_direction"] == "in" and op["amount"] is not None), Decimal(0))

    def ref(key):
        region_vals = [v for k, v in refs.items() if k.startswith(key + "#") and not k.endswith("#text")]
        if region_vals:
            if key in ("debit", "credit"):
                return sum(region_vals, Decimal(0))
            return region_vals[0] if len(regions) <= 1 else None
        return refs.get(f"{key}#text")

    def add(name, status, expected=None, actual=None, detail="", level="error"):
        checks.append({"check": name, "status": status, "level": level,
                       "expected": None if expected is None else str(expected),
                       "actual": None if actual is None else str(actual), "detail": detail})

    ref_out, ref_in = ref("debit"), ref("credit")
    with_refs = [i for i in range(len(regions)) if f"debit#{i}" in refs or f"credit#{i}" in refs]
    if with_refs and len(with_refs) < len(regions):
        names = {regions[i].grid.name for i in with_refs}
        cmp_out = sum((op["amount"] for op in ops if op["_direction"] == "out" and op["amount"] is not None and op["_grid"] in names), Decimal(0))
        cmp_in = sum((op["amount"] for op in ops if op["_direction"] == "in" and op["amount"] is not None and op["_grid"] in names), Decimal(0))
    else:
        cmp_out, cmp_in = total_out, total_in
    for name, expected, actual in (("обороты по дебету (списания)", ref_out, cmp_out),
                                   ("обороты по кредиту (зачисления)", ref_in, cmp_in)):
        if expected is None:
            add(name, "skip", detail="итог в документе не найден", level="info")
        elif abs(expected - actual) < Decimal("0.01"):
            add(name, "pass", expected, actual)
        else:
            add(name, "fail", expected, actual, detail=f"разница {expected - actual}")

    b_in, b_out = ref("balance_in"), ref("balance_out")
    if b_in is not None and b_out is not None:
        calc = b_in + total_in - total_out
        if abs(calc - b_out) < Decimal("0.01"):
            add("остатки: вход + зачисления − списания = исход", "pass", b_out, calc)
        else:
            add("остатки: вход + зачисления − списания = исход", "fail", b_out, calc, detail=f"разница {b_out - calc}")
    else:
        add("остатки", "skip", detail="входящий и/или исходящий остаток не найдены", level="info")

    for i, region in enumerate(regions):
        spec = specs[i]
        if not spec.balance:
            continue
        grid, bcol = region.grid, spec.balance[0]
        prev = None
        hits = total = 0
        for op in ops:
            if op["_grid"] != grid.name or not op["_rows"]:
                continue
            bal = parse_amount(grid.text_at(op["_rows"][0], bcol))
            if bal is None or op["amount"] is None or op["_direction"] is None:
                prev = bal
                continue
            if prev is not None:
                total += 1
                signed = op["amount"] if op["_direction"] == "in" else -op["amount"]
                hits += abs(prev + signed - bal) < Decimal("0.01")
            prev = bal
        if total >= 2:
            add(f"текущий остаток ({grid.name})", "pass" if hits == total else "fail",
                detail=f"сошлось {hits} из {total} переходов")

    for i, region in enumerate(regions):
        spec = specs[i]
        if spec.debit is None or spec.credit is None:
            continue
        grid = region.grid
        two_sided_spec = any(v is not None for v in spec.payer.values()) and any(v is not None for v in spec.payee.values())
        group_out = group_in = Decimal(0)
        total_out_sum = total_in_sum = Decimal(0)
        chains: dict[str, list[Decimal]] = {}
        ops_rows = {r: op for op in ops if op["_grid"] == grid.name for r in op["_rows"][:1]}
        bad = []
        confirmed = 0

        def kind_of(row: int) -> str:
            first = next((c.text for c in grid.row(row) if c.text.strip()), "")
            return re.sub(r"\s+", " ", re.sub(r"[\d.,:/\-]+", " ", first.casefold())).strip()

        for v in verdicts_by_region[i]:
            if v.r in ops_rows:
                op = ops_rows[v.r]
                if op["amount"] is not None:
                    k = 0 if op["_direction"] == "out" else 1 if op["_direction"] == "in" else None
                    if k is not None:
                        if k == 0:
                            group_out += op["amount"]; total_out_sum += op["amount"]
                        else:
                            group_in += op["amount"]; total_in_sum += op["amount"]
                        for ch in chains.values():
                            ch[k] += op["amount"]
            elif v.label == SUBTOTAL and v.reason.startswith("ключевое слово итога"):
                d = parse_amount(grid.text_at(v.r, spec.debit)) if grid.at(v.r, spec.debit) and grid.at(v.r, spec.debit).r == v.r else None
                c = parse_amount(grid.text_at(v.r, spec.credit)) if grid.at(v.r, spec.credit) and grid.at(v.r, spec.credit).r == v.r else None
                kind = kind_of(v.r)
                chain = chains.setdefault(kind, [total_out_sum, total_in_sum])
                candidates = [(group_out, group_in), (chain[0], chain[1]), (total_out_sum, total_in_sum)]
                if two_sided_spec or flipped:
                    candidates += [(gi, go) for go, gi in candidates]
                ok = any((d is None or abs(abs(d) - go) < Decimal("0.01")) and (c is None or abs(abs(c) - gi) < Decimal("0.01"))
                         for go, gi in candidates)
                if d is None and c is None:
                    pass
                elif ok:
                    confirmed += 1
                else:
                    bad.append(f"R{v.r}: итог {d}/{c}, операций {group_out}/{group_in}")
                group_out = group_in = Decimal(0)
                chains[kind] = [Decimal(0), Decimal(0)]
        if confirmed or bad:
            add(f"промежуточные итоги ({grid.name})", "pass" if not bad else "fail",
                detail=f"сошлось {confirmed}" + (f"; не сошлось: {'; '.join(bad[:5])}" if bad else ""))

    n = len(ops)
    no_date = sum(1 for op in ops if op["date"] is None)
    no_amount = sum(1 for op in ops if op["amount"] is None)
    no_dir = sum(1 for op in ops if op["_direction"] is None and op["amount"] is not None)
    bad_inn = sum(1 for op in ops for k in ("debit_inn", "credit_inn") if op[k] and not inn_ok(op[k]))
    both_owner = sum(1 for op in ops if owner.get("inn") and op["debit_inn"] == owner["inn"] == op["credit_inn"])
    add("операции найдены", "pass" if n else "fail", detail=f"{n} операций")
    add("дата у каждой операции", "pass" if not no_date else "fail", detail=f"без даты: {no_date}")
    add("сумма у каждой операции", "pass" if not no_amount else "fail", detail=f"без суммы: {no_amount}")
    add("направление определено", "pass" if not no_dir else "fail", detail=f"без направления: {no_dir}")
    add("ИНН проходят контрольную сумму", "pass" if not bad_inn else "fail", detail=f"неверных: {bad_inn}", level="warning")
    def _legit_no_inn(name: str | None) -> bool:
        if not name:
            return False
        t = name.strip()
        if re.search(r"[A-Za-z]{3,}", t) and not re.search(r"[А-Яа-яЁё]{3,}", t):
            return True
        words = t.replace("ё", "е").replace("Ё", "Е").split()
        return len(words) == 3 and all(re.fullmatch(r"[А-Яа-я]+(?:-[А-Яа-я]+)?", w) for w in words) and not OPF_RE.search(t)

    no_cp_inn_all = sum(1 for op in ops if not op["credit_inn" if op["_direction"] == "out" else "debit_inn"]
                        and op["_direction"] in ("out", "in"))
    no_cp_inn = sum(1 for op in ops if not op["credit_inn" if op["_direction"] == "out" else "debit_inn"]
                    and op["_direction"] in ("out", "in")
                    and not _legit_no_inn(op["credit_name" if op["_direction"] == "out" else "debit_name"]))
    n_dir = sum(1 for op in ops if op["_direction"] in ("out", "in"))
    has_inn_col = any(any(side.get("inn") is not None or side.get("composite") is not None for side in (sp.payer, sp.payee, sp.counterparty))
                      for sp in specs.values()) or not specs
    add("ИНН контрагента указан", "pass" if not no_cp_inn_all else "fail",
        detail=f"без ИНН: {no_cp_inn_all} из {n_dir}, из них не иностранные и не физлица: {no_cp_inn} — проверить вручную",
        level="error" if no_cp_inn and no_cp_inn / max(n_dir, 1) >= 0.2 and has_inn_col else "warning" if no_cp_inn_all else "info")
    add("владелец не на обеих сторонах", "pass" if not both_owner else "fail", detail=f"строк: {both_owner}", level="warning")
    owner_bad = owner_problems(owner)
    add("владелец определён", "pass" if (owner.get("inn") or owner.get("name")) and not owner_bad else "fail",
        detail="; ".join(owner_bad) or f"ИНН={owner.get('inn')}, наименование={owner.get('name')!r}")
    for conflict in owner.get("conflicts") or []:
        add("владелец: шапка и таблица согласованы", "fail", detail=conflict)
    if owner.get("name") and not owner.get("inn") and OPF_RE.search(owner["name"]):
        add("ИНН владельца-организации", "fail", detail=f"у «{owner['name']}» не найден ИНН")
    if "period" in refs and n:
        start, end = refs["period"]
        outside = sum(1 for op in ops if op["date"] and not (start <= op["date"] <= end))
        add("даты в периоде выписки", "pass" if not outside else "fail", detail=f"вне периода: {outside}", level="warning")
    uncertain = sum(1 for vs in verdicts_by_region.values() for v in vs if v.label == UNCERTAIN)
    if uncertain:
        add("неразмеченные строки", "fail", detail=f"{uncertain} строк исключены как спорные", level="warning")

    for code, message in semantic_issues(ops, owner):
        add(f"смысл полей: {code}", "fail", detail=message)
    errors = [c for c in checks if c["status"] == "fail" and c["level"] == "error"]
    arithmetic_pass = [c for c in checks if c["status"] == "pass" and (
        c["check"].startswith(("обороты", "остатки", "промежуточные", "текущий остаток")))]
    status = "needs_review" if errors else ("ok" if arithmetic_pass else "ok_unverified")
    return {"status": status, "checks": checks, "total_out": total_out, "total_in": total_in, "owner": owner,
            "errors": [f"{c['check']}: {c['detail']}" for c in errors]}


# GigaChat для таблицы
def llm_table_spec(region: Region, algo: TableSpec, llm: Any | None = None,
                   feedback: list[str] | None = None) -> TableSpec | None:
    cards = "\n".join(region.profiles[c].card() for c in region.columns if region.profiles[c].n > 0)
    grid = region.grid
    sample_rows = region.anchors[:3] + region.anchors[len(region.anchors) // 2: len(region.anchors) // 2 + 2] + region.anchors[-2:]
    rows_text = []
    for r in dict.fromkeys(sample_rows):
        rows_text.append(f"R{r}: " + " | ".join(f"C{c.c}={c.text[:40]}" for c in grid.row(r)))
    user = (
        f"КАРТОЧКИ КОЛОНОК:\n{cards}\n\nПРИМЕРЫ СТРОК:\n" + "\n".join(rows_text)
        + f"\n\nЧЕРНОВИК АЛГОРИТМА (может быть неверным): {json.dumps(_spec_brief(algo), ensure_ascii=False)}"
        + f"\nСомнения алгоритма: {'; '.join(algo.notes) or 'нет'}"
        + ("\n\nПРОБЛЕМЫ РЕЗУЛЬТАТА ПО ЧЕРНОВИКУ (исправь их):\n- " + "\n- ".join(feedback) if feedback else "")
    )

    def validate(data):
        if not isinstance(data, dict):
            return None, "ожидался JSON-объект"
        try:
            spec = spec_from_json(data)
        except (TypeError, ValueError) as exc:
            return None, f"неверный формат: {exc}"
        errors = check_spec(spec, region)
        if spec.amount is not None and spec.direction is not None:
            values = {e.strip().casefold().rstrip(".") for e in region.profiles[spec.direction].examples}
            listed = {x.casefold() for x in spec.direction_out + spec.direction_in}
            if not values <= listed:
                errors.append(f"значения направления {sorted(values - listed)} не отнесены ни к direction_out, ни к direction_in")
        return (spec, None) if not errors else (None, "; ".join(errors))

    spec = call_llm_json("table_spec", TABLE_SYSTEM, user, validate, llm=llm)
    if spec is not None:
        spec.source = "gigachat"
        spec.notes = [f"схема от GigaChat (уверенность {spec.confidence})"]
    return spec


def _spec_brief(spec: TableSpec) -> dict[str, Any]:
    return {"date": spec.date, "debit": spec.debit, "credit": spec.credit, "signed": spec.signed,
            "amount": spec.amount, "direction": spec.direction, "payer": spec.payer, "payee": spec.payee,
            "counterparty": spec.counterparty, "purpose": spec.purpose[0] if spec.purpose else None,
            "balance": spec.balance}


def spec_from_json(data: dict[str, Any]) -> TableSpec:
    def col(v):
        if v is None or v == "":
            return None
        if isinstance(v, str):
            v = v.strip().upper().lstrip("C")
        return int(v)

    def side(v):
        v = v or {}
        return {"inn": col(v.get("inn")), "name": col(v.get("name")), "composite": col(v.get("composite"))}

    purpose = data.get("purpose")
    purpose_cols = [col(p) for p in purpose] if isinstance(purpose, list) else ([col(purpose)] if purpose is not None else [])
    return TableSpec(
        date=col(data.get("date")), debit=col(data.get("debit")), credit=col(data.get("credit")),
        signed=col(data.get("signed")), amount=col(data.get("amount")), direction=col(data.get("direction")),
        direction_out=[str(x) for x in data.get("direction_out") or []],
        direction_in=[str(x) for x in data.get("direction_in") or []],
        payer=side(data.get("payer")), payee=side(data.get("payee")), counterparty=side(data.get("counterparty")),
        purpose=[p for p in purpose_cols if p is not None],
        balance=[col(b) for b in data.get("balance") or [] if col(b) is not None],
        confidence=float(data.get("confidence", 0.8) or 0.8),
    )


def _row_line(grid: Grid, r: int) -> str:
    return f"R{r}: " + " | ".join(f"C{c.c}={c.text[:50]}" for c in grid.row(r))


def llm_label_rows(region: Region, verdicts: list[RowVerdict], llm: Any | None = None) -> int:
    uncertain = [v for v in verdicts if v.label == UNCERTAIN][:ROW_UNCERTAIN_LIMIT]
    if not uncertain:
        return 0
    grid = region.grid
    by_r = {v.r: v for v in verdicts}
    context_rows = sorted({rr for v in uncertain for rr in range(v.r - 2, v.r + 3) if rr in by_r})
    header = " / ".join(grid.row_text(r) for r in region.header_rows) or "нет"
    user = (
        f"ШАПКА: {header}\n\nСТРОКИ (текущая метка алгоритма в скобках):\n"
        + "\n".join(f"{_row_line(grid, r)}   ({by_r[r].label})" for r in context_rows)
        + "\n\nРАЗМЕТИТЬ: " + ", ".join(f"R{v.r}" for v in uncertain)
    )
    allowed = {OPERATION, CONTINUATION, SUBTOTAL, TOTAL, BALANCE_IN, BALANCE_OUT, HEADER_REPEAT, SERVICE, NOTE}
    targets = {f"R{v.r}" for v in uncertain}

    def validate(data):
        if not isinstance(data, dict):
            return None, "ожидался JSON-объект {\"R57\": \"МЕТКА\"}"
        bad_rows = [k for k in data if k not in targets]
        bad_labels = [f"{k}={v}" for k, v in data.items() if v not in allowed]
        errors = []
        if bad_rows:
            errors.append(f"строк {bad_rows} нет в списке РАЗМЕТИТЬ")
        if bad_labels:
            errors.append(f"недопустимые метки: {bad_labels}")
        return (data, None) if not errors else (None, "; ".join(errors))

    answer = call_llm_json("row_labels", ROWS_SYSTEM, user, validate, llm=llm)
    if not answer:
        return 0
    changed = 0
    for key, label in answer.items():
        v = by_r[int(key[1:])]
        v.label, v.decided_by, v.confidence = label, "gigachat", 0.7
        v.reason = f"GigaChat: {label}"
        changed += 1
    return changed


def _subset_with_sum(items: list[tuple[int, Decimal]], target: Decimal, max_size: int = 3) -> list[int] | None:
    items = [(r, v) for r, v in items if v is not None and v != 0][:40]
    for size in range(1, max_size + 1):
        for combo in combinations(items, size):
            if abs(sum(v for _, v in combo) - target) < Decimal("0.01"):
                return [r for r, _ in combo]
    return None


def arithmetic_repair(region: Region, spec: TableSpec, verdicts: list[RowVerdict], diff: dict[str, Decimal]) -> list[str]:
    notes = []
    grid = region.grid
    for key, d in diff.items():
        col = getattr(spec, key, None)
        if col is None or abs(d) < Decimal("0.01"):
            continue
        amount_of = lambda r: (lambda cell: parse_amount(cell.text) if cell is not None and cell.r == r else None)(grid.at(r, col))
        if d > 0:
            pool = [(v.r, abs(amount_of(v.r) or 0)) for v in verdicts
                    if v.label in (UNCERTAIN, NOTE, SUBTOTAL) and not v.reason.startswith("ключевое слово итога")]
            rows = _subset_with_sum(pool, d)
            if rows:
                for v in verdicts:
                    if v.r in rows:
                        v.label, v.decided_by, v.reason = OPERATION, "repair", f"сверка оборотов ({key}): не хватало {d}"
                notes.append(f"{key}: строки {rows} возвращены в операции (не хватало {d})")
        else:
            pool = [(v.r, abs(amount_of(v.r) or 0)) for v in verdicts if v.label == OPERATION and v.confidence < 0.9]
            rows = _subset_with_sum(pool, -d) or _subset_with_sum(
                [(v.r, abs(amount_of(v.r) or 0)) for v in verdicts if v.label == OPERATION], -d, max_size=2)
            if rows:
                for v in verdicts:
                    if v.r in rows:
                        v.label, v.decided_by, v.reason = SUBTOTAL, "repair", f"сверка оборотов ({key}): лишние {-d}"
                notes.append(f"{key}: строки {rows} исключены из операций (лишние {-d})")
    return notes


def llm_repair(region: Region, spec: TableSpec, verdicts: list[RowVerdict], diff: dict[str, Decimal],
               llm: Any | None = None) -> dict[str, str] | None:
    grid = region.grid
    amount_cols = [c for c in (spec.debit, spec.credit, spec.signed, spec.amount) if c is not None]
    candidates = [v for v in verdicts if v.label not in (EMPTY, HEADER_REPEAT, COLUMN_NUMBERS)
                  and any((cell := grid.at(v.r, c)) is not None and cell.r == v.r and "amount" in region.kinds(cell) for c in amount_cols)]
    candidates = candidates[:50]
    if not candidates:
        return None
    user = (
        "РАСХОЖДЕНИЕ (итог документа − сумма операций): "
        + ", ".join(f"{k}: {v}" for k, v in diff.items() if abs(v) >= Decimal("0.01"))
        + f"\nКолонки сумм: дебет=C{spec.debit}, кредит=C{spec.credit}, со знаком=C{spec.signed}"
        + "\n\nКАНДИДАТЫ (текущая метка в скобках):\n"
        + "\n".join(f"{_row_line(grid, v.r)}   ({v.label}: {v.reason})" for v in candidates)
    )
    allowed = {OPERATION, SUBTOTAL, TOTAL, NOTE, BALANCE_IN, BALANCE_OUT}
    targets = {f"R{v.r}" for v in candidates}

    def validate(data):
        if not isinstance(data, dict) or not data:
            return None, "ожидался непустой JSON-объект {\"R57\": \"МЕТКА\"}"
        errs = [k for k in data if k not in targets] + [f"{k}={v}" for k, v in data.items() if v not in allowed]
        return (data, None) if not errs else (None, f"недопустимые строки/метки: {errs}")

    return call_llm_json("repair", REPAIR_SYSTEM, user, validate, llm=llm)


def spec_quality(doc: Document, region: Region, spec: TableSpec) -> tuple[float, dict[str, Any]]:
    verdicts = classify_rows(region, spec)
    raw = records_to_raw_ops(build_records(region, spec, verdicts), spec)
    n = len(raw)
    if not n:
        return -10.0, {"operations": 0}
    owner = detect_owner(doc, {region.grid.name: [region]}, raw, allow_llm=False)
    ops, _ = finalize_ops(copy.deepcopy(raw), owner)
    refs = find_references(doc, [region], {0: verdicts}, {0: spec})
    diffs = _diffs(refs, ops, 1)
    passes = sum(1 for v in diffs.values() if abs(v) < Decimal("0.01"))
    fails = len(diffs) - passes
    amount_rate = sum(1 for o in ops if o["amount"] is not None) / n
    dir_rate = sum(1 for o in ops if o["_direction"] and o["amount"] is not None) / n
    date_rate = sum(1 for o in ops if o["date"] is not None) / n
    both_sides = sum(1 for o in ops if (o["debit_inn"] or o["debit_name"]) and (o["credit_inn"] or o["credit_name"])) / n
    inns = [o[k] for o in ops for k in ("debit_inn", "credit_inn") if o[k]]
    inn_valid = sum(1 for x in inns if inn_ok(x)) / len(inns) if inns else 0.0
    owner_as_cp = sum(1 for o in raw if owner.get("inn") and o.get("counterparty_inn") == owner["inn"]) / n
    issues = semantic_issues(ops, owner)
    score = (3 * passes - 3 * fails + 2 * amount_rate + dir_rate + date_rate + 0.5 * both_sides
             + 0.5 * inn_valid - 2 * owner_as_cp - 1.5 * len(issues))
    return round(score, 3), {"operations": n, "turnover_pass": passes, "turnover_fail": fails,
                             "amount": round(amount_rate, 2), "direction": round(dir_rate, 2), "both_sides": round(both_sides, 2),
                             "inn_valid": round(inn_valid, 2), "owner_as_counterparty": round(owner_as_cp, 2),
                             "semantic_issues": [code for code, _ in issues]}


# документы без таблицы
def is_1c_format(lines: list[str]) -> bool:
    head = "\n".join(lines[:5])
    return "1CClientBankExchange" in head or sum(1 for line in lines if line.startswith("СекцияДокумент")) >= 1


def parse_1c(lines: list[str]) -> tuple[list[dict[str, Any]], dict[str, Any], dict[str, Any]]:
    header: dict[str, str] = {}
    accounts: list[dict[str, str]] = []
    docs: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    section = None
    for raw in lines:
        line = raw.strip()
        if line.startswith("СекцияДокумент"):
            current, section = {"_type": line.split("=", 1)[-1]}, "doc"
            continue
        if line == "КонецДокумента":
            if current is not None:
                docs.append(current)
            current, section = None, None
            continue
        if line == "СекцияРасчСчет":
            current, section = {}, "acc"
            continue
        if line == "КонецРасчСчет":
            accounts.append(current or {})
            current, section = None, None
            continue
        if "=" in line:
            key, value = line.split("=", 1)
            if current is not None:
                current[key.strip()] = value.strip()
            else:
                header[key.strip()] = value.strip()
    own_accounts = {header.get("РасчСчет")} | {a.get("РасчСчет") for a in accounts}
    own_accounts.discard(None)

    owner_inn = owner_name = None
    ops = []
    for d in docs:
        payer_acc = d.get("ПлательщикСчет") or d.get("ПлательщикРасчСчет")
        payee_acc = d.get("ПолучательСчет") or d.get("ПолучательРасчСчет")
        payer_name = d.get("Плательщик1") or d.get("Плательщик")
        payee_name = d.get("Получатель1") or d.get("Получатель")
        if payer_acc in own_accounts:
            direction = "out"
            owner_inn = owner_inn or d.get("ПлательщикИНН")
            owner_name = owner_name or payer_name
        elif payee_acc in own_accounts:
            direction = "in"
            owner_inn = owner_inn or d.get("ПолучательИНН")
            owner_name = owner_name or payee_name
        else:
            direction = None
        op_date = parse_date(d.get("ДатаСписано") or d.get("ДатаПоступило") or d.get("Дата"))
        ops.append({
            "date": op_date, "amount": parse_amount(d.get("Сумма")), "_direction": direction,
            "debit_inn": d.get("ПлательщикИНН"), "debit_name": clean_party_name(payer_name or "") or None,
            "credit_inn": d.get("ПолучательИНН"), "credit_name": clean_party_name(payee_name or "") or None,
            "purpose": d.get("НазначениеПлатежа"), "_rows": [], "_grid": "1c", "_date_inherited": False,
        })
    owner_name = header.get("ВладелецСчет") or owner_name
    owner = {"inn": owner_inn, "name": clean_party_name(owner_name) if owner_name else None,
             "inn_confidence": 0.95 if owner_inn else 0.0, "name_confidence": 0.95 if owner_name else 0.0,
             "confidence": 0.95 if owner_inn else 0.5, "method": "1c_account_match",
             "evidence": {"inn": f"счёт {', '.join(sorted(own_accounts))}", "name": None}}
    refs: dict[str, Any] = {"_sources": {}}
    acc = accounts[0] if len(accounts) == 1 else {}
    for key, field_ in (("debit#0", "ВсегоСписано"), ("credit#0", "ВсегоПоступило"),
                        ("balance_in#0", "НачальныйОстаток"), ("balance_out#0", "КонечныйОстаток")):
        value = acc.get(field_) or (header.get(field_) if len(accounts) <= 1 else None)
        if value is not None:
            refs[key] = parse_amount(value)
            refs["_sources"][key] = f"1C: {field_}={value}"
    return ops, owner, refs


def _grounded(value: str | None, lines_text: str, kind: str) -> bool:
    if not value:
        return False
    hay = lines_text.casefold().replace("ё", "е")
    if kind == "inn":
        return re.sub(r"\D", "", value) in re.sub(r"(?<=\d)[  ](?=\d)", "", hay)
    if kind == "amount":
        target = parse_amount(value)
        return target is not None and any(
            parse_amount(m) is not None and abs(abs(parse_amount(m)) - abs(target)) < Decimal("0.01")
            for m in re.findall(r"-?\d[\d  ]*(?:[.,-]\d{1,2})?", lines_text)
        )
    if kind == "date":
        d = parse_date(value)
        return d is not None and any(parse_date(m) == d for m in re.findall(
            r"\d{1,2}[./-]\d{1,2}[./-]\d{2,4}|\d{4}-\d{2}-\d{2}|\d{1,2}\s+[а-яa-z]{3,}\.?\s+\d{4}", lines_text, re.I))
    needle = value.casefold().replace("ё", "е")
    if needle in hay:
        return True
    tokens = [t for t in re.findall(r"[0-9a-zа-я]{3,}", needle)]
    return bool(tokens) and sum(t in hay for t in tokens) / len(tokens) >= 0.8


_GROUND_FIELDS = (("date", "date"), ("amount", "amount"), ("counterparty_inn", "inn"),
                  ("counterparty_name", "text"), ("purpose", "text"))


def _item_lines(item: dict[str, Any]) -> list[int]:
    return [int(str(x).lstrip("L")) for x in item.get("lines") or [] if str(x).lstrip("L").isdigit()]


def llm_text_extract(lines: list[str], llm: Any | None = None) -> tuple[list[dict[str, Any]], list[str]]:
    ops: list[dict[str, Any]] = []
    notes: list[str] = []
    step = TEXT_CHUNK_LINES - 5
    seen: set[tuple] = set()
    n_chunks = len(range(0, max(len(lines), 1), step))
    for k, start in enumerate(range(0, max(len(lines), 1), step), 1):
        if CLOCK.unavailable or CLOCK.left("algo") <= 1:
            notes.append(f"извлечение из текста остановлено на фрагменте {k}/{n_chunks}: "
                         f"{CLOCK.unavailable or 'бюджет времени исчерпан'}")
            log("текст", notes[-1])
            break
        chunk = [(i, lines[i]) for i in range(start, min(start + TEXT_CHUNK_LINES, len(lines))) if lines[i].strip()]
        if not chunk:
            continue
        if not any(parse_date(m) for _, text in chunk for m in re.findall(r"\d{1,2}[./]\d{1,2}[./]\d{2,4}", text)):
            continue
        log("текст", f"фрагмент {k}/{n_chunks} (строки {start}–{start + TEXT_CHUNK_LINES}), найдено операций: {len(ops)}")
        valid_lines = {i for i, _ in chunk}
        user = "ФРАГМЕНТ:\n" + "\n".join(f"L{i}: {text}" for i, text in chunk)
        last: dict[str, Any] = {"data": None}

        def validate(data):
            if not isinstance(data, list):
                return None, "ожидался JSON-массив операций"
            last["data"] = data
            errors = []
            for k, item in enumerate(data):
                if not isinstance(item, dict):
                    errors.append(f"элемент {k} — не объект")
                    continue
                refs = _item_lines(item)
                if not refs or any(i not in valid_lines for i in refs):
                    errors.append(f"операция {k}: lines={item.get('lines')} вне фрагмента")
                    continue
                if item.get("direction") not in ("in", "out", None):
                    errors.append(f"операция {k}: direction должно быть in или out")
                text = "\n".join(lines[i] for i in refs)
                for key, kind in _GROUND_FIELDS:
                    value = item.get(key)
                    if value is not None and not _grounded(str(value), text, kind):
                        errors.append(f"операция {k}: {key}={value!r} нет в строках {['L%d' % i for i in refs]} — "
                                      f"копируй значение дословно")
            return (data, None) if not errors else (None, "; ".join(errors[:6]))

        answer = call_llm_json("text_extract", TEXT_SYSTEM, user, validate, llm=llm)
        if answer is None and isinstance(last["data"], list):
            answer = last["data"]
            notes.append(f"L{start}–L{start + TEXT_CHUNK_LINES}: ответ GigaChat не прошёл проверку — "
                         f"неподтверждённые поля отброшены")
        for item in answer or []:
            if not isinstance(item, dict):
                continue
            refs = [i for i in _item_lines(item) if 0 <= i < len(lines)]
            text = "\n".join(lines[i] for i in refs)
            fields = {}
            for key, kind in _GROUND_FIELDS:
                value = item.get(key)
                fields[key] = str(value) if value is not None and _grounded(str(value), text, kind) else None
            op_date, amount = parse_date(fields["date"]), parse_amount(fields["amount"])
            if op_date is None or amount is None:
                notes.append(f"L{refs}: операция отброшена — нет подтверждённой даты или суммы")
                continue
            key = (tuple(refs), op_date, abs(amount))
            if key in seen:
                continue
            seen.add(key)
            inn = re.sub(r"\D", "", fields["counterparty_inn"] or "") or None
            ops.append({"date": op_date, "amount": abs(amount), "direction": item.get("direction"),
                        "counterparty_inn": inn, "counterparty_name": fields["counterparty_name"],
                        "payer_inn": None, "payer_name": None, "payee_inn": None, "payee_name": None,
                        "purpose": fields["purpose"], "_rows": refs, "_grid": "text",
                        "_date_inherited": False, "_provenance": {}})
    return ops, notes


KV_LABELS = [
    ("date", re.compile(r"^\s*дата\b", re.I)),
    ("direction", re.compile(r"^\s*(?:(?:вид|тип)\s+операц\w*|направлен\w*|дебет/кредит|дт/кт|тип|вид)\s*:?\s*$", re.I)),
    ("amount", re.compile(r"^\s*сумма|^\s*amount", re.I)),
    ("purpose", re.compile(r"^\s*(?:назначени\w+|примечани\w+|детал\w*|основани\w*|комментари\w*)", re.I)),
    ("counterparty_inn", re.compile(r"^\s*инн\b", re.I)),
    ("counterparty_name", re.compile(r"^\s*(?:контрагент\w*|корреспондент\w*|плательщик\w*|получател\w*)", re.I)),
    ("skip", re.compile(r"^\s*(?:бик|кпп|огрн|сч[её]т\w*|р/с|к/с|банк\w*|номер\w*|№|код\w*|название\s+банка)", re.I)),
]


def _kv_label(text: str) -> str | None:
    t = text.strip()
    if len(t) > 60:
        return None
    for kind, pattern in KV_LABELS:
        if pattern.match(t):
            return kind
    return None


def kv_blocks_parse(doc: Document) -> list[dict[str, Any]] | None:
    for grid in doc.grids:
        if grid.n_cols < 2 or grid.is_text_only():
            continue
        label_col, best = None, 0
        for c in range(grid.n_cols - 1):
            hits = sum(1 for cell in grid.cells
                       if cell.c == c and cell.r2 == cell.r and _kv_label(cell.text))
            if hits > best:
                label_col, best = c, hits
        if label_col is None or best < 6:
            continue
        value_cols = [c for c in range(grid.n_cols) if c != label_col]
        blocks: list[dict[str, Any]] = []
        current: dict[str, Any] | None = None

        def value_of(r: int) -> str | None:
            for c in value_cols:
                cell = grid.at(r, c)
                if cell is not None and cell.r == r and cell.text.strip():
                    return cell.text.strip()
            return None

        for r in range(grid.n_rows):
            label_cell = grid.at(r, label_col)
            label = _kv_label(label_cell.text) if label_cell is not None and label_cell.r == r else None
            value = value_of(r)
            if label == "date" and value and parse_date(value) is not None:
                current = {"date": parse_date(value), "_rows": [r]}
                blocks.append(current)
                continue
            if current is None:
                continue
            if label and label != "skip" and value:
                key = label
                if key in current and current[key] and value not in str(current[key]):
                    current[key] = f"{current[key]} {value}"
                else:
                    current[key] = value
                current["_rows"].append(r)
            elif label_cell is None and value is None:
                current = None
        ops = []
        for b in blocks:
            amount = parse_amount(str(b.get("amount", "")))
            if b["date"] is None or amount is None:
                continue
            raw_dir = str(b.get("direction", "")).strip().casefold().rstrip(".")
            direction = "out" if raw_dir in DIRECTION_OUT else "in" if raw_dir in DIRECTION_IN else None
            inn = re.sub(r"\D", "", str(b.get("counterparty_inn", ""))) or None
            if inn and len(inn) not in (10, 12):
                inn = None
            ops.append({"date": b["date"], "amount": abs(amount), "direction": direction,
                        "counterparty_inn": inn, "counterparty_name": b.get("counterparty_name"),
                        "payer_inn": None, "payer_name": None, "payee_inn": None, "payee_name": None,
                        "purpose": b.get("purpose"), "_rows": b["_rows"], "_grid": grid.name,
                        "_date_inherited": False, "_provenance": {}})
        if len(ops) >= 3:
            log("текст", f"{grid.name}: блоки «ключ-значение» — {len(ops)} операций")
            return ops
    return None


# главная функция
import copy


def _diffs(refs: dict[str, Any], ops: list[dict[str, Any]], n_regions: int, regions: list | None = None) -> dict[str, Decimal]:
    out = {}
    names = None
    if regions:
        with_refs = [i for i in range(len(regions)) if f"debit#{i}" in refs or f"credit#{i}" in refs]
        if with_refs and len(with_refs) < len(regions):
            names = {regions[i].grid.name for i in with_refs}
    for key, direction in (("debit", "out"), ("credit", "in")):
        vals = [v for k, v in refs.items() if k.startswith(key + "#") and not k.endswith("#text")]
        expected = sum(vals, Decimal(0)) if vals else refs.get(f"{key}#text")
        if expected is None:
            continue
        actual = sum((op["amount"] for op in ops if op["_direction"] == direction and op["amount"] is not None
                      and (names is None or op["_grid"] in names)), Decimal(0))
        out[key] = expected - actual
    return out


def _error_size(diff: dict[str, Decimal]) -> Decimal:
    return sum((abs(v) for v in diff.values()), Decimal(0))


ISSUE_HINTS = {
    "purpose_codes": "в назначении получились коды (01, 17…) — назначение должно быть длинным текстом",
    "bank_as_counterparty": "в контрагентах получились банки — возьми колонки самого плательщика/получателя, а не его банка",
    "owner_caption": "владелец — подпись формы, а не имя",
    "purpose_in_party": "в наименованиях сторон оказалось назначение платежа — стороны бери только из колонок плательщика/получателя, текст назначения — в purpose",
    "same_party_both_sides": "одна и та же сторона стоит и в дебете, и в кредите — у плательщика и получателя должны быть РАЗНЫЕ колонки (или одна колонка контрагента)",
    "owner_on_wrong_side": "владелец стоит не на своей стороне: при списании он в дебете, при зачислении — в кредите; проверь колонки дебет/кредит (flip_direction) и стороны",
    "owner_as_counterparty": "контрагентом оказывается сам владелец — для контрагента нужны колонки «плательщика/получателя» (противоположной стороны), а не владельца",
    "party_not_text": "в наименовании контрагента цифры (счёт, БИК, номер) — выбери колонку с текстовым наименованием",
    "party_inn_invalid": "в ИНН контрагента не ИНН (контрольная сумма не сходится) — выбери колонку ИНН (10 или 12 цифр), не КПП и не счёт",
    "subtotals_in_ops": "в операции попали итоговые строки",
}


def parse_statement(path: str | Path, llm: Any | None = None, use_llm: bool = True) -> tuple[pd.DataFrame, dict[str, Any]]:
    llm_eff = (llm if llm is not None else LLM_CLIENT) if use_llm else None
    table, report = _parse_statement_once(path, llm, use_llm)
    if llm_eff is None:
        return table, report
    calls = list(LLM_CALLS)
    base_table, base_report = _parse_statement_once(path, None, False)
    LLM_CALLS[:] = calls
    if base_report.get("score", -99.0) > report.get("score", -99.0) + 1e-9:
        base_report["llm_calls"] = calls
        base_report["gigachat"] = True
        base_report["notes"].append(f"GigaChat ухудшил результат (оценка {report.get('score')} против {base_report.get('score')} у алгоритма) — "
                                    f"возвращён результат алгоритма")
        log("итог", "результат с GigaChat хуже алгоритма — возвращён результат алгоритма")
        return base_table, base_report
    return table, report


def _parse_statement_once(path: str | Path, llm: Any | None = None, use_llm: bool = True) -> tuple[pd.DataFrame, dict[str, Any]]:
    llm = (llm if llm is not None else LLM_CLIENT) if use_llm else None
    LLM_CALLS.clear()
    CLOCK.reset()
    CLOCK.mark("этап 1/4 чтение", f"{Path(path).name}; GigaChat {'подключён' if llm is not None else 'выключен'}")
    doc = read_document(path)
    report: dict[str, Any] = {"file": str(path), "notes": [], "gigachat": llm is not None}

    regions_by_grid: dict[str, list[Region]] = {}
    regions: list[Region] = []
    extra_grids: list[Grid] = []
    for grid in doc.grids:
        found = find_regions(grid) if not grid.is_text_only() else []
        if not found and not grid.is_text_only():
            label_col = looks_transposed(grid)
            if label_col is not None:
                flipped = transposed_grid(grid, label_col)
                if flipped is not None:
                    found = find_regions(flipped)
                    if found:
                        extra_grids.append(flipped)
                        report["notes"].append(f"{grid.name}: таблица транспонирована (поля в строках)")
                        log("сетка", f"{grid.name}: транспонированная таблица — операции в колонках")
        if found:
            regions_by_grid[found[0].grid.name] = found
            regions.extend(found)
    if extra_grids:
        doc.grids.extend(extra_grids)

    specs: dict[int, TableSpec] = {}
    verdicts: dict[int, list[RowVerdict]] = {}

    CLOCK.mark("этап 2/4 разбор", f"таблиц операций найдено: {len(regions)}" if regions else "таблица операций не найдена — текстовый режим")
    if not regions:
        kv_ops = kv_blocks_parse(doc)
        if kv_ops:
            report["mode"] = "kv"
            raw_ops = kv_ops
            owner = detect_owner(doc, {}, raw_ops, llm)
            ops, notes = finalize_ops(raw_ops, owner)
            report["notes"] += notes
            refs = find_references(doc, [], {}, {})
            result = run_checks(ops, refs, {}, [], {}, owner)
        else:
            lines = [grid.row_text(r, sep="  ") for grid in doc.grids for r in range(grid.n_rows)]
            lines = [line for line in lines if line.strip()]
            if is_1c_format(lines):
                report["mode"] = "text_1c"
                ops, owner, refs = parse_1c(lines)
            else:
                report["mode"] = "text_llm"
                if llm is None:
                    raw_ops, notes = [], ["таблица не найдена, а GigaChat выключен — извлекать нечем"]
                else:
                    raw_ops, notes = llm_text_extract(lines, llm)
                report["notes"] += notes
                owner = detect_owner(doc, {}, raw_ops, llm)
                ops, notes = finalize_ops(raw_ops, owner)
                report["notes"] += notes
                refs = find_references(doc, [], {}, {})
            result = run_checks(ops, refs, {}, [], {}, owner)
    else:
        report["mode"] = "table"
        for i, region in enumerate(regions):
            spec = infer_table_spec(region)
            errors = check_spec(spec, region)
            log("схема", f"{region.grid.name}: строки {region.first_anchor}–{region.end_row}, "
                         f"шапка {region.header_rows}, уверенность {spec.confidence}"
                         + (f", ошибки: {errors}" if errors else ""))
            q_algo, d_algo = spec_quality(doc, region, spec) if llm is not None else (None, {})
            feedback = [f"{code}" for code in d_algo.get("semantic_issues", [])]
            if d_algo.get("turnover_fail"):
                feedback.append("сумма операций не сходится с итогами документа")
            if (spec.confidence < TABLE_LLM_THRESHOLD or errors or feedback) and llm is not None:
                llm_spec = llm_table_spec(region, spec, llm, feedback=[ISSUE_HINTS.get(f, f) for f in feedback])
                if llm_spec is not None:
                    q_llm, d_llm = spec_quality(doc, region, llm_spec)
                    better = q_llm > q_algo
                    verdict = "принята" if better else "отклонена"
                    report["notes"].append(f"{region.grid.name}: схема GigaChat {verdict} "
                                           f"(качество {q_llm} против {q_algo} у алгоритма)")
                    report.setdefault("spec_arbitration", []).append(
                        {"grid": region.grid.name, "algorithm": d_algo, "gigachat": d_llm,
                         "score_algorithm": q_algo, "score_gigachat": q_llm, "accepted": bool(better)})
                    log("схема", f"GigaChat {verdict}: {q_llm} vs {q_algo}")
                    if better:
                        spec = llm_spec
            specs[i] = spec
            verdicts[i] = classify_rows(region, spec)
            if llm is not None:
                n = llm_label_rows(region, verdicts[i], llm)
                if n:
                    report["notes"].append(f"{region.grid.name}: GigaChat разметил {n} спорных строк")

        state = {"flip": False}

        def build_raw() -> list[dict[str, Any]]:
            raw = []
            for i, region in enumerate(regions):
                raw += records_to_raw_ops(build_records(region, specs[i], verdicts[i]), specs[i])
            if state["flip"]:
                for op in raw:
                    if op["direction"]:
                        op["direction"] = "in" if op["direction"] == "out" else "out"
            return raw

        raw_ops = build_raw()
        CLOCK.mark("этап 3/4 владелец и сверка", f"операций собрано: {len(raw_ops)}")
        owner = detect_owner(doc, regions_by_grid, raw_ops, llm)
        ops, notes = finalize_ops(raw_ops, owner)
        report["notes"] += notes
        refs = find_references(doc, regions, verdicts, specs)
        result = run_checks(ops, refs, verdicts, regions, specs, owner)

        def _balance_hits(res):
            chk = next((c for c in res["checks"] if c["check"].startswith("текущий остаток")), None)
            m = re.search(r"сошлось (\d+) из (\d+)", chk["detail"]) if chk else None
            return (int(m.group(1)), int(m.group(2))) if m else None

        bh = _balance_hits(result)
        two_sided_ops = any(op.get("payer_inn") or op.get("payer_name") for op in raw_ops) and \
            any(op.get("payee_inn") or op.get("payee_name") for op in raw_ops)
        if bh and bh[1] >= 3 and bh[0] <= 0.2 * bh[1] and not two_sided_ops:
            state["flip"] = True
            flipped_raw = build_raw()
            flipped_ops, _ = finalize_ops(flipped_raw, owner)
            flipped_result = run_checks(flipped_ops, refs, verdicts, regions, specs, owner, flipped=True)
            fh = _balance_hits(flipped_result)
            if fh and fh[0] >= 0.8 * fh[1]:
                raw_ops, ops, result = flipped_raw, flipped_ops, flipped_result
                swapped_refs = {(("credit#" + k[6:]) if k.startswith("debit#") else ("debit#" + k[7:]) if k.startswith("credit#") else k): v
                                for k, v in refs.items()}
                alt = run_checks(ops, swapped_refs, verdicts, regions, specs, owner, flipped=True)
                if len([c for c in alt["checks"] if c["status"] == "fail"]) < len([c for c in result["checks"] if c["status"] == "fail"]):
                    refs, result = swapped_refs, alt
                report["notes"].append(f"дебет и кредит были перепутаны — направление перевёрнуто по колонке остатка ({fh[0]} из {fh[1]} переходов сошлись)")
            else:
                state["flip"] = False

        diff = _diffs(refs, ops, len(regions), regions)
        swapped = diff and all(k in diff for k in ("debit", "credit")) and _error_size(diff) >= Decimal("0.01")
        if swapped:
            ref_out = result["total_out"] + diff["debit"]
            ref_in = result["total_in"] + diff["credit"]
            if abs(ref_out - result["total_in"]) < Decimal("0.01") and abs(ref_in - result["total_out"]) < Decimal("0.01"):
                two_sided = any(op.get("payer_inn") or op.get("payer_name") for op in raw_ops) and \
                    any(op.get("payee_inn") or op.get("payee_name") for op in raw_ops)
                if two_sided:
                    refs = {(("credit#" + k[6:]) if k.startswith("debit#") else ("debit#" + k[7:]) if k.startswith("credit#") else k): v
                            for k, v in refs.items()}
                    result = run_checks(ops, refs, verdicts, regions, specs, owner)
                    report["notes"].append("итоги «по дебету»/«по кредиту» стояли в колонках наоборот — сверка выполнена с обменом; направления операций взяты по сторонам")
                else:
                    state["flip"] = True
                    raw_ops = build_raw()
                    ops, _ = finalize_ops(raw_ops, owner)
                    result = run_checks(ops, refs, verdicts, regions, specs, owner, flipped=state["flip"])
                    report["notes"].append("дебет и кредит были перепутаны — направление перевёрнуто по сверке оборотов")
                diff = _diffs(refs, ops, len(regions), regions)
        for iteration in range(REPAIR_MAX_ITERATIONS + 1):
            if _error_size(diff) < Decimal("0.01") or len(regions) != 1:
                break
            log("исправление", f"итерация {iteration + 1}/{REPAIR_MAX_ITERATIONS + 1}: расхождение с итогами "
                               f"{_error_size(diff)} ({'арифметика' if iteration == 0 else 'GigaChat'})")
            before = copy.deepcopy(verdicts)
            if iteration == 0:
                fix_notes = arithmetic_repair(regions[0], specs[0], verdicts[0], diff)
                source = "арифметика"
            else:
                if llm is None:
                    break
                answer = llm_repair(regions[0], specs[0], verdicts[0], diff, llm)
                fix_notes = []
                if answer:
                    by_r = {v.r: v for v in verdicts[0]}
                    for key, label in answer.items():
                        v = by_r[int(key[1:])]
                        v.label, v.decided_by, v.reason = label, "gigachat-repair", "GigaChat: исправление по сверке"
                    fix_notes = [f"GigaChat переразметил {list(answer)}"]
                source = "GigaChat"
            if not fix_notes:
                continue
            new_raw = build_raw()
            new_ops, notes = finalize_ops(new_raw, owner)
            new_diff = _diffs(refs, new_ops, len(regions), regions)
            if _error_size(new_diff) < _error_size(diff):
                raw_ops, ops, diff = new_raw, new_ops, new_diff
                result = run_checks(ops, refs, verdicts, regions, specs, owner, flipped=state["flip"])
                log("исправление", f"принято, расхождение теперь {_error_size(diff)}")
                report["notes"] += [f"исправление ({source}): {n}" for n in fix_notes]
            else:
                verdicts = before
                report["notes"].append(f"исправление ({source}) отклонено — расхождение не уменьшилось")

    CLOCK.mark("этап 4/4 итог алгоритма", f"операций {len(ops)}, статус проверок: {result['status']}")
    return compose_output(path, doc, report, regions, specs, verdicts, owner, ops, refs, result,
                          raw_ops=locals().get("raw_ops"), regions_by_grid=regions_by_grid)


def ops_to_table(ops: list[dict[str, Any]]) -> pd.DataFrame:
    rows = []
    for op in ops:
        row = {k: op.get(k) for k in OUTPUT_COLUMNS}
        if row["amount"] is not None:
            row["amount"] = -abs(row["amount"]) if op.get("_direction") == "out" else abs(row["amount"])
        rows.append(row)
    table = pd.DataFrame(rows, columns=OUTPUT_COLUMNS)
    table["date"] = pd.to_datetime(table["date"], errors="coerce").astype("datetime64[ns]")
    table["amount"] = pd.to_numeric(table["amount"].map(lambda v: float(v) if v is not None else None), errors="coerce").astype("Float64")
    for col in ("debit_inn", "debit_name", "credit_inn", "credit_name", "purpose"):
        table[col] = table[col].astype("string")
    return table


def compose_output(path, doc: Document, report: dict[str, Any], regions: list[Region], specs: dict[int, TableSpec],
                   verdicts: dict[int, list[RowVerdict]], owner: dict[str, Any], ops: list[dict[str, Any]],
                   refs: dict[str, Any], result: dict[str, Any], raw_ops: list[dict[str, Any]] | None = None,
                   regions_by_grid: dict[str, list[Region]] | None = None) -> tuple[pd.DataFrame, dict[str, Any]]:
    table = ops_to_table(ops)

    llm_failed = sorted({c["task"] for c in LLM_CALLS if any("raw" not in a for a in c["attempts"])})
    if llm_failed:
        result = dict(result)
        result["status"] = "needs_review"
        result["errors"] = list(result["errors"]) + [f"GigaChat не ответил ({', '.join(llm_failed)}): результат алгоритма без проверки моделью"]
        report["notes"].append("GigaChat не ответил — использован результат алгоритма")

    excluded = []
    for i, region in enumerate(regions):
        for v in verdicts[i]:
            if v.label not in (OPERATION, CONTINUATION, EMPTY):
                excluded.append({"sheet": region.grid.name, "row": v.r + 1, "label": v.label, "reason": v.reason,
                                 "decided_by": v.decided_by, "text": region.grid.row_text(v.r)[:200]})
    report.update({
        "status": result["status"],
        "errors": result["errors"],
        "owner": {k: owner.get(k) for k in ("inn", "name", "confidence", "inn_confidence", "name_confidence",
                                             "method", "evidence", "conflicts", "inn_candidates", "name_candidates", "llm")},
        "checks": result["checks"],
        "references": {k: (str(v) if not isinstance(v, (dict, tuple)) else v) for k, v in refs.items()},
        "specs": {regions[i].grid.name + f"#{i}": {**_spec_brief(s), "confidence": s.confidence,
                                                   "source": s.source, "notes": s.notes} for i, s in specs.items()},
        "excluded_rows": excluded,
        "operations_extra": [{"direction": op.get("_direction"), "rows": [r + 1 for r in op.get("_rows", [])],
                              "sheet": op.get("_grid"), "date_inherited": op.get("_date_inherited")} for op in ops],
        "llm_calls": list(LLM_CALLS),
        "timeline": list(CLOCK.timeline),
        "score": result_score(result, ops),
        "_state": {"doc": doc, "regions": regions, "specs": specs, "verdicts": verdicts, "owner": owner,
                   "ops": ops, "refs": refs, "raw_ops": raw_ops, "mode": report.get("mode"),
                   "regions_by_grid": regions_by_grid or {}},
    })
    log("итог", f"{Path(path).name}: {len(table)} операций, статус {report['status']}"
                + (f" — {report['errors']}" if report["errors"] else ""))
    return table, report


def result_score(result: dict[str, Any], ops: list[dict[str, Any]]) -> float:
    checks = result["checks"]
    n = len(ops)
    arithmetic = sum(1 for c in checks if c["status"] == "pass" and c["check"].startswith(
        ("обороты", "остатки", "промежуточные", "текущий остаток")))
    errors = [c for c in checks if c["status"] == "fail" and c["level"] == "error"]
    heavy = sum(1 for c in errors if c["check"].startswith(("обороты", "остатки", "смысл полей", "операции найдены")))
    if not n:
        return round(-10.0 - 3 * len(errors), 3)
    amount_rate = sum(1 for o in ops if o.get("amount") is not None) / n
    dir_rate = sum(1 for o in ops if o.get("_direction") in ("in", "out")) / n
    date_rate = sum(1 for o in ops if o.get("date") is not None) / n
    cp_rate = sum(1 for o in ops if (o.get("debit_name") or o.get("debit_inn")) and (o.get("credit_name") or o.get("credit_inn"))) / n
    owner_penalty = 6 * len(owner_problems(result.get("owner") or {})) if result.get("owner") else 0
    return round(3 * arithmetic - 3 * len(errors) - 2 * heavy - owner_penalty + 2 * amount_rate + dir_rate + date_rate + 0.5 * cp_rate, 3)


def export_result(table: pd.DataFrame, report: dict[str, Any], path: str | Path, output_dir: str | Path | None = None) -> Path:
    path = Path(path)
    out_dir = Path(output_dir) if output_dir else path.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / f"{path.stem}__normalized.xlsx"
    with pd.ExcelWriter(target) as writer:
        table.to_excel(writer, sheet_name="operations", index=False)
        pd.DataFrame(report["checks"]).to_excel(writer, sheet_name="checks", index=False)
        pd.DataFrame(report["excluded_rows"] or [{"row": None}]).to_excel(writer, sheet_name="excluded_rows", index=False)
        owner = report["owner"]
        pd.DataFrame([{"status": report["status"], "owner_inn": owner.get("inn"), "owner_name": owner.get("name"),
                       "owner_confidence": owner.get("confidence"), "owner_method": owner.get("method"),
                       "notes": "; ".join(report["notes"])}]).to_excel(writer, sheet_name="summary", index=False)
    return target


def run_folder(folder: str | Path, output_dir: str | Path | None = None, llm: Any | None = None) -> pd.DataFrame:
    rows = []
    files = [p for p in sorted(Path(folder).rglob("*"))
             if p.suffix.lower() in READERS and not p.name.endswith("__normalized.xlsx") and not p.name.startswith("~$")]
    batch_started = time.time()
    for n, path in enumerate(files, 1):
        print(f"\n===== файл {n}/{len(files)}: {path.name} (прошло {time.time() - batch_started:.0f} с) =====", flush=True)
        try:
            runner = globals().get("parse_statement_auto", parse_statement)
            table, report = runner(path, llm=llm)
            target = export_result(table, report, path, output_dir)
            rows.append({"file": str(path), "operations": len(table), "status": report["status"],
                         "owner_inn": report["owner"]["inn"], "owner_name": report["owner"]["name"],
                         "errors": "; ".join(report["errors"]), "output": str(target)})
        except Exception as exc:
            rows.append({"file": str(path), "operations": 0, "status": "failed", "errors": f"{type(exc).__name__}: {exc}"})
    return pd.DataFrame(rows)


# агент
import ast
import signal
import threading
from typing import Annotated, Optional, TypedDict

AGENT_ENABLED = True
AGENT_MAX_STEPS = 10
AGENT_PATIENCE = 5
READ_ONLY_TOOLS = {"get_overview", "show_cells", "search_cells", "column_cards", "show_result", "evaluate_result"}
AGENT_MAX_REPEATS = 2
AGENT_HISTORY_EXCHANGES = 6
AGENT_CODE_TIMEOUT = 10
AGENT_VERBOSE = True

try:
    from langchain_core.messages import AIMessage, HumanMessage, SystemMessage, ToolMessage
    from langchain_core.tools import StructuredTool
    from langgraph.graph import END, StateGraph
    from langgraph.graph.message import add_messages
    from pydantic import BaseModel, Field
    AGENT_AVAILABLE = True
except Exception as _agent_import_error:
    AGENT_AVAILABLE = False
    pass


def alog(message: str) -> None:
    if AGENT_VERBOSE:
        print(f"[агент] {message}")


AGENT_SYSTEM = """Ты — агент разбора банковских выписок. Файл уже разобран алгоритмом, но его самопроверка
нашла проблемы. Твоя задача — довести результат до правильного, используя инструменты.

ЦЕЛЬ. Таблица операций: date, debit_inn, debit_name, credit_inn, credit_name, amount, purpose.
debit_* — плательщик (с чьего счёта ушли деньги), credit_* — получатель. Владелец выписки — клиент банка,
по чьему счёту выписка: при списании он плательщик, при зачислении — получатель. Банк контрагента
(корр. счёт, БИК, «ПАО Сбербанк» в графе «реквизиты банка…») — НЕ контрагент.

КАК РАБОТАТЬ
В первом сообщении уже есть ВЫБОРКА: 50 первых операций и 50 случайных. Для каждой показаны исходные ячейки
таблицы (C<колонка>=значение) и то, что из них получилось (Д — дебет/плательщик, К — кредит/получатель, сумма, назначение).
Смотреть данные дополнительно почти не нужно — ты уже их видишь.
1. СРАЗУ сравни исходные ячейки с результатом и назови, что неверно: сторона пустая или чужая, одна и та же сторона
   в дебете и кредите, в стороне оказалось назначение или банк, суммы/даты не те, операции пропущены или лишние.
2. СРАЗУ исправь причиной, а не симптомом — первый же шаг должен быть изменением, а не просмотром:
   - неверные роли колонок → set_column_roles (номера колонок C из выборки; -1 снимает роль);
   - перепутаны списание/зачисление у всех операций → flip_direction;
   - итоги/остатки/служебные строки попали в операции или наоборот → relabel_rows;
   - таблица найдена не там → set_table_bounds; владелец неверный → set_owner;
   - необычная раскладка (блоки «метка: значение», повёрнутая таблица, текст) → run_parser_code (напиши parse)
     или extract_from_text.
3. После каждого изменения инструмент возвращает самопроверку, оценку и ИЗМЕНЕНИЯ НА ВЫБОРКЕ (было → стало).
   Проверяй по ним, что исправил именно то, что хотел. Стало хуже — верни или попробуй другую гипотезу.
4. Просмотр (show_cells, search_cells, column_cards, show_result) — только если в выборке не хватает данных,
   и не больше двух подряд. Дальше обязательно изменение.
5. finish — когда ошибок нет, или когда исправить нельзя (например, ИНН в документе отсутствует): назови причину.

ПРАВИЛА
- Один инструмент за ход. Номера строк и колонок бери только из того, что видел.
- Ничего не придумывай: ИНН, суммы, названия берутся из ячеек; инструменты проверяют это сами.
- Владелец — настоящий клиент банка: организация (ООО/АО/ИП…) или ФИО, с ИНН из 10/12 цифр. НЕ владелец: None, пустое,
  подпись формы («код налогового органа», «сведения о…»), налоговый орган, банк, контрагент. Если в документе владельца нет — так и скажи
  в finish, но не подставляй мусор. В конце сам перечитай результат глазами человека: владелец на своей стороне в каждой операции?
- Итоги, остатки, обороты, «Итого за день» — не операции.
- Назначение — длинный текст, а не коды вида операции («01», «17»).
- Не трать шаги на то, что уже проверено."""

AGENT_CODE_HELP = """Функция parse(rows) получает rows — список строк листа, каждая строка — список строк-ячеек
(пустая ячейка = ""). Верни список словарей:
  {"row": номер строки, где стоит дата операции, "rows": [все строки, откуда взяты значения] (необязательно),
   "date": "дд.мм.гггг" как в ячейке,
   "amount": "сумма как в ячейке", "direction": "out" (списание) или "in" (зачисление),
   "counterparty_name": ..., "counterparty_inn": ..., "purpose": ...}
Если в таблице обе стороны операции, вместо counterparty_* верни payer_name/payer_inn (плательщик)
и payee_name/payee_inn (получатель). Владелец выписки подставляется сам, его указывать не нужно.
Можно: циклы, условия, re, str-методы, parse_date(text), parse_amount(text), inn_ok(text).
Нельзя: import, open, eval/exec, атрибуты с «_». Время — не больше 10 секунд.
Каждое значение должно дословно встречаться в строках rows (±1 строка) — иначе поле
будет отброшено, а операция без подтверждённой даты или суммы — исключена."""


import builtins as _builtins

_SAFE_BUILTINS = {name: getattr(_builtins, name)
                  for name in ("len", "range", "enumerate", "zip", "min", "max", "sum", "sorted", "reversed",
                               "list", "dict", "set", "tuple", "str", "int", "float", "bool", "abs", "round",
                               "any", "all", "isinstance", "map", "filter", "print")}
_FORBIDDEN_NAMES = {"open", "exec", "eval", "compile", "__import__", "globals", "locals", "vars", "getattr",
                    "setattr", "delattr", "input", "breakpoint", "help", "exit", "quit", "memoryview", "type", "object"}


def check_agent_code(code: str) -> str | None:
    try:
        tree = ast.parse(code)
    except SyntaxError as exc:
        return f"синтаксическая ошибка: {exc.msg} (строка {exc.lineno})"
    for node in ast.walk(tree):
        if isinstance(node, (ast.Import, ast.ImportFrom)):
            return "import запрещён: re, parse_date, parse_amount, inn_ok уже доступны"
        if isinstance(node, ast.Attribute) and node.attr.startswith("_"):
            return f"атрибут {node.attr} запрещён"
        if isinstance(node, ast.Name) and (node.id in _FORBIDDEN_NAMES or node.id.startswith("__")):
            return f"имя {node.id} запрещено"
        if isinstance(node, (ast.Global, ast.Nonlocal, ast.AsyncFunctionDef, ast.Await, ast.With)):
            return f"конструкция {type(node).__name__} запрещена"
    if not any(isinstance(n, ast.FunctionDef) and n.name == "parse" for n in tree.body):
        return "нужна функция parse(rows)"
    return None


class _Timeout(Exception):
    pass


def run_agent_code(code: str, rows: list[list[str]]) -> tuple[Any, str | None]:
    error = check_agent_code(code)
    if error:
        return None, error
    env = {"__builtins__": _SAFE_BUILTINS, "re": re, "parse_date": parse_date, "parse_amount": parse_amount,
           "inn_ok": inn_ok}
    use_alarm = threading.current_thread() is threading.main_thread() and hasattr(signal, "SIGALRM")

    def _handler(signum, frame):
        raise _Timeout()

    try:
        if use_alarm:
            old = signal.signal(signal.SIGALRM, _handler)
            signal.alarm(AGENT_CODE_TIMEOUT)
        exec(compile(code, "<agent>", "exec"), env)
        result = env["parse"]([list(r) for r in rows])
    except _Timeout:
        return None, f"превышено время {AGENT_CODE_TIMEOUT} с"
    except Exception as exc:
        return None, f"{type(exc).__name__}: {exc}"
    finally:
        if use_alarm:
            signal.alarm(0)
            signal.signal(signal.SIGALRM, old)
    if not isinstance(result, list):
        return None, "parse должна вернуть список словарей"
    return result, None


class Workbench:
    def __init__(self, path: str | Path, table: pd.DataFrame, report: dict[str, Any], llm: Any | None):
        st = report["_state"]
        self.path = path
        self.doc: Document = st["doc"]
        self.regions: list[Region] = list(st["regions"])
        self.specs: dict[int, TableSpec] = dict(st["specs"])
        self.verdicts: dict[int, list[RowVerdict]] = {k: list(v) for k, v in st["verdicts"].items()}
        self.owner: dict[str, Any] = dict(st["owner"])
        self.owner_fixed = False
        self.ops_override: list[dict[str, Any]] | None = None
        self.llm = llm
        self.base_report = report
        self.current = (table, report)
        self.best = (report.get("score", -99.0), table, report, "алгоритм")
        self.trace: list[dict[str, Any]] = []
        self.finished: str | None = None
        self.reads_in_row = 0
        self.flip = False
        self.since_best = 0
        self.seen_calls: Counter = Counter()
        self.grids = {g.name: g for g in doc_grids(self.doc)}

    @staticmethod
    def _op_key(extra: dict[str, Any]) -> tuple:
        return (extra.get("sheet"), (extra.get("rows") or [None])[0])

    def _cells_line(self, sheet: str | None, row1: int | None, cap: int = 230) -> str:
        grid = self.grids.get(sheet) or (next(iter(self.grids.values())) if len(self.grids) == 1 else None)
        if grid is None or row1 is None or not 0 <= row1 - 1 < grid.n_rows:
            return "(строка недоступна)"
        text = " | ".join(f"C{c.c}={c.text.replace(chr(10), ' ')[:24]}" for c in grid.row(row1 - 1) if c.text.strip())
        return f"R{row1 - 1}: " + (text if len(text) <= cap else text[:cap] + "…")

    @staticmethod
    def _result_line(row) -> str:
        nz = lambda v: "" if str(v) in ("<NA>", "nan", "None", "NaT") else str(v)
        return (f"Д: {nz(row['debit_inn'])} {nz(row['debit_name'])[:28]} | К: {nz(row['credit_inn'])} {nz(row['credit_name'])[:28]}"
                f" | {nz(row['amount'])} | {nz(row['purpose'])[:30]}")

    def init_sample(self, n_top: int = 50, n_random: int = 50, seed: int = 11) -> None:
        import random
        table, rep = self.current
        extra = rep.get("operations_extra") or []
        n = min(len(table), len(extra))
        self.sample_keys: list[tuple] = []
        self.sample_raw_rows: list[tuple] = []
        if n:
            top = list(range(min(n_top, n)))
            rest = list(range(len(top), n))
            rnd = sorted(random.Random(seed).sample(rest, min(n_random, len(rest))))
            self.sample_keys = [self._op_key(extra[i]) for i in top + rnd]
            self.sample_n_top = len(top)
        else:
            rows = [(g.name, r + 1) for g in self.grids.values() for r in range(g.n_rows) if g.row(r)]
            top = rows[:n_top]
            rnd = sorted(random.Random(seed).sample(rows[n_top:], min(n_random, max(len(rows) - n_top, 0))))
            self.sample_raw_rows = top + rnd
            self.sample_n_top = len(top)
        self.sample_prev = self._snapshot()

    def _snapshot(self) -> dict[tuple, str]:
        table, rep = self.current
        extra = rep.get("operations_extra") or []
        by_key = {self._op_key(extra[i]): self._result_line(table.iloc[i]) for i in range(min(len(table), len(extra)))}
        return {k: by_key.get(k, "— операции нет —") for k in self.sample_keys}

    def sample_text(self) -> str:
        if getattr(self, "sample_raw_rows", None):
            lines = [f"ВЫБОРКА: операций алгоритм НЕ нашёл; вот {len(self.sample_raw_rows)} строк документа "
                     f"({self.sample_n_top} первых и остальные случайные):"]
            lines += [f"«{sheet}» {self._cells_line(sheet, r1)}" for sheet, r1 in self.sample_raw_rows]
            return "\n".join(lines)
        if not getattr(self, "sample_keys", None):
            return "ВЫБОРКА пуста"
        snap = self._snapshot()
        lines = [f"ВЫБОРКА: {self.sample_n_top} первых и {len(self.sample_keys) - self.sample_n_top} случайных операций "
                 f"(исходные ячейки → результат):"]
        for k in self.sample_keys:
            lines.append(self._cells_line(k[0], k[1]))
            lines.append("   → " + snap[k])
        return "\n".join(lines)

    def sample_diff(self) -> str:
        if not getattr(self, "sample_keys", None):
            return ""
        now = self._snapshot()
        changed = [k for k in self.sample_keys if now[k] != self.sample_prev.get(k)]
        out = [f"ИЗМЕНЕНИЯ НА ВЫБОРКЕ: изменилось {len(changed)} из {len(self.sample_keys)}"]
        for k in changed[:8]:
            out.append(f"R{(k[1] or 1) - 1}: было {self.sample_prev.get(k)}\n      стало {now[k]}")
        if len(changed) > 8:
            out.append(f"… и ещё {len(changed) - 8}")
        self.sample_prev = now
        return "\n".join(out)

    def evaluate(self, action: str) -> str:
        report = {"file": str(self.path), "notes": [f"агент: {action}"], "gigachat": self.llm is not None,
                  "mode": "agent"}
        if self.ops_override is not None:
            raw = self.ops_override
            regions, specs, verdicts = [], {}, {}
        else:
            regions, specs, verdicts = self.regions, self.specs, self.verdicts
            raw = []
            for i, region in enumerate(regions):
                raw += records_to_raw_ops(build_records(region, specs[i], verdicts[i]), specs[i])
        regions_by_grid: dict[str, list[Region]] = defaultdict(list)
        for region in regions:
            regions_by_grid[region.grid.name].append(region)
        if self.flip:
            raw = copy.deepcopy(raw)
            for op in raw:
                if op.get("direction"):
                    op["direction"] = "in" if op["direction"] == "out" else "out"
        if not self.owner_fixed:
            self.owner = detect_owner(self.doc, dict(regions_by_grid), raw, allow_llm=False)
        ops, notes = finalize_ops(copy.deepcopy(raw), self.owner)
        refs = find_references(self.doc, regions, verdicts, specs)
        result = run_checks(ops, refs, verdicts, regions, specs, self.owner)
        table, rep = compose_output(self.path, self.doc, report, regions, specs, verdicts, self.owner, ops, refs,
                                    result, raw_ops=raw, regions_by_grid=dict(regions_by_grid))
        self.current = (table, rep)
        improved = rep["score"] > self.best[0]
        if improved:
            self.best = (rep["score"], table, rep, action)
        diff = self.sample_diff() if getattr(self, "sample_keys", None) else ""
        return self.summary(rep, improved) + ("\n" + diff if diff else "")

    def summary(self, rep: dict[str, Any], improved: bool | None = None) -> str:
        errors = rep["errors"] or ["нет"]
        lines = [f"статус: {rep['status']}; операций: {len(rep['operations_extra'])}; оценка: {rep['score']}"
                 + (f" (лучшая пока: {self.best[0]})" if improved is not None else "")]
        if improved:
            lines.append("это новый лучший вариант — сохранён")
        lines.append("ошибки самопроверки: " + " | ".join(e[:160] for e in errors[:6]))
        owner = rep["owner"]
        lines.append(f"владелец: ИНН={owner.get('inn')}, наименование={owner.get('name')!r}")
        return "\n".join(lines)

    def region(self, idx: int) -> Region:
        if not 0 <= idx < len(self.regions):
            raise ValueError(f"таблицы {idx} нет; есть 0..{len(self.regions) - 1}")
        return self.regions[idx]

    def t_overview(self, with_result: bool = True) -> str:
        table, rep = self.current
        parts = ["ЛИСТЫ: " + "; ".join(f"«{g.name}» строк {g.n_rows}, колонок {g.n_cols}" for g in self.grids.values())]
        if self.ops_override is not None:
            parts.append("ТЕКУЩИЙ РЕЖИМ: операции из текста/кода агента")
        for i, region in enumerate(self.regions):
            spec = self.specs[i]
            parts.append(f"ТАБЛИЦА {i}: лист «{region.grid.name}», шапка строки {region.header_rows}, данные "
                         f"R{region.first_anchor}–R{region.end_row}; схема: {json.dumps(_spec_brief(spec), ensure_ascii=False)}")
        if not self.regions and self.ops_override is None:
            parts.append("ТАБЛИЦА ОПЕРАЦИЙ НЕ НАЙДЕНА алгоритмом (возможно, операции блоками, текстом или таблица повёрнута)")
        parts.append("САМОПРОВЕРКА: " + self.summary(rep))
        if rep["errors"]:
            parts.append("ОШИБКИ ПОЛНОСТЬЮ:\n" + "\n".join(f"- {e}" for e in rep["errors"][:10]))
        if with_result:
            parts.append("ТЕКУЩИЙ РЕЗУЛЬТАТ (что сейчас получилось; полный список — show_result):\n" + self.t_show_result(0, 8))
        return "\n".join(parts)

    def t_show_result(self, start: int = 0, count: int = 15) -> str:
        table, rep = self.current
        extra = rep.get("operations_extra") or []
        if table.empty:
            return "операций нет"
        lines = []
        for i in range(max(start, 0), min(start + min(count, 40), len(table))):
            row = table.iloc[i]
            rows = ",".join(f"R{r}" for r in (extra[i]["rows"] if i < len(extra) else []))
            lines.append(f"#{i} [{rows}] {row['date'].date() if str(row['date']) != 'NaT' else None} | "
                         f"ДЕБЕТ(плательщик) {row['debit_inn']} {str(row['debit_name'])[:40]} → "
                         f"КРЕДИТ(получатель) {row['credit_inn']} {str(row['credit_name'])[:40]} | {row['amount']} | "
                         f"{str(row['purpose'])[:50]}")
        return f"всего операций {len(table)}\n" + "\n".join(lines)

    def t_show(self, sheet: str, row_from: int, row_to: int, col_from: int = 0, col_to: int | None = None) -> str:
        grid = self.grid(sheet)
        row_to = min(row_to, row_from + 39, grid.n_rows - 1)
        out = []
        for r in range(max(row_from, 0), row_to + 1):
            cells = [c for c in grid.row(r) if c.c >= col_from and (col_to is None or c.c <= col_to)]
            if cells:
                out.append(f"R{r}: " + " | ".join(f"C{c.c}{'-' + str(c.c2) if c.c2 > c.c else ''}={c.text[:60]}" for c in cells))
        text = "\n".join(out) or "пусто"
        return text[:9000]

    def t_search(self, pattern: str, sheet: str | None = None) -> str:
        try:
            rx = re.compile(pattern, re.I)
        except re.error as exc:
            return f"неверное регулярное выражение: {exc}"
        hits = []
        for grid in self.grids.values():
            if sheet and grid.name != sheet:
                continue
            for c in grid.cells:
                if rx.search(c.text):
                    hits.append(f"«{grid.name}» R{c.r}C{c.c}: {c.text[:80]}")
                    if len(hits) >= 30:
                        return "\n".join(hits) + "\n… (показаны первые 30)"
        return "\n".join(hits) or "ничего не найдено"

    def t_cards(self, table: int) -> str:
        region = self.region(table)
        return "\n".join(region.profiles[c].card() for c in region.columns if region.profiles[c].n > 0)

    def t_set_bounds(self, sheet: str, header_rows: list[int], first_data_row: int, last_data_row: int) -> str:
        grid = self.grid(sheet)
        kinds = KindCache()
        anchors, date_col = find_anchor_rows(grid, kinds)
        anchors = [r for r in anchors if first_data_row <= r <= last_data_row]
        if not anchors:
            anchors = [r for r in range(first_data_row, last_data_row + 1) if grid.row(r)]
        cols = sorted({c.c for r in list(anchors) + list(header_rows) for c in grid.row(r)})
        region = Region(grid=grid, header_rows=sorted(header_rows), first_anchor=first_data_row, last_anchor=last_data_row,
                        end_row=min(last_data_row + 10, grid.n_rows - 1), columns=cols,
                        headers=header_paths(grid, sorted(header_rows), cols), anchors=anchors, kinds=kinds)
        region.profiles = build_profiles(region)
        spec = infer_table_spec(region)
        self.regions = [r for r in self.regions if r.grid.name != sheet] + [region]
        idx = len(self.regions) - 1
        self.specs = {i: (self.specs[i] if i < idx and i in self.specs else spec) for i in range(len(self.regions))}
        self.specs[idx] = spec
        self.verdicts = {i: (self.verdicts[i] if i < idx and i in self.verdicts else classify_rows(self.regions[i], self.specs[i]))
                         for i in range(len(self.regions))}
        self.verdicts[idx] = classify_rows(region, spec)
        self.ops_override = None
        return (f"таблица {idx} задана; схема алгоритма: {json.dumps(_spec_brief(spec), ensure_ascii=False)}\n"
                + self.evaluate(f"границы таблицы {sheet} R{first_data_row}–R{last_data_row}"))

    def t_set_roles(self, table: int, **roles) -> str:
        region = self.region(table)
        old = self.specs[table]
        data = _spec_brief(old)
        for key in ("date", "debit", "credit", "signed", "amount", "direction", "purpose"):
            if key in roles and roles[key] is not None:
                data[key] = None if roles[key] < 0 else roles[key]
        for side in ("payer", "payee", "counterparty"):
            data[side] = dict(data[side])
            for part in ("inn", "name", "composite"):
                v = roles.get(f"{side}_{part}")
                if v is not None:
                    data[side][part] = None if v < 0 else v
        if roles.get("direction_out"):
            data["direction_out"] = [x.strip() for x in roles["direction_out"].split(",")]
        if roles.get("direction_in"):
            data["direction_in"] = [x.strip() for x in roles["direction_in"].split(",")]
        spec = spec_from_json({**data, "direction_out": data.get("direction_out") or old.direction_out,
                               "direction_in": data.get("direction_in") or old.direction_in})
        errors = check_spec(spec, region)
        if errors:
            return "схема НЕ применена, проверка не прошла: " + "; ".join(errors)
        spec.source = "agent"
        self.specs[table] = spec
        self.verdicts[table] = classify_rows(region, spec)
        self.ops_override = None
        return self.evaluate(f"роли колонок таблицы {table}: {json.dumps(_spec_brief(spec), ensure_ascii=False)}")

    def t_relabel(self, table: int, labels: str) -> str:
        self.region(table)
        by_r = {v.r: v for v in self.verdicts[table]}
        changed, bad = [], []
        for item in re.split(r"[,;\n]+", labels):
            m = re.match(r"\s*R?(\d+)\s*[=:]\s*([A-Z_]+)\s*$", item)
            if not m:
                if item.strip():
                    bad.append(item.strip())
                continue
            r, label = int(m.group(1)), m.group(2)
            if r not in by_r or label not in ROW_LABELS:
                bad.append(item.strip())
                continue
            by_r[r].label, by_r[r].decided_by, by_r[r].reason = label, "agent", "агент"
            changed.append(f"R{r}={label}")
        if bad:
            return f"не распознано (формат R57=SUBTOTAL, строки только из таблицы {table}): {bad}"
        return self.evaluate(f"переразметка строк {changed}")

    def t_set_owner(self, inn: str | None = None, name: str | None = None) -> str:
        problems = owner_problems({"inn": re.sub(r"\D", "", inn or "") if inn else None, "name": name})
        if name and not problems and not (OPF_RE.search(name) or PERSON_RE.match(name.strip())
                                          or PERSON_INITIALS_RE.match(name.strip())):
            problems.append(f"«{name}» не похоже на организацию, ИП или ФИО (нет формы собственности «ООО/АО/ИП…» и не ФИО)")
        if name and not any(c.text.strip().casefold() == name.strip().casefold() or names_match(name, c.text)
                            for g in self.grids.values() for c in g.cells):
            problems.append(f"наименования «{name}» нет в документе ни одной целой ячейкой/строкой — подстрока слова не считается")
        if problems:
            return ("владелец НЕ изменён: " + "; ".join(problems)
                    + ". Выбери настоящего клиента банка (ООО/АО/ИП/ФИО) из шапки: show_cells / search_cells по «в отношении», «владелец», «ИНН».")
        if inn:
            digits = re.sub(r"\D", "", inn)
            if not inn_ok(digits):
                problems.append(f"ИНН {digits} не проходит контрольную сумму")
            elif digits not in self._document_digits():
                problems.append(f"ИНН {digits} не найден в документе")
            else:
                inn = digits
        if name:
            if not any(names_match(name, c.text) or name.casefold() in c.text.casefold()
                       for g in self.grids.values() for c in g.cells):
                problems.append(f"наименование «{name}» не найдено в документе")
        if problems:
            return "владелец НЕ изменён: " + "; ".join(problems)
        self.owner = {**self.owner, "inn": inn or self.owner.get("inn"), "name": name or self.owner.get("name"),
                      "method": "agent", "conflicts": [], "inn_confidence": 0.9, "name_confidence": 0.9, "confidence": 0.9}
        self.owner_fixed = True
        return self.evaluate(f"владелец ИНН={inn} наименование={name!r}")

    def t_extract_text(self, sheet: str, row_from: int, row_to: int) -> str:
        if self.llm is None:
            return "GigaChat не подключён — извлечение из текста недоступно, используй run_parser_code"
        grid = self.grid(sheet)
        lines = [grid.row_text(r, sep="  ") for r in range(row_from, min(row_to, grid.n_rows - 1) + 1)]
        raw, notes = llm_text_extract(lines, self.llm)
        for op in raw:
            op["_rows"] = [row_from + i for i in op["_rows"]]
            op["_grid"] = sheet
        if not raw:
            return "операции не извлечены: " + "; ".join(notes[:5])
        self.ops_override = raw
        return self.evaluate(f"извлечение из текста R{row_from}–R{row_to}") + ("\n" + "; ".join(notes[:5]) if notes else "")

    def t_run_code(self, sheet: str, code: str) -> str:
        grid = self.grid(sheet)
        rows = [[grid.text_at(r, c) if (cell := grid.at(r, c)) is not None and cell.r == r and cell.c == c else ""
                 for c in range(grid.n_cols)] for r in range(grid.n_rows)]
        result, error = run_agent_code(code, rows)
        if error:
            return f"код не выполнен: {error}\n\n{AGENT_CODE_HELP}"
        raw, dropped = [], []
        for k, item in enumerate(result[:5000]):
            if not isinstance(item, dict) or not isinstance(item.get("row"), int):
                dropped.append(f"#{k}: нет целого поля row")
                continue
            r = item["row"]
            used = [x for x in (item.get("rows") or [r]) if isinstance(x, int)] or [r]
            window = sorted({rr for x in used for rr in range(x - 1, x + 2) if 0 <= rr < len(rows)})
            text = "\n".join(" | ".join(rows[rr]) for rr in window)
            fields = {}
            for key, kind in (("date", "date"), ("amount", "amount"), ("counterparty_inn", "inn"),
                              ("counterparty_name", "text"), ("purpose", "text"), ("payer_inn", "inn"),
                              ("payer_name", "text"), ("payee_inn", "inn"), ("payee_name", "text")):
                v = item.get(key)
                fields[key] = str(v) if v not in (None, "") and _grounded(str(v), text, kind) else None
            op_date, amount = parse_date(fields["date"]), parse_amount(fields["amount"])
            if op_date is None or amount is None:
                dropped.append(f"R{r}: нет подтверждённой даты или суммы")
                continue
            direction = item.get("direction") if item.get("direction") in ("in", "out") else ("out" if amount < 0 else None)
            raw.append({"date": op_date, "amount": abs(amount), "direction": direction,
                        "counterparty_inn": re.sub(r"\D", "", fields["counterparty_inn"] or "") or None,
                        "counterparty_name": fields["counterparty_name"],
                        "payer_inn": re.sub(r"\D", "", fields["payer_inn"] or "") or None, "payer_name": fields["payer_name"],
                        "payee_inn": re.sub(r"\D", "", fields["payee_inn"] or "") or None, "payee_name": fields["payee_name"],
                        "purpose": fields["purpose"], "_rows": [r],
                        "_grid": sheet, "_date_inherited": False, "_provenance": {}})
        if not raw:
            return "парсер не дал ни одной подтверждённой операции: " + "; ".join(dropped[:8])
        self.ops_override = raw
        note = f"\nотброшено {len(dropped)}: " + "; ".join(dropped[:5]) if dropped else ""
        return self.evaluate(f"парсер агента: {len(raw)} операций") + note

    def t_flip_direction(self) -> str:
        self.flip = not self.flip
        return self.evaluate(f"направление всех операций {'перевёрнуто' if self.flip else 'возвращено'}")

    def t_evaluate(self) -> str:
        return self.summary(self.current[1])

    def t_finish(self, reason: str) -> str:
        bad = owner_problems(self.current[1].get("owner") or {})
        if bad and not getattr(self, "_finish_refused", False):
            self._finish_refused = True
            return ("finish НЕ принят: результат содержит заведомо неверного владельца — " + "; ".join(bad)
                    + ". Найди настоящего владельца (set_owner) или верни прежнего; если в документе его нет — вызови finish ещё раз и объясни.")
        self.finished = reason
        return "готово"

    def grid(self, sheet: str) -> Grid:
        if sheet in self.grids:
            return self.grids[sheet]
        if len(self.grids) == 1:
            return next(iter(self.grids.values()))
        raise ValueError(f"листа «{sheet}» нет; есть: {list(self.grids)}")

    def nothing_to_find(self) -> str | None:
        errors = self.base_report.get("errors") or []
        if not errors or not all(e.startswith("ИНН контрагента указан") for e in errors):
            return None
        valid = {m for m in re.findall(r"(?<!\d)(\d{10}|\d{12})(?!\d)", self._document_digits()) if inn_ok(m)}
        valid.discard(self.owner.get("inn"))
        if valid:
            return None
        return "единственная проблема — у контрагентов нет ИНН, и в документе нет ни одного ИНН, кроме владельца; искать нечего"

    def _document_digits(self) -> str:
        if not hasattr(self, "_digits_cache"):
            parts = []
            for g in self.grids.values():
                for r in range(g.n_rows):
                    row = g.row(r)
                    parts.append("".join(re.sub(r"\D", "", c.text) for c in row if re.fullmatch(r"[\d ]+", c.text)))
                    parts.extend(re.sub(r"(?<=\d)[  ](?=\d)", "", c.text) for c in row)
            self._digits_cache = " ".join(parts)
        return self._digits_cache


def doc_grids(doc: Document) -> list[Grid]:
    return list(doc.grids)


if AGENT_AVAILABLE:
    class _NoArgs(BaseModel):
        pass

    class _ShowArgs(BaseModel):
        sheet: str = Field(description="имя листа из get_overview")
        row_from: int = Field(description="первая строка (R-номер)")
        row_to: int = Field(description="последняя строка, не больше row_from+39")
        col_from: int = Field(0, description="первая колонка (C-номер)")
        col_to: Optional[int] = Field(None, description="последняя колонка или null")

    class _SearchArgs(BaseModel):
        pattern: str = Field(description="регулярное выражение, например 'ИНН|владел'")
        sheet: Optional[str] = Field(None, description="лист или null — все листы")

    class _ResultArgs(BaseModel):
        start: int = Field(0, description="с какой операции показать")
        count: int = Field(15, description="сколько операций (до 40)")

    class _TableArgs(BaseModel):
        table: int = Field(0, description="номер таблицы из get_overview")

    class _BoundsArgs(BaseModel):
        sheet: str = Field(description="лист")
        header_rows: list[int] = Field(description="строки шапки таблицы операций (можно пустой список)")
        first_data_row: int = Field(description="первая строка с операцией")
        last_data_row: int = Field(description="последняя строка с операцией")

    class _RolesArgs(BaseModel):
        table: int = Field(0, description="номер таблицы")
        date: Optional[int] = Field(None, description="колонка даты операции")
        debit: Optional[int] = Field(None, description="колонка суммы списания со счёта владельца")
        credit: Optional[int] = Field(None, description="колонка суммы зачисления на счёт владельца")
        signed: Optional[int] = Field(None, description="одна колонка суммы со знаком (минус = списание)")
        amount: Optional[int] = Field(None, description="одна сумма без знака (нужна direction)")
        direction: Optional[int] = Field(None, description="колонка направления (Д/К, списание/зачисление)")
        direction_out: Optional[str] = Field(None, description="значения направления для списания через запятую")
        direction_in: Optional[str] = Field(None, description="значения направления для зачисления через запятую")
        counterparty_name: Optional[int] = Field(None, description="колонка наименования контрагента")
        counterparty_inn: Optional[int] = Field(None, description="колонка ИНН контрагента")
        counterparty_composite: Optional[int] = Field(None, description="одна ячейка с ИНН+наименованием контрагента")
        payer_name: Optional[int] = Field(None, description="наименование плательщика (если в таблице обе стороны)")
        payer_inn: Optional[int] = Field(None, description="ИНН плательщика")
        payee_name: Optional[int] = Field(None, description="наименование получателя")
        payee_inn: Optional[int] = Field(None, description="ИНН получателя")
        purpose: Optional[int] = Field(None, description="колонка назначения платежа (длинный текст)")

    class _RelabelArgs(BaseModel):
        table: int = Field(0, description="номер таблицы")
        labels: str = Field(description="через запятую: R57=SUBTOTAL, R60=OPERATION. Метки: " + ", ".join(ROW_LABELS))

    class _OwnerArgs(BaseModel):
        inn: Optional[str] = Field(None, description="ИНН владельца, как в документе")
        name: Optional[str] = Field(None, description="наименование владельца, как в документе")

    class _TextArgs(BaseModel):
        sheet: str = Field(description="лист")
        row_from: int = Field(description="первая строка")
        row_to: int = Field(description="последняя строка")

    class _CodeArgs(BaseModel):
        sheet: str = Field(description="лист, строки которого получит parse(rows)")
        code: str = Field(description="код Python с функцией parse(rows). " + AGENT_CODE_HELP)

    class _FinishArgs(BaseModel):
        reason: str = Field(description="почему закончил: что исправлено или почему дальше исправить нельзя")


def make_agent_tools(wb: Workbench) -> list:
    def safe(fn):
        def wrapper(**kwargs):
            try:
                return fn(**kwargs)
            except Exception as exc:
                return f"ошибка инструмента: {type(exc).__name__}: {exc}"
        return wrapper

    spec = [
        ("get_overview", "Обзор: листы, найденные таблицы, текущая схема колонок и ошибки самопроверки.", _NoArgs, wb.t_overview),
        ("show_result", "Текущий результат: операции с исходными строками, плательщиком, получателем, суммой и назначением.", _ResultArgs, wb.t_show_result),
        ("show_cells", "Показать ячейки листа с координатами R<строка> C<колонка> (до 40 строк).", _ShowArgs, wb.t_show),
        ("search_cells", "Найти ячейки по регулярному выражению.", _SearchArgs, wb.t_search),
        ("column_cards", "Карточки колонок таблицы: шапка, доли типов значений, примеры.", _TableArgs, wb.t_cards),
        ("set_table_bounds", "Заново задать границы таблицы операций на листе (если алгоритм нашёл не ту таблицу).", _BoundsArgs, wb.t_set_bounds),
        ("set_column_roles", "Назначить роли колонкам таблицы (номера колонок C). Незаданные роли остаются как были; -1 — убрать роль.", _RolesArgs, wb.t_set_roles),
        ("relabel_rows", "Переразметить строки таблицы (операция, итог, остаток…).", _RelabelArgs, wb.t_relabel),
        ("set_owner", "Задать владельца выписки (значения должны быть в документе).", _OwnerArgs, wb.t_set_owner),
        ("extract_from_text", "Извлечь операции из диапазона строк как из текста (для операций блоками/абзацами).", _TextArgs, wb.t_extract_text),
        ("run_parser_code", "Выполнить свою функцию parse(rows) для листа, если таблица нестандартная.", _CodeArgs, wb.t_run_code),
        ("flip_direction", "Перевернуть списание/зачисление у ВСЕХ операций (если дебет и кредит поняты наоборот; повторный вызов возвращает как было).", _NoArgs, wb.t_flip_direction),
        ("evaluate_result", "Текущая самопроверка результата.", _NoArgs, wb.t_evaluate),
        ("finish", "Закончить работу.", _FinishArgs, wb.t_finish),
    ]
    return [StructuredTool.from_function(func=safe(fn), name=name, description=desc, args_schema=schema)
            for name, desc, schema, fn in spec]


if AGENT_AVAILABLE:
    class AgentState(TypedDict):
        messages: Annotated[list, add_messages]
        steps: int
        nudges: int


def _action_from_text(text: str, tool_names: set[str]) -> dict | None:
    data, err = extract_json(text or "")
    if err or not isinstance(data, dict):
        return None
    name = data.get("tool") or data.get("name") or data.get("function")
    args = data.get("args") or data.get("arguments") or data.get("parameters") or {}
    if isinstance(args, str):
        args, _ = extract_json(args)
    if name in tool_names and isinstance(args, dict):
        return {"name": name, "args": args, "id": f"text_{abs(hash(text)) % 10**8}", "type": "tool_call"}
    return None


def _trim(messages: list) -> list:
    head, tail = messages[:1], messages[1:]
    keep: list = []
    pairs = 0
    i = len(tail) - 1
    while i >= 0 and pairs < AGENT_HISTORY_EXCHANGES:
        keep.insert(0, tail[i])
        if isinstance(tail[i], AIMessage):
            pairs += 1
        i -= 1
    while keep and isinstance(keep[0], ToolMessage):
        keep.pop(0)
    return head + keep


def build_agent_graph(wb: Workbench, llm: Any):
    tools = make_agent_tools(wb)
    by_name = {t.name: t for t in tools}
    model = llm.bind_tools(tools) if hasattr(llm, "bind_tools") else llm

    def agent_node(state: AgentState):
        msgs = [SystemMessage(content=AGENT_SYSTEM)] + _trim(state["messages"])
        try:
            ai = guarded_call(lambda: model.invoke(msgs),
                              f"агент, итерация {state['steps'] + 1}/{AGENT_MAX_STEPS}: выбор действия", "agent")
        except Exception as exc:
            alog(f"модель не ответила: {type(exc).__name__}: {exc} — останавливаюсь, оставляю лучшее из найденного")
            wb.finished = f"модель не ответила: {type(exc).__name__}: {exc}"
            return {"messages": [AIMessage(content="ошибка модели")]}
        if not getattr(ai, "tool_calls", None):
            action = _action_from_text(getattr(ai, "content", ""), set(by_name))
            if action:
                ai = AIMessage(content="", tool_calls=[action])
        if getattr(ai, "tool_calls", None) and len(ai.tool_calls) > 1:
            ai = AIMessage(content=ai.content, tool_calls=ai.tool_calls[:1])
        return {"messages": [ai]}

    def tools_node(state: AgentState):
        ai = state["messages"][-1]
        call = ai.tool_calls[0]
        tool = by_name.get(call["name"])
        best_before = wb.best[0]
        call_key = (call["name"], json.dumps(call.get("args") or {}, sort_keys=True, ensure_ascii=False, default=str))
        wb.seen_calls[call_key] += 1
        if tool is None:
            observation = f"инструмента {call['name']} нет; есть: {sorted(by_name)}"
        elif wb.seen_calls[call_key] > AGENT_MAX_REPEATS:
            observation = (f"этот вызов уже выполнялся {wb.seen_calls[call_key] - 1} раз(а) с тем же результатом — "
                           f"попробуй другое действие или вызови finish")
            wb.finished = "агент зациклился на одном и том же вызове"
        else:
            observation = tool.invoke(call.get("args") or {})
        wb.trace.append({"step": state["steps"] + 1, "tool": call["name"],
                         "args": {k: (v if not isinstance(v, str) or len(v) < 200 else v[:200] + "…")
                                  for k, v in (call.get("args") or {}).items()},
                         "observation": str(observation)[:400], "score": wb.current[1].get("score"),
                         "status": wb.current[1].get("status")})
        if call["name"] in READ_ONLY_TOOLS:
            wb.reads_in_row += 1
        else:
            wb.reads_in_row = 0
        if wb.reads_in_row == 2:
            observation = str(observation) + ("\n\n⚠ Ты второй раз подряд только смотришь данные — всё нужное уже есть в выборке. Следующим шагом ИЗМЕНИ что-то "
                                              "(set_column_roles, relabel_rows, set_owner, flip_direction, set_table_bounds, "
                                              "run_parser_code) или вызови finish с объяснением, что исправить нельзя.")
        elif wb.reads_in_row >= 4:
            wb.finished = "агент только изучал данные (4 просмотра подряд) и не предложил исправления"
        rep_now = wb.current[1]
        if wb.best[0] > best_before:
            wb.since_best = 0
            verdict = f"УЛУЧШЕНИЕ {best_before} → {wb.best[0]}"
        else:
            wb.since_best += 1
            verdict = f"без улучшения ({wb.since_best}/{AGENT_PATIENCE})"
        wb.trace[-1]["verdict"] = verdict
        args_short = ", ".join(f"{k}={str(v)[:40]!r}" for k, v in (call.get("args") or {}).items())
        alog(f"итерация {state['steps'] + 1}/{AGENT_MAX_STEPS}: {call['name']}({args_short}) → статус {rep_now['status']}, "
             f"операций {len(rep_now['operations_extra'])}, оценка {rep_now.get('score')}; {verdict}; "
             f"ошибок: {len(rep_now['errors'])} [{CLOCK.elapsed():.0f} с от старта, агенту осталось {max(CLOCK.left('agent'), 0):.0f} с]")
        return {"messages": [ToolMessage(content=str(observation)[:9000], tool_call_id=call["id"], name=call["name"])],
                "steps": state["steps"] + 1}

    def nudge_node(state: AgentState):
        return {"messages": [HumanMessage(content="Вызови один из инструментов (или finish с причиной). "
                                                  "Ответ — только вызов функции.")],
                "nudges": state["nudges"] + 1}

    def after_agent(state: AgentState) -> str:
        if wb.finished:
            return END
        if getattr(state["messages"][-1], "tool_calls", None):
            return "tools"
        return "nudge" if state["nudges"] < 1 else END

    def after_tools(state: AgentState) -> str:
        rep = wb.current[1]
        if wb.finished or state["steps"] >= AGENT_MAX_STEPS:
            wb.finished = wb.finished or f"исчерпан лимит итераций ({AGENT_MAX_STEPS})"
            return END
        if wb.since_best >= AGENT_PATIENCE:
            wb.finished = f"{AGENT_PATIENCE} итераций подряд без улучшения — стоп"
            return END
        if CLOCK.left("agent") <= 1:
            wb.finished = f"исчерпан бюджет времени агента ({AGENT_BUDGET} с)"
            return END
        if rep["status"] in ("ok", "ok_unverified") and wb.trace and wb.trace[-1]["tool"] not in ("get_overview",):
            wb.finished = "самопроверка пройдена"
            return END
        return "agent"

    graph = StateGraph(AgentState)
    graph.add_node("agent", agent_node)
    graph.add_node("tools", tools_node)
    graph.add_node("nudge", nudge_node)
    graph.set_entry_point("agent")
    graph.add_conditional_edges("agent", after_agent, {"tools": "tools", "nudge": "nudge", END: END})
    graph.add_conditional_edges("tools", after_tools, {"agent": "agent", END: END})
    graph.add_edge("nudge", "agent")
    return graph.compile()


def run_agent(path: str | Path, table: pd.DataFrame, report: dict[str, Any], llm: Any) -> tuple[pd.DataFrame, dict[str, Any]]:
    CLOCK.start_agent()
    wb = Workbench(path, table, report, llm)
    hopeless = wb.nothing_to_find()
    if hopeless:
        alog(f"не запускаю: {hopeless}")
        out = dict(report)
        out["agent"] = {"used": False, "skipped": hopeless}
        out["notes"] = list(report.get("notes", [])) + [f"агент не запускался: {hopeless}"]
        return table, out
    app = build_agent_graph(wb, llm)
    alog(f"запуск: ошибки алгоритма — {'; '.join(report['errors'][:3]) or 'нет'}; оценка алгоритма {report.get('score')}; "
         f"лимиты: {AGENT_MAX_STEPS} итераций, {AGENT_BUDGET} с, стоп после {AGENT_PATIENCE} итераций без улучшения")
    wb.init_sample()
    first = ("Файл: " + Path(str(path)).name + "\n" + wb.t_overview(with_result=False) + "\n\n" + wb.sample_text()
             + "\n\nСравни исходные ячейки с результатом, назови, что неверно, и сразу исправь инструментом "
               "(не просмотром). Каждое изменение покажет, что поменялось на выборке.")
    started = time.time()
    try:
        app.invoke({"messages": [HumanMessage(content=first)], "steps": 0, "nudges": 0},
                   config={"recursion_limit": 4 * AGENT_MAX_STEPS + 10})
    except Exception as exc:
        wb.finished = f"сбой агента: {type(exc).__name__}: {exc}"
    score, best_table, best_report, action = wb.best
    best_report = dict(best_report)
    base_score = report.get("score", -99.0)
    best_report["agent"] = {
        "used": True, "steps": len(wb.trace), "finish_reason": wb.finished or "бюджет шагов исчерпан",
        "baseline_score": base_score, "best_score": score, "improved": score > base_score,
        "best_action": action, "seconds": round(time.time() - started, 1), "trace": wb.trace,
    }
    if score > base_score:
        best_report["notes"] = list(report.get("notes", [])) + [f"агент улучшил результат ({base_score} → {score}): {action}"]
    else:
        best_report["notes"] = list(best_report.get("notes", [])) + ["агент не нашёл варианта лучше алгоритма"]
    alog(f"ИТОГ: оценка {base_score} → {score} ({'улучшено' if score > base_score else 'не улучшено, оставлен результат алгоритма'}); "
         f"статус {best_report['status']}; причина остановки: {best_report['agent']['finish_reason']}; "
         f"итераций {len(wb.trace)}, {best_report['agent']['seconds']} с")
    for t in wb.trace:
        alog(f"   {t['step']:>2}. {t['tool']:<18} оценка {t['score']}  {t.get('verdict', '')}")
    return best_table, best_report


REVIEW_SYSTEM = """Ты — проверяющий результата разбора банковской выписки. Решай по СМЫСЛУ, а не по формату.
Тебе дают владельца выписки, выборку операций и итоги самопроверки. Таблица: debit_* — плательщик, credit_* — получатель.
Владелец — клиент банка, по чьему счёту выписка: при списании он плательщик (debit), при зачислении — получатель (credit).

Проверь и назови ВСЕ найденные проблемы:
1. Владелец — настоящий клиент банка (организация ООО/АО/ИП… или ФИО) с ИНН. НЕ годятся: None/пусто, подпись формы
   («код налогового органа», «сведения о…»), налоговый орган, банк, контрагент из операций.
2. В каждой операции владелец стоит на своей стороне, а противоположная сторона — другое лицо (не владелец, не банк).
3. В полях сторон — наименования и ИНН, а не назначение платежа, не счета/БИК и не пустые значения без причины.
4. Назначение — осмысленный текст платежа.
5. Суммы и даты правдоподобны; нет итоговых строк среди операций.
Не придирайся к мелочам: иностранный контрагент без ИНН, физлицо без КПП — нормально.
Ответ — ТОЛЬКО JSON: {"verdict": "ok" | "suspect", "problems": ["кратко, что именно неверно"]}.
verdict=ok — только если по выборке результат можно считать законченным."""


def llm_final_review(table: pd.DataFrame, report: dict[str, Any], llm: Any) -> dict[str, Any]:
    import random
    CLOCK.start_review()
    owner = report.get("owner") or {}
    n = len(table)
    idx = list(range(min(8, n))) + sorted(random.Random(3).sample(range(min(8, n), n), min(8, max(n - 8, 0))))
    nz = lambda v: "" if str(v) in ("<NA>", "nan", "None", "NaT") else str(v)
    lines = [f"#{i}: {nz(table.iloc[i]['date'])[:10]} | Д: {nz(table.iloc[i]['debit_inn'])} {nz(table.iloc[i]['debit_name'])[:45]} | "
             f"К: {nz(table.iloc[i]['credit_inn'])} {nz(table.iloc[i]['credit_name'])[:45]} | {nz(table.iloc[i]['amount'])} | "
             f"{nz(table.iloc[i]['purpose'])[:60]}" for i in idx]
    user = (f"ВЛАДЕЛЕЦ: ИНН={owner.get('inn')!r}, наименование={owner.get('name')!r}\n"
            f"ОПЕРАЦИЙ ВСЕГО: {n}; статус самопроверки: {report.get('status')}\n"
            f"ОШИБКИ САМОПРОВЕРКИ: {'; '.join(e[:150] for e in (report.get('errors') or [])[:5]) or 'нет'}\n\nВЫБОРКА:\n" + "\n".join(lines))

    def validate(data):
        if not isinstance(data, dict) or data.get("verdict") not in ("ok", "suspect"):
            return None, 'ожидался JSON {"verdict": "ok" | "suspect", "problems": [...]}'
        problems = data.get("problems") or []
        if not isinstance(problems, list):
            return None, "problems должно быть списком строк"
        return {"verdict": data["verdict"], "problems": [str(p)[:200] for p in problems][:6]}, None

    answer = call_llm_json("final_review", REVIEW_SYSTEM, user, validate, llm=llm, phase="review")
    return answer or {"verdict": "не получен", "problems": []}


def apply_final_review(table: pd.DataFrame, report: dict[str, Any], llm: Any) -> dict[str, Any]:
    report = dict(report)
    errors = list(report.get("errors") or [])
    bad = owner_problems(report.get("owner") or {})
    review: dict[str, Any] = {"owner_problems": bad}
    if llm is not None and len(table):
        review.update(llm_final_review(table, report, llm))
    suspicious = bad or review.get("verdict") == "suspect"
    if review.get("verdict") == "suspect":
        errors.append("итоговая проверка моделью: " + ("; ".join(review["problems"]) or "результат вызывает сомнения"))
    if suspicious and report.get("status") != "needs_review":
        report["status"] = "needs_review"
    if llm is not None and len(table) and review.get("verdict") == "не получен":
        report["notes"] = list(report.get("notes", [])) + ["итоговая проверка моделью не выполнена (нет ответа или ответ не по формату)"]
    report["errors"] = errors
    report["review"] = review
    log("проверка", f"итоговая проверка: {review.get('verdict', 'без модели')}"
                    + (f"; владелец: {'; '.join(bad)}" if bad else "")
                    + (f"; замечания: {'; '.join(review.get('problems') or [])}" if review.get("problems") else ""))
    return report


def parse_statement_auto(path: str | Path, llm: Any | None = None, agent_llm: Any | None = None,
                         use_llm: bool = True) -> tuple[pd.DataFrame, dict[str, Any]]:
    llm = (llm if llm is not None else LLM_CLIENT) if use_llm else None
    table, report = parse_statement(path, llm=llm, use_llm=use_llm)
    agent_llm = agent_llm if agent_llm is not None else llm
    if isinstance(agent_llm, CachedLLM):
        agent_llm = agent_llm.client
    if report["status"] == "needs_review" and CLOCK.unavailable:
        log("агент", f"не запускаю: GigaChat недоступен ({CLOCK.unavailable}); результат алгоритма оставлен со статусом needs_review")
        report["agent"] = {"used": False, "skipped": CLOCK.unavailable}
    elif report["status"] == "needs_review" and AGENT_ENABLED and AGENT_AVAILABLE and agent_llm is not None:
        log("агент", f"самопроверка не пройдена ({'; '.join(report['errors'][:3])}) — запускаю агента")
        CLOCK.mark("этап агента", "алгоритм не справился — подключён агент")
        table, report = run_agent(path, table, report, agent_llm)
        CLOCK.mark("конец", f"всего {CLOCK.elapsed():.0f} с, статус {report['status']}")
    else:
        report["agent"] = {"used": False}
    report = apply_final_review(table, report, llm)
    return table, report

