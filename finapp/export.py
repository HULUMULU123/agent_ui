"""Выгрузка результата в Excel: выписка с цветовой разметкой, контрагенты, статистика (формат ноутбука агента)."""
from __future__ import annotations

import math
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .analytics import (CONFIDENCE_LABELS, SOURCE_LABELS, TOOL_LABELS, counterparty_table, kpis,
                        LEVEL_LABELS)
from .store import REVIEW_DECISIONS

RISK_LABELS = {0: "Риска нет", 1: "Требует проверки"}
FINAL_COLUMNS = {
    "idx": "ID операции", "date": "Дата", "debit_inn": "ИНН плательщика", "debit_name": "Плательщик",
    "credit_inn": "ИНН получателя", "credit_name": "Получатель", "direction": "Направление", "amount": "Сумма, руб.",
    "purpose": "Назначение платежа", "period": "Период до заявления", "operation_type": "Тип операции",
    "status_label": "Статус", "risk_level": "Риск (0/1)", "risk_factors": "Факторы риска",
    "counterparty_category": "Категория контрагента", "legal_qualification": "Маршрут проверки",
    "legal_basis": "Правовые основания", "court_basis": "Судебная практика",
    "challenge_criteria": "Что установить", "connections_basis": "Связи сторон",
    "decision_argumentation": "Обоснование", "risk_explanation": "Пояснение по риску",
    "recommendation": "Рекомендация", "recommendation_agent": "Рекомендация агента (до правки)", "requested_documents": "Документы к запросу",
    "used_tools": "Использованные источники", "analysis_source": "Источник анализа",
    "propagation_confidence": "Уверенность переноса", "matched_representative_idx": "Перенесено с операции",
    "review_decision": "Решение сотрудника", "review_comment": "Комментарий сотрудника",
}
COLUMN_WIDTH_LIMITS = {
    "Назначение платежа": (30, 60), "Обоснование": (40, 80), "Пояснение по риску": (30, 60),
    "Рекомендация": (30, 70), "Документы к запросу": (30, 60), "Правовые основания": (30, 60),
    "Судебная практика": (30, 60), "Что установить": (30, 60), "Факторы риска": (22, 45),
    "Маршрут проверки": (22, 45), "Рекомендация агента (до правки)": (30, 70), "Связи сторон": (18, 40), "Комментарий сотрудника": (24, 50),
}
DEFAULT_WIDTH_LIMITS = (10, 30)
STATUS_FILL = {"Потенциально оспоримая": "FFF8CBAD", "Требует проверки": "FFFFE699",
               "Не классифицирована": "FFEDEDED", "Типовая операция": "FFE2F0D9"}
HEADER_FILL = "FF1F4E78"
THIN = Side(style="thin", color="FF9E9E9E")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LINE_HEIGHT_PT = 15
MAX_ROW_HEIGHT_PT = 409


def _humanize(value: Any) -> Any:
    if isinstance(value, (list, tuple)):
        return "\n".join(f"• {v}" for v in value if v is not None and str(v) != "")
    if isinstance(value, pd.Timestamp):
        return None if pd.isna(value) else value.to_pydatetime().date()
    if isinstance(value, (bool, np.bool_)):
        return "Да" if value else "Нет"
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def build_statement(df: pd.DataFrame, reviews: dict[str, dict] | None = None) -> pd.DataFrame:
    out = df.copy()
    reviews = reviews or {}
    out["review_decision"] = out["idx"].map(lambda i: REVIEW_DECISIONS.get(reviews.get(str(i), {}).get("decision", ""), "")
                                            if str(i) in reviews else "")
    out["review_comment"] = out["idx"].map(lambda i: reviews.get(str(i), {}).get("comment", ""))
    edited = out["idx"].map(lambda i: reviews.get(str(i), {}).get("recommendation", ""))
    out["recommendation_agent"] = out["recommendation"].where(edited.astype(bool), "")
    out["recommendation"] = edited.where(edited.astype(bool), out["recommendation"])
    out["analysis_source"] = out["analysis_source"].map(SOURCE_LABELS).fillna(out["analysis_source"])
    propagated = out["analysis_source"].eq(SOURCE_LABELS["propagated"])
    out["propagation_confidence"] = out["propagation_confidence"].map(CONFIDENCE_LABELS).where(propagated)
    out["matched_representative_idx"] = out["matched_representative_idx"].where(propagated)
    out["used_tools"] = out["used_tools"].map(lambda v: [TOOL_LABELS.get(t, t) for t in v] if isinstance(v, list) else v)
    out["amount"] = pd.to_numeric(out["amount"], errors="coerce")
    for col in ("idx", "debit_inn", "credit_inn"):
        out[col] = out[col].astype("string")
    out = out[[c for c in FINAL_COLUMNS if c in out.columns]]
    out = out.astype(object).where(out.notna(), None).map(_humanize)
    return out.rename(columns=FINAL_COLUMNS)


def _text_lines(value: Any, width: int) -> int:
    if value is None:
        return 1
    text = value.strftime("%d.%m.%Y") if isinstance(value, (date, datetime)) else str(value)
    return sum(max(1, math.ceil(len(line) / max(width - 1, 1))) for line in text.split("\n"))


def _format_statement_sheet(ws, statement: pd.DataFrame) -> None:
    headers = list(statement.columns)
    status_col = headers.index(FINAL_COLUMNS["status_label"]) + 1 if FINAL_COLUMNS["status_label"] in headers else None
    wrap_top = Alignment(wrap_text=True, vertical="top")
    widths = {}
    for j, name in enumerate(headers, start=1):
        low, high = COLUMN_WIDTH_LIMITS.get(name, DEFAULT_WIDTH_LIMITS)
        longest = max([len(name)] + [max((len(line) for line in str(v).split("\n")), default=0)
                                     for v in statement[name] if v is not None])
        widths[j] = min(max(longest + 2, low), high)
        ws.column_dimensions[get_column_letter(j)].width = widths[j]
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
    ws.row_dimensions[1].height = LINE_HEIGHT_PT * max(_text_lines(c.value, widths[c.column]) for c in ws[1])
    number_formats = {FINAL_COLUMNS["date"]: "DD.MM.YYYY", FINAL_COLUMNS["amount"]: "#,##0.00"}
    for row in ws.iter_rows(min_row=2):
        fill_color = STATUS_FILL.get(row[status_col - 1].value) if status_col else None
        lines = 1
        for cell in row:
            name = headers[cell.column - 1]
            cell.border = BORDER
            cell.alignment = wrap_top
            if name in number_formats:
                cell.number_format = number_formats[name]
            elif name in (FINAL_COLUMNS["idx"], FINAL_COLUMNS["debit_inn"], FINAL_COLUMNS["credit_inn"]):
                cell.number_format = "@"
            if fill_color:
                cell.fill = PatternFill("solid", fgColor=fill_color)
            lines = max(lines, _text_lines(cell.value, widths[cell.column]))
        ws.row_dimensions[row[0].row].height = min(LINE_HEIGHT_PT * lines, MAX_ROW_HEIGHT_PT)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def _format_plain_sheet(ws) -> None:
    for cell in ws[1]:
        if cell.value is not None:
            cell.font = Font(bold=True, color="FFFFFFFF")
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for j in range(1, ws.max_column + 1):
        letter = get_column_letter(j)
        longest = max((len(str(c.value)) for c in ws[letter] if c.value is not None), default=10)
        ws.column_dimensions[letter].width = min(max(longest + 2, 12), 55)
    ws.freeze_panes = "A2"


def _counterparties_sheet(df: pd.DataFrame, risk_db: Any | None) -> pd.DataFrame:
    table = counterparty_table(df, risk_db)
    if table.empty:
        return pd.DataFrame({"Контрагенты": ["Не определены"]})
    return pd.DataFrame({
        "ИНН контрагента": table["inn"], "Наименование": table["name"], "Операций": table["n_ops"],
        "Оборот (всего)": table["turnover"].round(2), "Исходящие (убытие)": table["outflow"].round(2),
        "Входящие (прибытие)": table["inflow"].round(2), "Потенциально оспоримых": table["n_disputable"],
        "Требуют проверки": table["n_attention"], "Статус риска": table["level"].map(LEVEL_LABELS),
        "Историческая память: уровень": table["history_grade"],
        "Первая операция": pd.to_datetime(table["first"]).dt.date, "Последняя операция": pd.to_datetime(table["last"]).dt.date,
    })


def _overview_blocks(df: pd.DataFrame) -> list[tuple[str, pd.DataFrame]]:
    k = kpis(df)
    kpi = pd.DataFrame({"Показатель": ["Период выписки", "Всего операций", "Уникальных контрагентов", "Поступления", "Списания",
                                       "Потенциально оспоримых", "Требуют проверки", "Не классифицировано", "Сумма потенциально оспоримых"],
                        "Значение": [f"{pd.Timestamp(k['period_start']):%d.%m.%Y} — {pd.Timestamp(k['period_end']):%d.%m.%Y}"
                                     if pd.notna(k["period_start"]) else "—", k["total"], k["counterparties"],
                                     round(k["inflow"], 2), round(k["outflow"], 2), k["disputable"], k["attention"],
                                     k["unclassified"], round(k["risky_amount"], 2)]})
    by_type = (df.groupby("operation_type", dropna=False)
               .agg(Операций=("idx", "count"), Сумма=("amount_abs", lambda s: round(s.sum(), 2)),
                    **{"Из них с риском": ("risk_level", lambda s: int((s == 1).sum()))})
               .reset_index().rename(columns={"operation_type": "Тип операции"}).sort_values("Операций", ascending=False))
    by_status = df["status_label"].value_counts().rename_axis("Статус").reset_index(name="Операций")
    by_source = df["source_label"].value_counts().rename_axis("Источник анализа").reset_index(name="Операций")
    return [("Ключевые показатели", kpi), ("По типу операции", by_type), ("По статусу", by_status),
            ("По источнику анализа", by_source)]


def export_workbook(df: pd.DataFrame, path: Path, *, reviews: dict[str, dict] | None = None, risk_db: Any | None = None,
                    only_idx: set[str] | None = None, title: str = "Выписка") -> Path:
    """df — результат analytics.enrich(); only_idx ограничивает лист «Выписка» (отчёт по выбранным операциям)."""
    sheet_df = df if only_idx is None else df[df["idx"].astype(str).isin(only_idx)]
    statement = build_statement(sheet_df, reviews)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        statement.to_excel(writer, sheet_name=title[:31], index=False)
        _counterparties_sheet(df, risk_db).to_excel(writer, sheet_name="Контрагенты", index=False)
        row = 0
        for block_title, block in _overview_blocks(df):
            if block.empty:
                continue
            pd.DataFrame({block_title: []}).to_excel(writer, sheet_name="Статистика по выписке", index=False, startrow=row)
            block.to_excel(writer, sheet_name="Статистика по выписке", index=False, startrow=row + 1)
            row += len(block) + 3
    wb = load_workbook(path)
    _format_statement_sheet(wb[title[:31]], statement)
    _format_plain_sheet(wb["Контрагенты"])
    ws = wb["Статистика по выписке"]
    for row_cells in ws.iter_rows():
        first = next((c for c in row_cells if c.value is not None), None)
        if first is not None and first.column == 1 and all(c.value is None for c in row_cells[1:]):
            first.font = Font(bold=True, size=12)
    for col in ("A", "B", "C", "D"):
        ws.column_dimensions[col].width = 32
    wb.save(path)
    return path


def export_to_temp(df: pd.DataFrame, name: str, **kwargs: Any) -> str:
    safe = "".join(ch if ch.isalnum() or ch in "-_." else "_" for ch in name)[:80] or "report"
    directory = Path(tempfile.mkdtemp(prefix="finapp_export_"))
    return str(export_workbook(df, directory / f"{safe}.xlsx", **kwargs))
